"""تقرير جرد الأصول الثابتة (Excel + PDF): ملخص + العجز + الزيادة (+ الموجود اختياريًا)."""
import datetime
import os

from .asset_labels import _register_fonts
from .invoice_pdf import _shape_ar

STATE_LABELS = {"open": "مفتوح", "closed": "مقفول"}
METHOD_LABELS = {"barcode": "باركود", "manual": "تأكيد يدوي"}
REASON_LABELS = {
    "out_of_scope": "أصل مسجّل لكن خارج نطاق الجرد",
    "disposed": "أصل مستبعد من السجلات لكنه موجود فعليًا",
    "unknown": "كود غير مسجّل في النظام",
}
COLORS = {"shortage": "C0392B", "surplus": "D68910", "found": "1E8449", "summary": "1F2937"}


def _fmt_dt(d):
    return d.strftime("%Y-%m-%d %H:%M") if d else "-"


def _summary_pairs(rep):
    inv, sm = rep["inventory"], rep["summary"]
    pairs = [
        ("رقم الجرد", inv["number"]), ("اسم الجرد", inv["name"]), ("تاريخ البدء", _fmt_dt(inv["date"])),
        ("النطاق", inv["scope_label"]), ("الحالة", STATE_LABELS.get(inv["state"], inv["state"])),
        ("أُنشئ بواسطة", inv["created_by"] or "-"),
        ("عدد الأصول المتوقعة (حسب السجلات)", sm["expected"]),
        ("تم العثور عليها", f'{sm["found"]} من {sm["expected"]} - نسبة {sm["match_pct"]:g}%'),
        ("منها مؤكّدة يدويًا (بدون باركود)", sm["manual_confirmed"]),
        ("العجز - أصول مسجّلة ولم يتم العثور عليها", sm["shortage_count"]),
        ("قيمة العجز بالتكلفة", sm["shortage_cost"]),
        ("قيمة العجز بصافي القيمة الدفترية", sm["shortage_book_value"]),
        ("الزيادة - أصول موجودة وغير متوقعة", sm["surplus_count"]),
        ("   منها أكواد غير مسجّلة في النظام", sm["surplus_unknown"]),
        ("   منها أصول مسجّلة خارج النطاق/مستبعدة", sm["surplus_known"]),
    ]
    if sm.get("disposed_count"):
        pairs += [("تم استبعاد أصول العجز بقيد محاسبي", f'{sm["disposed_count"]} أصل - خسارة {sm["disposed_loss"]:,.2f}'),
                  ("عجز متبقٍ لم يُستبعد", f'{sm["pending_count"]} أصل')]
    return pairs


# ------------------------------------------------------------------ Excel
def export_inventory_excel(rep, company, output_dir, include_found=False):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    os.makedirs(output_dir, exist_ok=True)
    stamp = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
    path = os.path.join(output_dir, f"asset_inventory_{rep['inventory']['number']}_{stamp}.xlsx")
    thin = Side(style="thin", color="D0D0D8")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    wb = Workbook()

    def header(ws, row, titles, color):
        for c, t in enumerate(titles, 1):
            cell = ws.cell(row, c, t)
            cell.font = Font(bold=True, color="FFFFFF"); cell.fill = PatternFill("solid", fgColor=color)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True); cell.border = box

    def title_block(ws, subtitle, ncols):
        ws.sheet_view.rightToLeft = True
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)
        ws.cell(1, 1, company.name if company and company.name else "اسم الشركة").font = Font(size=16, bold=True)
        ws.cell(1, 1).alignment = Alignment(horizontal="center")
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ncols)
        ws.cell(2, 1, subtitle).font = Font(size=13, bold=True)
        ws.cell(2, 1).alignment = Alignment(horizontal="center")
        ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=ncols)
        ws.cell(3, 1, f'{rep["inventory"]["number"]} - {rep["inventory"]["name"]}   |   {rep["inventory"]["scope_label"]}   |   تاريخ الطباعة: {datetime.datetime.now():%Y-%m-%d %H:%M}')
        ws.cell(3, 1).alignment = Alignment(horizontal="center")

    def widths(ws, ws_widths):
        for i, w in enumerate(ws_widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w

    # الملخص
    ws = wb.active; ws.title = "الملخص"
    title_block(ws, "ملخص جرد الأصول الثابتة", 2)
    header(ws, 5, ["البند", "القيمة"], COLORS["summary"])
    for i, (k, v) in enumerate(_summary_pairs(rep), 6):
        a, b = ws.cell(i, 1, k), ws.cell(i, 2, v)
        a.border = b.border = box; a.alignment = Alignment(horizontal="right"); b.alignment = Alignment(horizontal="center")
        if isinstance(v, float):
            b.number_format = "#,##0.00"
        if k.startswith("العجز") or k.startswith("قيمة العجز"):
            a.font = b.font = Font(bold=True, color=COLORS["shortage"])
        elif k.startswith("الزيادة"):
            a.font = b.font = Font(bold=True, color=COLORS["surplus"])
    widths(ws, [52, 30])

    def table_sheet(name, subtitle, titles, rows, color, money_cols=(), total_cols=()):
        w = wb.create_sheet(name)
        title_block(w, subtitle, len(titles))
        header(w, 5, titles, color)
        for r, row in enumerate(rows, 6):
            for c, v in enumerate(row, 1):
                cell = w.cell(r, c, v); cell.border = box
                cell.alignment = Alignment(horizontal="center" if c in money_cols else "right")
                if c in money_cols:
                    cell.number_format = "#,##0.00"
        if rows:
            tr = 6 + len(rows)
            w.cell(tr, 1, f"الإجمالي ({len(rows)})").font = Font(bold=True)
            for c in total_cols:
                cell = w.cell(tr, c, round(sum(r[c - 1] for r in rows), 2)); cell.font = Font(bold=True); cell.number_format = "#,##0.00"
            for c in range(1, len(titles) + 1):
                w.cell(tr, c).border = box; w.cell(tr, c).fill = PatternFill("solid", fgColor="F0F0F5")
        else:
            w.cell(6, 1, "لا يوجد").font = Font(italic=True, color="808080")
        w.freeze_panes = "A6"
        return w

    w = table_sheet("العجز", "العجز - أصول مسجّلة ولم يتم العثور عليها",
                    ["الكود", "اسم الأصل", "الفئة", "التكلفة", "صافي القيمة الدفترية", "الاستبعاد"],
                    [[r["code"], r["name"], r["category"], r["cost"], round(r["book_value"], 2),
                      (f'مستبعد ({r["disposal_ref"]}) - خسارة {r["disposal_loss"]:,.2f}' if r["disposed_at"] else "-")] for r in rep["shortage"]],
                    COLORS["shortage"], money_cols=(4, 5), total_cols=(4, 5))
    widths(w, [16, 38, 22, 16, 22, 40])
    w = table_sheet("الزيادة", "الزيادة - أصول موجودة وغير متوقعة في الجرد",
                    ["الكود", "اسم الأصل", "السبب", "مرات المسح"],
                    [[r["code"], r["name"] or "-", REASON_LABELS.get(r["reason"], ""), r["scan_count"]] for r in rep["surplus"]],
                    COLORS["surplus"])
    widths(w, [18, 38, 44, 14])
    if include_found:
        w = table_sheet("الموجود", "أصول تم العثور عليها",
                        ["الكود", "اسم الأصل", "الفئة", "طريقة الإثبات", "وقت المسح"],
                        [[r["code"], r["name"], r["category"], METHOD_LABELS.get(r["method"], "-"), _fmt_dt(r["scanned_at"])] for r in rep["found"]],
                        COLORS["found"])
        widths(w, [16, 38, 22, 16, 22])
    wb.save(path)
    return path


# ------------------------------------------------------------------ PDF
def export_inventory_pdf(rep, company, output_dir, include_found=False):
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER, TA_RIGHT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle, KeepTogether

    os.makedirs(output_dir, exist_ok=True)
    stamp = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
    path = os.path.join(output_dir, f"asset_inventory_{rep['inventory']['number']}_{stamp}.pdf")
    font, bold = _register_fonts()
    accent = getattr(company, "print_color", None) or "#6C4EE3"
    P = lambda t, st: Paragraph(_shape_ar(t) if not str(t).isascii() else str(t), st)

    title = ParagraphStyle("t", fontName=bold, fontSize=17, alignment=TA_CENTER, leading=22)
    sub = ParagraphStyle("s", fontName=bold, fontSize=12, alignment=TA_CENTER, leading=16)
    small = ParagraphStyle("sm", fontName=font, fontSize=8.5, alignment=TA_CENTER, textColor=colors.grey, leading=11)
    h2 = ParagraphStyle("h2", fontName=bold, fontSize=11.5, alignment=TA_RIGHT, leading=15)
    cell_r = ParagraphStyle("cr", fontName=font, fontSize=8.5, alignment=TA_RIGHT, leading=11)
    cell_c = ParagraphStyle("cc", fontName=font, fontSize=8.5, alignment=TA_CENTER, leading=11)
    head = ParagraphStyle("hd", fontName=bold, fontSize=8.5, alignment=TA_CENTER, textColor=colors.white, leading=11)
    bold_c = ParagraphStyle("bc", fontName=bold, fontSize=8.5, alignment=TA_CENTER, leading=11)
    bold_r = ParagraphStyle("br", fontName=bold, fontSize=8.5, alignment=TA_RIGHT, leading=11)

    def footer(canvas, doc):
        canvas.saveState(); canvas.setFont(font, 8); canvas.setFillGray(0.4)
        canvas.drawString(28, 16, f"Page {doc.page}")
        canvas.drawRightString(A4[0] - 28, 16, _shape_ar(f'جرد الأصول {rep["inventory"]["number"]} - {datetime.datetime.now():%Y-%m-%d %H:%M}'))
        canvas.restoreState()

    def section(name, color, titles, rows, widths, aligns, total=None):
        """rows: قايمة صفوف منطقية (الأول = اليمين) - بتتعكس عشان الجدول RTL."""
        out = [Spacer(1, 10), P(name, h2), Spacer(1, 4)]
        if not rows:
            out.append(P("لا يوجد", cell_r)); return out
        data = [[P(t, head) for t in titles[::-1]]]
        for r in rows:
            data.append([P(v, cell_c if al == "c" else cell_r) for v, al in zip(r[::-1], aligns[::-1])])
        if total:
            data.append([P(v, bold_c if al == "c" else bold_r) for v, al in zip(total[::-1], aligns[::-1])])
        t = Table(data, colWidths=widths[::-1], repeatRows=1, hAlign="RIGHT")
        style = [("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#" + color)), ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#D0D0D8")),
                 ("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
                 ("ROWBACKGROUNDS", (0, 1), (-1, -1 if not total else -2), [colors.white, colors.HexColor("#F7F6FC")])]
        if total:
            style += [("BACKGROUND", (0, -1), (-1, -1), colors.HexColor("#ECECF3")), ("LINEABOVE", (0, -1), (-1, -1), 0.8, colors.HexColor("#" + color))]
        t.setStyle(TableStyle(style)); out.append(t); return out

    inv, sm = rep["inventory"], rep["summary"]
    story = [P(company.name if company and company.name else "اسم الشركة", title), P("تقرير جرد الأصول الثابتة", sub),
             P(f'{inv["number"]} - {inv["name"]}', small), Spacer(1, 10)]
    info = _summary_pairs(rep)
    data = []
    for k, v in info:
        style_v = cell_c
        if k.startswith("العجز") or k.startswith("قيمة العجز"):
            k_st, v_st = ParagraphStyle("k1", parent=bold_r, textColor=colors.HexColor("#C0392B")), ParagraphStyle("v1", parent=bold_c, textColor=colors.HexColor("#C0392B"))
        elif k.startswith("الزيادة"):
            k_st, v_st = ParagraphStyle("k2", parent=bold_r, textColor=colors.HexColor("#D68910")), ParagraphStyle("v2", parent=bold_c, textColor=colors.HexColor("#D68910"))
        else:
            k_st, v_st = cell_r, style_v
        vv = f"{v:,.2f}" if isinstance(v, float) else v
        data.append([P(str(vv), v_st), P(str(k), k_st)])
    st = Table(data, colWidths=[200, 339], hAlign="RIGHT")
    st.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#D0D0D8")), ("BACKGROUND", (1, 0), (1, -1), colors.HexColor("#F7F6FC")),
                            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3)]))
    story.append(st)

    story += section(f'العجز - أصول مسجّلة ولم يتم العثور عليها ({sm["shortage_count"]})', COLORS["shortage"],
                     ["الكود", "اسم الأصل", "الفئة", "التكلفة", "القيمة الدفترية", "الاستبعاد"],
                     [[r["code"], r["name"], r["category"] or "-", f'{r["cost"]:,.2f}', f'{r["book_value"]:,.2f}',
                       (f'مستبعد {r["disposal_ref"]}' if r["disposed_at"] else "-")] for r in rep["shortage"]],
                     [68, 150, 82, 72, 82, 85], ["r", "r", "r", "c", "c", "c"],
                     total=["الإجمالي", "", "", f'{sm["shortage_cost"]:,.2f}', f'{sm["shortage_book_value"]:,.2f}', ""] if rep["shortage"] else None)
    story += section(f'الزيادة - أصول موجودة وغير متوقعة ({sm["surplus_count"]})', COLORS["surplus"],
                     ["الكود", "اسم الأصل", "السبب", "مرات المسح"],
                     [[r["code"], r["name"] or "-", REASON_LABELS.get(r["reason"], ""), str(r["scan_count"])] for r in rep["surplus"]],
                     [90, 150, 230, 69], ["r", "r", "r", "c"])
    if include_found:
        story += section(f'أصول تم العثور عليها ({sm["found"]})', COLORS["found"],
                         ["الكود", "اسم الأصل", "الفئة", "الإثبات", "وقت المسح"],
                         [[r["code"], r["name"], r["category"] or "-", METHOD_LABELS.get(r["method"], "-"), _fmt_dt(r["scanned_at"])] for r in rep["found"]],
                         [80, 180, 100, 70, 109], ["r", "r", "r", "c", "c"])
    sig = Table([[P("المدير المالي", bold_c), P("مراجعة", bold_c), P("إعداد لجنة الجرد", bold_c)],
                 [P("", cell_c), P("", cell_c), P("", cell_c)]], colWidths=[180, 180, 179], rowHeights=[16, 34], hAlign="RIGHT")
    sig.setStyle(TableStyle([("LINEBELOW", (0, 1), (-1, 1), 0.6, colors.grey)]))
    story += [Spacer(1, 22), KeepTogether([sig])]
    SimpleDocTemplate(path, pagesize=A4, rightMargin=28, leftMargin=28, topMargin=28, bottomMargin=30,
                      title="جرد الأصول").build(story, onFirstPage=footer, onLaterPages=footer)
    return path
