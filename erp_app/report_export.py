"""تصدير التقارير الرسمية إلى Excel وPDF مع رأس الشركة والشعار والتاريخ."""
import os
import datetime


def _cost_ok():
    """هل المستخدم الحالي مسموح له يشوف التكلفة/قيم المخزون؟ (صلاحيات الحقول) - لو لأ بتتفضّى أعمدة التكلفة في التصدير."""
    try:
        from . import field_access
        return field_access.can_see("cost_price")
    except Exception:
        return True


def _company():
    try:
        from .services import get_company_settings
        return get_company_settings()
    except Exception:
        return None


def _fmt_date(v):
    if not v:
        return ""
    if hasattr(v, "strftime"):
        return v.strftime("%Y-%m-%d")
    return str(v)


def _fmt(v):
    if isinstance(v, (int, float)):
        return f"{(v if v else 0.0):,.2f}"
    return _fmt_date(v) if hasattr(v, "strftime") else str(v or "")


_NO_MAIN_MARK = {'stock','payroll','allowances_insurance','depreciation','sales','purchases','receipts_payments','custody_statement','vat_return','vat_return_summary','vat_sales_details','vat_purchases_details','hr_documents','prepaid_schedule','expiry_report','purchase_price_comparison','stock_balance','inventory_adjustment','pos_session','sales_analysis','purchase_analysis','kitchen_events','void_report','audit_log','recipe_costing','restaurant_dashboard','discount_report','data_integrity','receivables_aging'}


def _main_row_flags(raw, report_type):
    """لكل صف بيانات (بعد صف العنوان) هل هو حساب رئيسي؟ بنقارن أي خلية
    نصية بأكواد الحسابات الرئيسية (الخلية = الكود بالظبط أو بتبدأ بـ'الكود - ')."""
    if report_type in _NO_MAIN_MARK:
        return [False] * len(raw)
    try:
        from .services import main_account_codes
        codes = main_account_codes()
    except Exception:
        codes = set()
    flags = [False]
    for row in raw[1:]:
        hit = False
        for cell in row:
            if isinstance(cell, str) and cell:
                c = cell.strip()
                if c in codes or any(c.startswith(code + ' - ') for code in codes):
                    hit = True; break
        flags.append(hit)
    return flags


_ROW_BG = {'section': 'E5E7EB', 'total': 'EEF2FF', 'net': 'FEF3C7', 'warn': 'FEE2E2'}


def _rows(data, report_type):
    if report_type in ('sales_analysis', 'purchase_analysis', 'kitchen_events', 'void_report', 'audit_log', 'recipe_costing', 'restaurant_dashboard', 'discount_report', 'data_integrity', 'receivables_aging'):
        return data['table']      # الجدول جاهز (عناوين + صفوف + إجمالي) من شاشة التقارير، مع row_styles وcol_weights
    if report_type == 'pos_session':
        d = data['session']
        t = d['totals']
        fmt = lambda dt_: dt_.strftime('%Y-%m-%d %H:%M') if dt_ else '-'
        rows = [['البند', 'التفاصيل', 'الكمية / العدد', 'قبل الضريبة', 'الضريبة', 'الإجمالي']]
        styles = {}
        def section(title):
            rows.append([title, '', '', '', '', '']); styles[len(rows) - 1] = 'section'
        section('بيانات الجلسة')
        rows += [['الماكينة', d['terminal'], '', '', '', ''], ['الكاشير', d['employee'] or d['opened_by_name'], '', '', '', ''], ['فُتحت', fmt(d['opened_at']) + f" - {d['opened_by_name']}", '', '', '', ''],
                 ['أُغلقت', (fmt(d['closed_at']) + f" - {d['closed_by']}") if d['state'] == 'closed' else 'لسه مفتوحة', '', '', '', ''], ['العهدة الافتتاحية', '', '', '', '', d['opening_cash']]]
        section('ملخص المبيعات')
        rows.append(['فواتير غير ملغاة', '', t['count'], t['subtotal'], t['tax'], t['total']])
        if t['voided_count']:
            rows.append(['فواتير ملغاة (مستبعدة)', '', t['voided_count'], '', '', t['voided_total']])
        rt = d['returns_totals']
        if rt['count']:
            rows.append(['مرتجعات على فواتير الجلسة', '', rt['count'], rt['subtotal'], rt['tax'], rt['total']]); styles[len(rows) - 1] = 'warn'
            rows.append(['صافي المبيعات بعد المرتجعات', '', '', d['net_subtotal'], d['net_tax'], d['net_total']]); styles[len(rows) - 1] = 'net'
        section('طرق الدفع')
        for m in d['payment_methods']:
            rows.append([m['name'] + (' (نقدي)' if m['is_cash'] else ''), '', m['count'], '', '', m['total']])
        section('جرد النقدية')
        rows += [['العهدة الافتتاحية', '', '', '', '', d['opening_cash']], ['المبيعات النقدية', '', '', '', '', d['cash_sales']]]
        if d['cash_refunds']:
            rows.append(['(-) مرتجعات مردودة نقدًا من الدرج', '', '', '', '', -d['cash_refunds']])
        rows.append(['النقدية المتوقعة', '', '', '', '', d['expected_cash_final']])
        if d['state'] == 'closed':
            rows += [['النقدية المعدودة', '', '', '', '', d['counted_cash']], ['الفرق (+ زيادة / - عجز)', d['close_notes'] or '', '', '', '', d['difference']]]
            styles[len(rows) - 1] = 'warn' if abs(d['difference'] or 0) > 0.005 else 'net'
        section(f"مبيعات الأصناف ({len(d['items'])} صنف) - الكمية المباعة / المرتجع / الصافي")
        for i in d['items']:
            rows.append([f"{i['sku']} - {i['name']}", f"{i['category']}  |  مباع {i['qty']:g}  مرتجع {i['returned_qty']:g}  صافي {i['net_qty']:g}", i['net_qty'], i['net'], i['tax'], i['net_total']])
        section(f"فواتير الجلسة ({len(d['invoices'])})")
        label = {'full': ' (مرتجعة كليًا)', 'partial': ' (مرتجع جزئي)', 'none': ''}
        for inv in d['invoices']:
            rows.append([inv['number'] + (' (ملغاة)' if inv['voided'] else label.get(inv.get('return_status', 'none'), '')), f"{inv['customer']} - {inv['payment_method']} - {inv['date']:%H:%M}", '', inv['subtotal'], inv['tax'], inv['total']])
            if inv['voided'] or inv.get('return_status') == 'full':
                styles[len(rows) - 1] = 'warn'
            elif inv.get('return_status') == 'partial':
                styles[len(rows) - 1] = 'net'
        if d['returns']:
            section(f"المرتجعات ({len(d['returns'])})")
            for r in d['returns']:
                rows.append([f"{r['number']} على {r['invoice_number']}" + (' (كلي)' if r['is_full'] else ' (جزئي)'), f"{r['customer']} - {r['refund_method']} - {r['reason']}", '', r['subtotal'], r['tax'], r['total']])
            section('الأصناف المرتجعة')
            for i in d['returned_items']:
                rows.append([f"{i['sku']} - {i['name']}", i['category'], i['qty'], '', '', i['total']])
        if d.get('cancellations'):
            section(f"إلغاءات المطعم ({d['cancel_totals']['count']} بند - لم تدخل في المبيعات)")
            for c in d['cancellations']:
                rows.append([f"{c['kind_label']} - {c['order_number']} ({c['where']})", f"{c['product']} × {c['quantity']:g} - {c['reason']} - {c['by']}" + (f" / اعتماد {c['approved_by']}" if c['approved_by'] else ''), c['quantity'], '', '', c['amount']]); styles[len(rows) - 1] = 'warn'
        rows.append(['إجمالي المبيعات', '', t['count'], t['subtotal'], t['tax'], t['total']]); styles[len(rows) - 1] = 'total'
        data['row_styles'] = styles
        data['col_weights'] = [3.0, 3.0, 1.2, 1.4, 1.2, 1.5]
        return rows
    if report_type == 'inventory_adjustment':
        d = data['adjustment']
        rows = [['الكود', 'الصنف', 'الوحدة', 'الرصيد الدفتري', 'الرصيد الفعلي', 'الفرق', 'تكلفة الوحدة', 'قيمة الفرق', 'النوع']]
        styles = {}
        for l in d['lines']:
            rows.append([l['sku'], l['name'], l['unit'], l['theoretical'], l['counted'], l['difference'], l['unit_cost'] if _cost_ok() else None, l['value'] if _cost_ok() else None, 'زيادة' if l['difference'] > 0 else 'عجز'])
            styles[len(rows) - 1] = 'net' if l['difference'] < 0 else 'normal'
        rows.append(['', f"الإجمالي ({len(d['lines'])} صنف)", '', '', '', '', f"زيادة {d['surplus_value']:,.2f}" if _cost_ok() else '', d['net_value'] if _cost_ok() else None, f"عجز {d['shortage_value']:,.2f}" if _cost_ok() else ''])
        styles[len(rows) - 1] = 'total'
        data['row_styles'] = {k: v for k, v in styles.items() if v != 'normal'}
        data['col_weights'] = [1.1, 3.0, 0.9, 1.3, 1.3, 1.1, 1.3, 1.4, 1.4]
        return rows
    if report_type == 'stock_balance':
        rows_in = data.get('stock_rows', [])
        by_wh = bool(data.get('stock_by_warehouse'))
        head = ['الكود', 'الصنف'] + (['المخزن'] if by_wh else []) + ['الوحدة', 'الرصيد', 'تكلفة الوحدة', 'قيمة الرصيد', 'حد إعادة الطلب', 'الحالة']
        rows = [head]
        styles = {}
        for r in rows_in:
            p = r['product']
            status = 'رصيد سالب' if r['balance'] < -1e-9 else ('ناقص - يحتاج طلب' if r['is_low'] else 'متوفر')
            rows.append([p.sku, p.name] + ([r.get('location_name', '')] if by_wh else []) + [p.unit or '', r['balance'], (p.cost_price or 0.0) if _cost_ok() else None, r['value'] if _cost_ok() else None, p.reorder_point or 0.0, status])
            if r['balance'] < -1e-9:
                styles[len(rows) - 1] = 'warn'
            elif r['is_low']:
                styles[len(rows) - 1] = 'net'
        total_value = sum(r['value'] for r in rows_in) if _cost_ok() else None
        rows.append(['', f'الإجمالي ({len(rows_in)} صنف)'] + ([''] if by_wh else []) + ['', '', '', total_value, '', ''])
        styles[len(rows) - 1] = 'total'
        data['row_styles'] = styles
        data['col_weights'] = ([1.2, 3.0] + ([1.8] if by_wh else []) + [1.0, 1.2, 1.2, 1.5, 1.3, 1.8])
        return rows
    if report_type == 'purchase_price_comparison':
        r = data['ppc']
        by_supplier = r['group_by'] == 'supplier'
        head = ['الفئة / الصنف'] + list(r['columns']) + (['أرخص مورد', 'الفرق %'] if by_supplier else ['التغيّر %']) + ['إجمالي الكمية', 'إجمالي القيمة']
        rows = [head]
        styles = {}
        for c in r['categories']:
            vals = [(v if v else '') for v in c['values']]
            rows.append([f"{c['name']} ({c['count']} صنف)"] + vals + (['', ''] if by_supplier else [f"{c['change_pct']:+.1f}%" if c['change_pct'] is not None else '-']) + ['', c['total_value']])
            styles[len(rows) - 1] = 'section'
            for p in c['products']:
                rows.append([f"   {p['sku']} - {p['name']} ({p['unit']})"] + [(x if x is not None else '-') for x in p['prices']]
                            + ([r['columns'][p['best_index']] if p['best_index'] is not None else '-', f"{p['spread_pct']:.1f}%" if p['spread_pct'] is not None else '-'] if by_supplier
                               else [f"{p['change_pct']:+.1f}%" if p['change_pct'] is not None else '-']) + [p['total_qty'], p['total_value']])
        rows.append(['الإجمالي (قيمة المشتريات)'] + [(v if v else '') for v in r['totals']] + ([''] * (2 if by_supplier else 1)) + ['', r['grand_total']])
        styles[len(rows) - 1] = 'total'
        data['row_styles'] = styles
        n = len(head)
        data['col_weights'] = [3.0] + [1.2] * (n - 1)
        return rows
    if report_type == 'expiry_report':
        labels = {'expired': 'منتهي', 'near': 'قارب على الانتهاء', 'ok': 'سليم', 'none': 'بدون صلاحية'}
        rows = [['الكود', 'الصنف', 'المخزن', 'التشغيلة', 'تاريخ الصلاحية', 'المتبقي (يوم)', 'الكمية', 'تكلفة الوحدة', 'القيمة', 'الحالة']]
        styles = {}
        for r in data.get('expiry_rows', []):
            d = r['days_left']
            rows.append([r['product'].sku, r['product'].name, r['location'].name if r['location'] else '-', r['batch_no'] or '-', r['expiry_date'] or '-',
                         ('متأخر ' + str(-d)) if d is not None and d < 0 else (str(d) if d is not None else '-'), r['remaining'], r['unit_cost'] if _cost_ok() else None, r['value'] if _cost_ok() else None, labels.get(r['status'], r['status'])])
            if r['status'] == 'expired':
                styles[len(rows) - 1] = 'warn'
            elif r['status'] == 'near':
                styles[len(rows) - 1] = 'net'
        data['row_styles'] = styles
        data['col_weights'] = [1.1, 2.4, 1.5, 1.5, 1.4, 1.1, 1.0, 1.1, 1.2, 1.4]
        return rows
    if report_type == 'prepaid_schedule':
        pp = data['prepaid']
        labels = {'posted': 'مرحّل', 'pending': 'قادم', 'void': 'ملغى'}
        rows = [['م', 'من تاريخ', 'إلى تاريخ', 'تاريخ الاستحقاق', 'المبلغ', 'الحالة', 'تاريخ الترحيل']]
        styles = {}
        now = datetime.datetime.now()
        for l in pp['lines']:
            rows.append([str(l['seq']), l['period_from'], l['period_to'], l['entry_date'], l['amount'], labels.get(l['state'], l['state']), l['posted_at'] or '-'])
            if l['state'] == 'pending' and l['entry_date'] <= now:
                styles[len(rows) - 1] = 'net'
        rows.append(['', '', '', 'الإجمالي', sum(l['amount'] for l in pp['lines'] if l['state'] != 'void'), f"مرحّل {pp['posted']:,.2f}", f"متبقي {pp['remaining']:,.2f}"])
        styles[len(rows) - 1] = 'total'
        data['row_styles'] = styles
        data['col_weights'] = [0.6, 1.4, 1.4, 1.5, 1.3, 1.5, 1.8]
        return rows
    if report_type == 'hr_documents':
        labels = {'expired': 'منتهي', 'expiring': 'قارب على الانتهاء', 'valid': 'ساري', 'archived': 'مؤرشف'}
        rows = [['النوع', 'المستند', 'الخاص بـ', 'رقم المستند', 'جهة الإصدار', 'تاريخ الإصدار', 'تاريخ الانتهاء', 'المتبقي (يوم)', 'الحالة', 'التكلفة']]
        styles = {}
        for r in data.get('hr_docs', []):
            rows.append([r['type_name'], r['title'], r['holder'], r['doc_number'] or '-', r['issuer'] or '-', r['issue_date'], r['expiry_date'], (f"متأخر {-int(r['days_left'])}" if r['days_left'] < 0 else str(int(r['days_left']))), labels.get(r['status'], r['status']), r['cost'] or 0.0])
            if r['status'] == 'expired':
                styles[len(rows) - 1] = 'warn'
            elif r['status'] == 'expiring':
                styles[len(rows) - 1] = 'net'
        data['row_styles'] = styles
        data['col_weights'] = [1.6, 2.5, 1.7, 1.5, 1.4, 1.2, 1.2, 0.9, 1.4, 1.4]
        return rows
    if report_type in ('vat_return', 'vat_return_summary'):
        v = data['vat']; sm = v['summary']; sl = v['sales']; pu = v['purchases']; lg = v['ledger']
        rows = [['البند', 'المبلغ الخاضع / القيمة', 'مبلغ الضريبة']]; styles = {}
        def add(label, base='', tax='', style=None):
            rows.append([label, base, tax])
            if style: styles[len(rows) - 1] = style
        add('أولًا: ضريبة المخرجات (المبيعات)', '', '', 'section')
        for r in sl['rows']:
            add(r['label'], r['base'], r['tax'])
        add('إجمالي المبيعات وضريبة المخرجات', sl['total_base'], sl['total_tax'], 'total')
        add(f"   منها مبيعات الشركات B2B - {sl['b2b']['count']} فاتورة", sl['b2b']['base'], sl['b2b']['tax'])
        add(f"   منها مبيعات الأفراد B2C - {sl['b2c']['count']} فاتورة", sl['b2c']['base'], sl['b2c']['tax'])
        add('ثانيًا: ضريبة المدخلات (المشتريات)', '', '', 'section')
        for r in pu['rows']:
            add(f"{r['label']} - {r['count']} مستند", r['base'], r['tax'])
        add('إجمالي المشتريات وضريبة المدخلات', pu['total_base'], pu['total_tax'], 'total')
        add('ثالثًا: صافي الضريبة', '', '', 'section')
        add('ضريبة المخرجات', '', sm['output_tax']); add('ناقص: ضريبة المدخلات', '', sm['input_tax'])
        status = {'payable': 'صافي الضريبة المستحقة للسداد', 'refundable': 'رصيد دائن لصالح المنشأة (استرداد / ترحيل)', 'zero': 'صافي الضريبة (لا يوجد مستحق)'}[sm['status']]
        add(status, '', abs(sm['net_tax']), 'net')
        if report_type == 'vat_return':      # النسخة الكاملة بس: المطابقة مع الدفتر والتنبيهات (النسخة المختصرة = أولًا وثانيًا وثالثًا فقط)
            add('رابعًا: المطابقة مع دفتر الأستاذ (قيود مرحّلة)', '', '', 'section')
            add('ضريبة المخرجات - حسب الفواتير', '', lg['output_docs']); add('ضريبة المخرجات - حسب الدفتر', '', lg['output_gl']); add('الفرق', '', lg['output_diff'])
            add('ضريبة المدخلات - حسب المستندات', '', lg['input_docs']); add('ضريبة المدخلات - حسب الدفتر', '', lg['input_gl']); add('الفرق', '', lg['input_diff'])
            for w in v['warnings']:
                add('تنبيه: ' + w, '', '', 'warn')
        data['row_styles'] = styles
        data['col_weights'] = [6, 2.2, 2.2]
        return rows
    if report_type == 'vat_sales_details':
        rows = [['التاريخ', 'رقم الفاتورة', 'العميل', 'الرقم الضريبي', 'النوع', 'المبلغ الخاضع', 'الضريبة', 'الإجمالي']]
        for x in data['vat']['sales']['invoices']:
            rows.append([x['date'], x['number'], x['party'], x['vat_number'] or '-', x['kind'], x['base'], x['tax'], x['total']])
        sl = data['vat']['sales']; rows.append(['', '', '', '', 'الإجمالي', sl['total_base'], sl['total_tax'], round(sl['total_base'] + sl['total_tax'], 2)])
        data['row_styles'] = {len(rows) - 1: 'total'}
        return rows
    if report_type == 'vat_purchases_details':
        rows = [['التاريخ', 'المصدر', 'رقم الفاتورة', 'المورد', 'الرقم الضريبي', 'المبلغ', 'الضريبة', 'الإجمالي']]
        for x in data['vat']['purchases']['documents']:
            rows.append([x['date'], x['source'], x['number'], x['party'], x['vat_number'] or '-', x['base'], x['tax'], x['total']])
        pu = data['vat']['purchases']; rows.append(['', '', '', '', 'الإجمالي', pu['total_base'], pu['total_tax'], round(pu['total_base'] + pu['total_tax'], 2)])
        data['row_styles'] = {len(rows) - 1: 'total'}
        return rows
    if report_type == 'custody_statement':
        b = data.get('custody_balances') or {}
        rows = [['التاريخ','النوع','البيان','المرجع','الحساب','وارد','منصرف','الرصيد']]
        for r in data.get('custody_rows', []):
            rows.append([r['date'], r['type'], r['description'] or '', r['ref'], r['account'], r['incoming'], r['outgoing'], r['balance']])
        rows.append(['','','','','الإجمالي', b.get('funded', 0.0), round((b.get('spent', 0.0) or 0) + (b.get('returned', 0.0) or 0), 2), b.get('balance', 0.0)])
        return rows
    if report_type == 'ledger':
        return [['التاريخ','رقم القيد','الكود','الحساب','البيان','مدين','دائن','الرصيد']] + [[r['date'],r['number'],r['code'],r['account'],r['description'],r['debit'],r['credit'],r.get('balance','')] for r in data.get('ledger',[])]
    if report_type == 'stock':
        return [['التاريخ','الكود','الصنف','المخزن','الحركة','الكمية','تكلفة الوحدة','القيمة','المرجع']] + [[r['date'],r['sku'],r['product'],r['location'],r['direction'],r['quantity'],r['unit_cost'],r['value'],r['reference']] for r in data.get('moves',[])]
    if report_type == 'trial':
        header = ['الكود','الحساب','أول المدة مدين','أول المدة دائن','حركة الفترة مدين','حركة الفترة دائن','الرصيد الختامي مدين','الرصيد الختامي دائن']
        rows = [header]
        for r in data.get('trial_balance', []):
            rows.append([
                r.get('code',''), r.get('name',''),
                r.get('opening_debit',0), r.get('opening_credit',0),
                r.get('period_debit',0), r.get('period_credit',0),
                r.get('closing_debit',0), r.get('closing_credit',0),
            ])
        return rows
    if report_type == 'sales':
        return [['التاريخ','الفاتورة','العميل','الإجمالي']] + [[r['date'],r['number'],r['party'],r['total']] for r in data.get('sales',[])]
    if report_type == 'purchases':
        return [['التاريخ','الفاتورة','المورد','الإجمالي']] + [[r['date'],r['number'],r['party'],r['total']] for r in data.get('purchases',[])]
    if report_type == 'income':
        rows=[['قائمة الدخل التفصيلية','','','','']]
        rows += [['الإيرادات','الكود','الحساب','المبلغ','']]
        rows += [['إيراد','%s'%r['code'],r['name'],r['amount'],''] for r in data.get('income',[])]
        rows += [['إجمالي الإيرادات','','',data.get('total_income',0),'']]
        rows += [['المصروفات','الكود','الحساب','المبلغ','']]
        rows += [['مصروف','%s'%r['code'],r['name'],r['amount'],''] for r in data.get('expense',[])]
        rows += [['إجمالي المصروفات','','',data.get('total_expense',0),'']]
        rows += [['صافي الربح / (الخسارة)','','',data.get('net_profit',0),'']]
        return rows
    if report_type == 'balance_sheet':
        rows=[['الميزانية العمومية التفصيلية','','']]
        for key,title,total in [('assets','الأصول',data.get('total_assets',0)),('liabilities','الالتزامات',data.get('total_liabilities',0)),('equity','حقوق الملكية',data.get('total_equity',0))]:
            rows += [[title,'الحساب','المبلغ']]
            rows += [[r['code'],r['name'],r['amount']] for r in data.get(key,[])]
            rows += [[f'إجمالي {title}','','%s'%total]]
        rows += [['فرق التوازن','','%s'%data.get('balance_gap',0)]]
        return rows
    if report_type == 'cash_flow':
        rows=[['قائمة التدفقات النقدية','','']]
        rows += [['الرصيد الافتتاحي','',data.get('opening_cash',0)]]
        for key,title,net_key in [('operating','الأنشطة التشغيلية','net_operating'),('investing','الأنشطة الاستثمارية','net_investing'),('financing','الأنشطة التمويلية','net_financing')]:
            rows += [[title,'البيان','المبلغ']]
            rows += [['',r['description'],r['amount']] for r in data.get(key,[])]
            rows += [[f'صافي {title}','',data.get(net_key,'%s'%sum(r['amount'] for r in data.get(key,[])))]]
        rows += [['صافي التغير في النقدية','',data.get('net',0)],['الرصيد الختامي','',data.get('closing_cash',0)]]
        if abs(data.get('unexplained',0)) > 0.01:
            rows += [['فرق غير مفسَّر','',data.get('unexplained',0)]]
        return rows
    if report_type == 'party_statement':
        return [['التاريخ','المستند','البيان','مدين','دائن','الرصيد']] + [[r['date'],r['number'],r['description'],r['debit'],r['credit'],r['balance']] for r in data.get('rows',[])]
    if report_type == 'income_compare':
        return [['الكود','الحساب','الفترة 1','الفترة 2','الفرق','% التغير']] + [[r['code'],r['name'],r['period1'],r['period2'],r['difference'],r['change_pct']] for r in data.get('rows',[])]
    if report_type == 'balance_compare':
        return [['الكود','الحساب','النوع','تاريخ 1','تاريخ 2','الفرق']] + [[r['code'],r['name'],r['type'],r['period1'],r['period2'],r['difference']] for r in data.get('rows',[])]
    if report_type == 'equity':
        return [['الكود','الحساب','التغير']] + [[r['code'],r['name'],r['amount']] for r in data.get('movements',[])] + [['الربح / الخسارة','',data.get('net_profit',0)],['حقوق الملكية الختامية','',data.get('closing_equity',0)]]
    if report_type in ('branch_income','consolidated_income'):
        rows=[['الإيرادات','الكود','الحساب','المبلغ']]
        rows += [['إيراد',r['code'],r['name'],r['amount']] for r in data.get('income',[])]
        rows += [['إجمالي الإيرادات','','',data.get('total_income',0)],['المصروفات','','','']]
        rows += [['مصروف',r['code'],r['name'],r['amount']] for r in data.get('expense',[])]
        rows += [['إجمالي المصروفات','','',data.get('total_expense',0)],['صافي الربح / (الخسارة)','','',data.get('net_profit',0)]]
        return rows
    if report_type in ('branch_balance','consolidated_balance'):
        rows=[['النوع','الكود','الحساب','المبلغ']]
        for key,title,total in [('assets','الأصول',data.get('total_assets',0)),('liabilities','الالتزامات',data.get('total_liabilities',0)),('equity','حقوق الملكية',data.get('total_equity',0))]:
            rows += [[title,'','','']]+[[title,r['code'],r['name'],r['amount']] for r in data.get(key,[])] + [[f'إجمالي {title}','','',total]]
        rows += [['فرق التوازن','','',data.get('balance_gap',0)]]
        return rows
    if report_type == 'budget_vs_actual':
        return [['الكود','الحساب','النوع','الموازنة','الفعلي','الانحراف','% التحقيق']] + [[r['code'],r['account'],r['type'],r['budget'],r['actual'],r['variance'],r['achievement_pct']] for r in data.get('rows',[])]
    if report_type == 'cost_center':
        return [['الكود','مركز التكلفة','الإيرادات','المصروفات','صافي الربح']] + [[r['code'],r['name'],r['income'],r['expense'],r['profit']] for r in data.get('rows',data.get('cost_centers',[]))]
    if report_type == 'general_journal':
        return [['رقم القيد','التاريخ','المصدر','الحساب','البيان','الفرع','مركز التكلفة','مدين','دائن']] + [[r['number'],r['date'],r['source'],r['account'],r['description'],r['branch'],r['cost_center'],r['debit'],r['credit']] for r in data.get('rows',[])]
    if report_type == 'receipts_payments':
        return [['الرقم','التاريخ','النوع','الطرف','المبلغ','الحساب']] + [[r['number'],r['date'],r['type'],r['party'],r['amount'],r['account']] for r in data.get('rows',[])]
    if report_type == 'payroll':
        return [['كود الموظف','الاسم','الأساسي','بدل سكن','بدل نقل','بدلات أخرى','التأمينات','خصومات أخرى','الصافي']] + [
            [r['code'],r['name'],r['basic'],r['housing'],r['transport'],r['other_allow'],r['gosi'],r['other_ded'],r['net']]
            for r in data.get('rows',[])
        ]
    if report_type == 'allowances_insurance':
        return [['كود الموظف','الاسم','مخصص إضافي','ملاحظة المخصص','تأمينات صاحب العمل','مخصص الإجازة','مخصص نهاية الخدمة','الإجمالي']] + [
            [r['code'],r['name'],r['extra_allowance'],r['extra_note'],r['employer_gosi'],r.get('leave_provision',0),r.get('eos_provision',0),r['total']]
            for r in data.get('rows',[])
        ]
    if report_type == 'depreciation':
        return [['كود الأصل','اسم الأصل','قيمة القسط']] + [[r['code'],r['name'],r['amount']] for r in data.get('rows',[])]
    return []


def _report_title(report_type):
    return {
        'trial':'ميزان المراجعة التفصيلي','ledger':'دفتر الأستاذ العام','stock':'تقرير حركة المخزون',
        'sales':'تقرير المبيعات','purchases':'تقرير المشتريات','income':'قائمة الدخل التفصيلية',
        'balance_sheet':'الميزانية العمومية التفصيلية','cash_flow':'قائمة التدفقات النقدية','party_statement':'كشف الحساب المتحرك','income_compare':'قائمة الدخل المقارنة','balance_compare':'الميزانية المقارنة','equity':'قائمة التغيرات في حقوق الملكية','budget_vs_actual':'مقارنة Actual vs Budget','branch_income':'قائمة دخل الفرع','branch_balance':'ميزانية الفرع','consolidated_income':'قائمة الدخل الموحدة','consolidated_balance':'الميزانية الموحدة','cost_center':'تقرير مراكز التكلفة','general_journal':'اليومية العامة','receipts_payments':'سندات القبض والصرف','payroll':'كشف رواتب','allowances_insurance':'كشف مخصصات وتأمينات','depreciation':'كشف إهلاك أصول ثابتة','custody_statement':'كشف عهدة موظف','vat_return':'تقرير ضريبة القيمة المضافة','vat_return_summary':'إقرار ضريبة القيمة المضافة','hr_documents':'كشف الرخص والمستندات وتواريخ انتهائها','prepaid_schedule':'جدول تقسيم مصروف مدفوع مقدمًا','expiry_report':'تقرير الصلاحية والتشغيلات','stock_balance':'تقرير أرصدة المخزون','inventory_adjustment':'سند تسوية جردية','pos_session':'تقرير جلسة نقاط البيع','purchase_price_comparison':'تقرير مقارنة أسعار الشراء','sales_analysis':'تقرير المبيعات التحليلي','purchase_analysis':'تقرير المشتريات التحليلي','kitchen_events':'تقرير الإلغاءات والمرتجعات (المطعم)','void_report':'تقرير الملغاة والمرتجعات','audit_log':'سجل التدقيق','recipe_costing':'تقرير تكلفة الوصفات','restaurant_dashboard':'لوحة المطعم','discount_report':'تقرير الخصومات والكوبونات','data_integrity':'تقرير فحص سلامة البيانات','receivables_aging':'تقرير أعمار الديون','vat_sales_details':'تفاصيل فواتير المبيعات الضريبية','vat_purchases_details':'تفاصيل مستندات المشتريات والمدخلات'
    }.get(report_type,'تقرير مالي') + (f" - {data_party_name}" if False else '')


def _period_text(data):
    df=data.get('date_from'); dt=data.get('date_to')
    if df or dt: return f"الفترة: {_fmt_date(df)} إلى {_fmt_date(dt)}"
    if data.get('as_of'): return f"حتى تاريخ: {_fmt_date(data['as_of'])}"
    return ""


def export_excel(data, report_type, output_dir):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.drawing.image import Image as XLImage
    os.makedirs(output_dir,exist_ok=True)
    stamp=datetime.datetime.now().strftime('%Y%m%d_%H%M%S')
    path=os.path.join(output_dir,f'report_{report_type}_{stamp}.xlsx')
    company=_company(); wb=Workbook(); ws=wb.active; ws.title=_report_title(report_type)[:31]
    rows=_rows(data,report_type)
    max_cols=max((len(r) for r in rows),default=3)
    ws.merge_cells(start_row=1,start_column=1,end_row=1,end_column=max_cols)
    ws.cell(1,1,company.name if company else 'اسم الشركة')
    ws.cell(1,1).font=Font(size=18,bold=True); ws.cell(1,1).alignment=Alignment(horizontal='center')
    ws.merge_cells(start_row=2,start_column=1,end_row=2,end_column=max_cols)
    party_suffix = f" - {data.get('party_name')}" if report_type == 'party_statement' and data.get('party_name') else ''
    if report_type == 'custody_statement' and data.get('custody'):
        party_suffix = f" - {data['custody']['employee_name']} ({data['custody']['number']})"
    if data.get('title_suffix'):
        party_suffix += f" - {data['title_suffix']}"
    ws.cell(2,1,_report_title(report_type)+party_suffix); ws.cell(2,1).font=Font(size=14,bold=True); ws.cell(2,1).alignment=Alignment(horizontal='center')
    ws.merge_cells(start_row=3,start_column=1,end_row=3,end_column=max_cols)
    ws.cell(3,1,(_period_text(data)+'   |   تاريخ الطباعة: '+datetime.datetime.now().strftime('%Y-%m-%d %H:%M'))); ws.cell(3,1).alignment=Alignment(horizontal='center')
    row0=5
    for ri,row in enumerate(rows,row0):
        for ci,val in enumerate(row,1):
            ws.cell(ri,ci,val)
    header_row=row0
    for cell in ws[header_row]:
        cell.font=Font(bold=True,color='FFFFFF'); cell.fill=PatternFill('solid',fgColor='1F2937'); cell.alignment=Alignment(horizontal='center')
    thin=Side(style='thin',color='B7B7B7')
    for row in ws.iter_rows(min_row=header_row,max_row=ws.max_row,min_col=1,max_col=max_cols):
        for cell in row:
            cell.border=Border(left=thin,right=thin,top=thin,bottom=thin); cell.alignment=Alignment(horizontal='right',vertical='center')
    main_flags=_main_row_flags(rows,report_type)
    for i,is_main in enumerate(main_flags):
        if is_main:
            for ci in range(1,max_cols+1):
                c=ws.cell(row0+i,ci); c.font=Font(bold=True,color='3B2A99'); c.fill=PatternFill('solid',fgColor='EEF2FF')
    for i,k in (data.get('row_styles') or {}).items():
        for ci in range(1,max_cols+1):
            c=ws.cell(row0+i,ci); c.font=Font(bold=True); c.fill=PatternFill('solid',fgColor=_ROW_BG[k])
    if any(main_flags):
        ws.cell(4,1,'الحسابات الرئيسية مميّزة بخط عريض وخلفية فاتحة'); ws.cell(4,1).font=Font(size=9,italic=True,color='6B7280')
    ws.freeze_panes=f'A{header_row+1}'
    for col in range(1,max_cols+1):
        vals=[str(ws.cell(r,col).value or '') for r in range(1,ws.max_row+1)]
        ws.column_dimensions[get_column_letter(col)].width=min(42,max(12,max(map(len,vals))+2))
    if company and getattr(company,'logo_path',None) and os.path.exists(company.logo_path):
        try:
            img=XLImage(company.logo_path); img.width=min(img.width,110); img.height=min(img.height,70); ws.add_image(img,'A1')
        except Exception: pass
    ws.sheet_view.rightToLeft=True
    wb.save(path); return path


def _shape_ar(text):
    text=str(text)
    try:
        import arabic_reshaper
        from bidi.algorithm import get_display
        import re
        # الأرقام والتواريخ (زي 2026-10-04 أو 07:24) لازم تفضل متجهة من
        # الشمال لليمين حتى جوه جملة عربية، وإلا مجموعاتها بتتعكس (04-10-2026)
        text=re.sub(r'\d[\d\-/:.,]*\d', lambda m: '\u202a'+m.group(0)+'\u202c', text)
        return get_display(arabic_reshaper.reshape(text))
    except Exception:
        return text


def export_pdf(data, report_type, output_dir):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, Image as RLImage
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.pdfbase import pdfmetrics
    from reportlab.lib.enums import TA_RIGHT, TA_CENTER, TA_LEFT
    os.makedirs(output_dir,exist_ok=True)
    stamp=datetime.datetime.now().strftime('%Y%m%d_%H%M%S')
    path=os.path.join(output_dir,f'report_{report_type}_{stamp}.pdf')
    company=_company(); font_candidates=[
        '/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf',
        'C:/Windows/Fonts/arial.ttf','C:/Windows/Fonts/tahoma.ttf','C:/Windows/Fonts/segoeui.ttf']
    font=next((x for x in font_candidates if os.path.exists(x)),None)
    font_name='ERPFont'
    if font:
        try: pdfmetrics.registerFont(TTFont(font_name,font))
        except Exception: font_name='Helvetica'
    bold_candidates=['/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf','C:/Windows/Fonts/arialbd.ttf','C:/Windows/Fonts/tahomabd.ttf','C:/Windows/Fonts/segoeuib.ttf']
    bold_path=next((x for x in bold_candidates if os.path.exists(x)),None)
    bold_name=font_name
    if bold_path:
        try: pdfmetrics.registerFont(TTFont('ERPFontBold',bold_path)); bold_name='ERPFontBold'
        except Exception: bold_name=font_name
    BRAND=colors.HexColor('#1F2937'); ACCENT=colors.HexColor('#1A7A5E')
    normal=ParagraphStyle('ERPNormal',fontName=font_name,fontSize=8,leading=11,alignment=TA_RIGHT)
    title_style=ParagraphStyle('ERPTitle',fontName=font_name,fontSize=17,leading=21,alignment=TA_CENTER)
    sub_style=ParagraphStyle('ERPSub',fontName=font_name,fontSize=10,leading=13,alignment=TA_CENTER)
    doc=SimpleDocTemplate(path,pagesize=landscape(A4),rightMargin=28,leftMargin=28,topMargin=24,bottomMargin=28)
    story=[]
    company_name=(company.name if company and company.name else 'اسم الشركة')
    name_style=ParagraphStyle('ERPCo',fontName=bold_name,fontSize=15,leading=19,alignment=TA_RIGHT,textColor=BRAND)
    meta_style=ParagraphStyle('ERPMeta',fontName=font_name,fontSize=8,leading=11,alignment=TA_RIGHT,textColor=colors.HexColor('#6B7280'))
    small_left=ParagraphStyle('ERPSmallL',fontName=font_name,fontSize=8,leading=11,alignment=TA_LEFT,textColor=colors.HexColor('#6B7280'))
    banner_style=ParagraphStyle('ERPBanner',fontName=bold_name,fontSize=13,leading=17,alignment=TA_CENTER,textColor=colors.white)
    # اليمين: الشعار + اسم الشركة + بيانات الاتصال | الشمال: تاريخ الطباعة والفترة
    right_cells=[]
    logo=getattr(company,'logo_path',None) if company else None
    if logo and os.path.exists(logo):
        try:
            img=RLImage(logo); img._restrictSize(80,48); right_cells.append(img)
        except Exception: pass
    right_cells.append(Paragraph(_shape_ar(company_name),name_style))
    if company:
        for label,val in (('الرقم الضريبي',company.vat_number),('هاتف',company.phone),('العنوان',company.address)):
            if val: right_cells.append(Paragraph(_shape_ar(f'{label}: {val}'),meta_style))
    left_cells=[Paragraph(_shape_ar('تاريخ الطباعة: '+datetime.datetime.now().strftime('%Y-%m-%d %H:%M')),small_left)]
    period=_period_text(data)
    if period: left_cells.append(Paragraph(_shape_ar(period),small_left))
    page_w=landscape(A4)[0]-56
    head=Table([[left_cells,right_cells]],colWidths=[page_w*0.4,page_w*0.6])
    head.setStyle(TableStyle([('VALIGN',(0,0),(-1,-1),'TOP'),('LEFTPADDING',(0,0),(-1,-1),0),('RIGHTPADDING',(0,0),(-1,-1),0),('BOTTOMPADDING',(0,0),(-1,-1),6),('LINEBELOW',(0,0),(-1,0),1.2,ACCENT)]))
    title_text=_report_title(report_type)+(f" - {data.get('party_name')}" if report_type=='party_statement' and data.get('party_name') else '')
    if report_type=='custody_statement' and data.get('custody'):
        title_text += f" - {data['custody']['employee_name']} ({data['custody']['number']})"
    if data.get('title_suffix'):
        title_text += f" - {data['title_suffix']}"
    banner=Table([[Paragraph(_shape_ar(title_text),banner_style)]],colWidths=[page_w])
    banner.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,-1),BRAND),('TOPPADDING',(0,0),(-1,-1),6),('BOTTOMPADDING',(0,0),(-1,-1),6)]))
    story += [head,Spacer(1,8),banner,Spacer(1,10)]
    raw=_rows(data,report_type)
    # الجداول في reportlab بترتّب الأعمدة من الشمال لليمين حسب ترتيب
    # القايمة، بغض النظر عن محاذاة النص جوه كل خلية - فـ_rows() بتبني كل
    # صف بالترتيب المنطقي العادي (التاريخ، رقم القيد، ...، الرصيد) واللي
    # كان بيطلع معكوس بصريًا (التاريخ على أقصى الشمال بدل اليمين). بنعكس
    # كل صف هنا مرة واحدة بدل ما نعدّل كل نوع تقرير لوحده.
    main_flags=_main_row_flags(raw,report_type)
    raw=[list(reversed(row)) for row in raw]
    head_cell=ParagraphStyle('ERPHeadCell',parent=normal,fontName=bold_name,textColor=colors.white)
    main_cell=ParagraphStyle('ERPMainCell',parent=normal,fontName=bold_name,textColor=colors.HexColor('#0F5A45'))
    row_styles=data.get('row_styles') or {}
    bold_cell=ParagraphStyle('ERPBoldCell',parent=normal,fontName=bold_name)
    table_rows=[]
    for i,row in enumerate(raw):
        st = head_cell if i==0 else (bold_cell if i in row_styles else (main_cell if main_flags[i] else normal))
        table_rows.append([Paragraph(_shape_ar(_fmt(x)),st) for x in row])
    if any(main_flags):
        legend=ParagraphStyle('ERPLegend',parent=normal,fontSize=7.5,textColor=colors.HexColor('#0F5A45'))
        story.append(Paragraph(_shape_ar('الحسابات الرئيسية مميّزة بخط عريض وخلفية فاتحة'),legend)); story.append(Spacer(1,4))
    if table_rows:
        cw=None
        if data.get('col_weights') and len(data['col_weights'])==len(raw[0]):
            tw=float(sum(data['col_weights'])); cw=[landscape(A4)[0]*0.0+(landscape(A4)[0]-56)*w/tw for w in reversed(data['col_weights'])]
        t=Table(table_rows,repeatRows=1,hAlign='RIGHT',colWidths=cw)
        t.setStyle(TableStyle([
            ('BACKGROUND',(0,0),(-1,0),colors.HexColor('#1F2937')),('TEXTCOLOR',(0,0),(-1,0),colors.white),
            ('GRID',(0,0),(-1,-1),0.35,colors.HexColor('#B7B7B7')),('FONTNAME',(0,0),(-1,-1),font_name),
            ('FONTSIZE',(0,0),(-1,-1),7),('VALIGN',(0,0),(-1,-1),'MIDDLE'),('ALIGN',(0,0),(-1,-1),'RIGHT'),
            ('ROWBACKGROUNDS',(0,1),(-1,-1),[colors.white,colors.HexColor('#F8FAFC')]),('TOPPADDING',(0,0),(-1,-1),4),('BOTTOMPADDING',(0,0),(-1,-1),4)
        ]))
        main_cmds=[('BACKGROUND',(0,i),(-1,i),colors.HexColor('#EAF5F0')) for i,f in enumerate(main_flags) if f and i>0]
        main_cmds+=[('BACKGROUND',(0,i),(-1,i),colors.HexColor('#'+_ROW_BG[k])) for i,k in row_styles.items()]
        if main_cmds: t.setStyle(TableStyle(main_cmds))
        story.append(t)
    def footer(canvas,doc):
        w=doc.pagesize[0]
        canvas.saveState()
        canvas.setStrokeColor(colors.HexColor('#D1D5DB')); canvas.setLineWidth(0.5); canvas.line(28,24,w-28,24)
        canvas.setFont(font_name,7); canvas.setFillColor(colors.HexColor('#6B7280'))
        canvas.drawRightString(w-28,13,_shape_ar(company_name+' - تقرير رسمي'))
        canvas.drawCentredString(w/2,13,_shape_ar('صفحة %d'%doc.page))
        canvas.drawString(28,13,datetime.datetime.now().strftime('%Y-%m-%d %H:%M'))
        canvas.restoreState()
    doc.build(story,onFirstPage=footer,onLaterPages=footer)
    return path


def export_voucher_pdf(doc, output_dir, prefix='voucher'):
    """سند مطبوع (A4 طولي) برأس الشركة: عنوان السند + رقمه وتاريخه، جدول بيانات
    (البند/القيمة)، جدول تفاصيل اختياري، إجماليات، وخانات توقيع.
    doc = {title, number, date, fields:[(label,value)], table:{headers,rows}|None, totals:[(label,value)], signatures:[str]}"""
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, Image as RLImage
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.pdfbase import pdfmetrics
    from reportlab.lib.enums import TA_RIGHT, TA_CENTER, TA_LEFT
    os.makedirs(output_dir, exist_ok=True)
    stamp = datetime.datetime.now().strftime('%Y%m%d_%H%M%S')
    path = os.path.join(output_dir, f'{prefix}_{doc.get("number","")}_{stamp}.pdf')
    company = _company()
    font_name = 'Helvetica'
    for cand in ['/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf', 'C:/Windows/Fonts/arial.ttf', 'C:/Windows/Fonts/tahoma.ttf', 'C:/Windows/Fonts/segoeui.ttf']:
        if os.path.exists(cand):
            try: pdfmetrics.registerFont(TTFont('ERPFont', cand)); font_name = 'ERPFont'
            except Exception: pass
            break
    bold_name = font_name
    for cand in ['/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf', 'C:/Windows/Fonts/arialbd.ttf', 'C:/Windows/Fonts/tahomabd.ttf', 'C:/Windows/Fonts/segoeuib.ttf']:
        if os.path.exists(cand):
            try: pdfmetrics.registerFont(TTFont('ERPFontBold', cand)); bold_name = 'ERPFontBold'
            except Exception: pass
            break
    BRAND = colors.HexColor('#1F2937'); ACCENT = colors.HexColor('#1A7A5E'); GREY = colors.HexColor('#6B7280')
    st = lambda name, **kw: ParagraphStyle(name, **kw)
    normal = st('VN', fontName=font_name, fontSize=9, leading=13, alignment=TA_RIGHT)
    boldp = st('VB', fontName=bold_name, fontSize=9, leading=13, alignment=TA_RIGHT)
    head = st('VH', fontName=bold_name, fontSize=9, leading=13, alignment=TA_RIGHT, textColor=colors.white)
    company_name = company.name if company and company.name else 'اسم الشركة'
    page_w = A4[0] - 56
    story = []
    right = [Paragraph(_shape_ar(company_name), st('VC', fontName=bold_name, fontSize=15, leading=19, alignment=TA_RIGHT, textColor=BRAND))]
    if company:
        for label, val in (('الرقم الضريبي', company.vat_number), ('هاتف', company.phone), ('العنوان', company.address)):
            if val: right.append(Paragraph(_shape_ar(f'{label}: {val}'), st('VM', fontName=font_name, fontSize=8, leading=11, alignment=TA_RIGHT, textColor=GREY)))
    d = doc.get('date')
    left = [Paragraph(_shape_ar('رقم السند: ' + str(doc.get('number', ''))), st('VL', fontName=bold_name, fontSize=10, leading=14, alignment=TA_LEFT)),
            Paragraph(_shape_ar('التاريخ: ' + (d.strftime('%Y-%m-%d') if hasattr(d, 'strftime') else str(d or ''))), st('VL2', fontName=font_name, fontSize=9, leading=13, alignment=TA_LEFT))]
    hdr = Table([[left, right]], colWidths=[page_w * 0.4, page_w * 0.6])
    hdr.setStyle(TableStyle([('VALIGN', (0, 0), (-1, -1), 'TOP'), ('LEFTPADDING', (0, 0), (-1, -1), 0), ('RIGHTPADDING', (0, 0), (-1, -1), 0), ('BOTTOMPADDING', (0, 0), (-1, -1), 6), ('LINEBELOW', (0, 0), (-1, 0), 1.2, ACCENT)]))
    banner = Table([[Paragraph(_shape_ar(doc.get('title', '')), st('VT', fontName=bold_name, fontSize=15, leading=19, alignment=TA_CENTER, textColor=colors.white))]], colWidths=[page_w])
    banner.setStyle(TableStyle([('BACKGROUND', (0, 0), (-1, -1), BRAND), ('TOPPADDING', (0, 0), (-1, -1), 7), ('BOTTOMPADDING', (0, 0), (-1, -1), 7)]))
    story += [hdr, Spacer(1, 8), banner, Spacer(1, 12)]

    fields = doc.get('fields', [])
    if fields:
        rows = [[Paragraph(_shape_ar(str(v if v not in (None, '') else '-')), normal), Paragraph(_shape_ar(l), boldp)] for l, v in fields]
        ft_ = Table(rows, colWidths=[page_w * 0.68, page_w * 0.32])
        ft_.setStyle(TableStyle([('GRID', (0, 0), (-1, -1), 0.4, colors.HexColor('#C9CDD3')), ('BACKGROUND', (1, 0), (1, -1), colors.HexColor('#F3F4F6')),
                                 ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'), ('TOPPADDING', (0, 0), (-1, -1), 5), ('BOTTOMPADDING', (0, 0), (-1, -1), 5)]))
        story += [ft_, Spacer(1, 12)]

    table = doc.get('table')
    if table and table.get('rows'):
        headers = list(reversed(table['headers']))
        body = [[Paragraph(_shape_ar(h), head) for h in headers]]
        for r in table['rows']:
            body.append([Paragraph(_shape_ar(_fmt(x)), normal) for x in reversed(r)])
        n = len(headers)
        weights = table.get('col_weights')
        if weights and len(weights) == n:
            tot_w = float(sum(weights)); widths = [page_w * w / tot_w for w in reversed(weights)]
        else:
            widths = [page_w / n] * n
        tt = Table(body, repeatRows=1, colWidths=widths)
        tt.setStyle(TableStyle([('LEFTPADDING', (0, 0), (-1, -1), 3), ('RIGHTPADDING', (0, 0), (-1, -1), 3), ('BACKGROUND', (0, 0), (-1, 0), BRAND), ('GRID', (0, 0), (-1, -1), 0.4, colors.HexColor('#C9CDD3')),
                                ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F9FAFB')]), ('VALIGN', (0, 0), (-1, -1), 'MIDDLE')]))
        story += [Paragraph(_shape_ar(table.get('caption', 'التفاصيل')), boldp), Spacer(1, 4), tt, Spacer(1, 12)]

    totals = doc.get('totals', [])
    if totals:
        trows = [[Paragraph(_shape_ar(_fmt(v)), boldp), Paragraph(_shape_ar(l), boldp)] for l, v in totals]
        tot = Table(trows, colWidths=[page_w * 0.3, page_w * 0.3], hAlign='RIGHT')
        tot.setStyle(TableStyle([('GRID', (0, 0), (-1, -1), 0.4, colors.HexColor('#C9CDD3')), ('BACKGROUND', (1, 0), (1, -1), colors.HexColor('#EAF5F0')),
                                 ('TOPPADDING', (0, 0), (-1, -1), 5), ('BOTTOMPADDING', (0, 0), (-1, -1), 5)]))
        story += [tot, Spacer(1, 26)]

    sigs = doc.get('signatures', [])
    if sigs:
        cells = [[Paragraph(_shape_ar(sg), st('VS', fontName=bold_name, fontSize=9, leading=12, alignment=TA_CENTER)) for sg in sigs],
                 [Paragraph('', normal) for _ in sigs]]
        sg = Table(cells, colWidths=[page_w / len(sigs)] * len(sigs), rowHeights=[18, 38])
        sg.setStyle(TableStyle([('BOX', (0, 0), (-1, -1), 0.5, colors.HexColor('#9CA3AF')), ('INNERGRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#9CA3AF')),
                                ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#F3F4F6')), ('VALIGN', (0, 0), (-1, -1), 'MIDDLE')]))
        story.append(sg)

    def footer(canvas, docu):
        w = docu.pagesize[0]
        canvas.saveState(); canvas.setStrokeColor(colors.HexColor('#D1D5DB')); canvas.setLineWidth(0.5); canvas.line(28, 24, w - 28, 24)
        canvas.setFont(font_name, 7); canvas.setFillColor(GREY)
        canvas.drawRightString(w - 28, 13, _shape_ar(company_name + ' - ' + doc.get('title', '')))
        canvas.drawString(28, 13, datetime.datetime.now().strftime('%Y-%m-%d %H:%M'))
        canvas.restoreState()
    SimpleDocTemplate(path, pagesize=A4, rightMargin=28, leftMargin=28, topMargin=24, bottomMargin=34).build(story, onFirstPage=footer, onLaterPages=footer)
    return path
