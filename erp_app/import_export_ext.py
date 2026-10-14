"""توسيع الاستيراد والتصدير ليغطي كل جداول النظام:
  1) كيانات للاستيراد والتصدير بقوالب Excel عربي (فئات، فروع، مراكز تكلفة، مخازن، طرق دفع، قوائم أسعار وأسعارها، موظفين،
     فئات أصول، موازنات، مكونات تصنيع، فترات محاسبية، مجموعات كاشير، أقسام وطاولات مطعم، أنواع ومستندات وقوالب متابعة...).
  2) تصدير أي جدول من قاعدة البيانات (كل الـ 76 جدول) أو كلهم في ملف واحد، مع ترجمة الأعمدة المرتبطة (العميل، الصنف، الحساب...)
     لأسماء مقروءة بجانب المعرّف."""
import datetime
import openpyxl
from sqlalchemy import select

from . import services
from .import_export import ENTITY_SPECS, _style_header, _row_values, _get, _yes_no, _tick
from .db import Base, SessionLocal

GROUPS = {"products": "الأصناف والمخزون", "units": "الأصناف والمخزون", "unit_factors": "الأصناف والمخزون", "adjustments": "الأصناف والمخزون",
          "customers": "المبيعات", "suppliers": "المشتريات", "accounts": "المحاسبة", "journal": "المحاسبة"}


# ----------------------------------------------------------------------
# مساعدات التحويل
# ----------------------------------------------------------------------
def _s(v):
    return str(v).strip() if v not in (None, "") else ""


def _f(v, default=0.0):
    if v in (None, ""):
        return default
    try:
        return float(str(v).replace(",", ""))
    except ValueError:
        raise ValueError(f'القيمة "{v}" مش رقم صحيح')


def _i(v, default=0):
    return int(_f(v, default))


def _date(v, required=False, label="التاريخ"):
    if v in (None, ""):
        if required:
            raise ValueError(f"{label} مطلوب")
        return None
    if isinstance(v, datetime.datetime):
        return v
    if isinstance(v, datetime.date):
        return datetime.datetime.combine(v, datetime.time())
    from .theme import parse_date_text
    d = parse_date_text(v)
    if not d:
        raise ValueError(f'{label} "{v}" مش بصيغة تاريخ صحيحة (مثلًا 2026-01-31 أو 31/01/2026)')
    return datetime.datetime.combine(d, datetime.time())


def _yn(b):
    return "نعم" if b else "لا"


def _fmt_date(d):
    return d.strftime("%Y-%m-%d") if d else ""


class _Ctx:
    """مخازن بحث مؤقتة أثناء الاستيراد (كود/اسم -> سجل) عشان ما نسألش القاعدة لكل صف."""

    def __init__(self):
        self._c = {}

    def _get(self, key, loader):
        if key not in self._c:
            self._c[key] = loader()
        return self._c[key]

    def accounts(self):
        return self._get("acc", lambda: {a.code: a for a in services.list_accounts()})

    def accounts_by_id(self):
        return self._get("acc_id", lambda: {a.id: a for a in self.accounts().values()})

    def account_id(self, code, required=False, label="الحساب"):
        code = _s(code)
        if not code:
            if required:
                raise ValueError(f"{label} مطلوب")
            return None
        a = self.accounts().get(code)
        if not a:
            raise ValueError(f'{label}: الكود "{code}" مش موجود في شجرة الحسابات')
        return a.id

    def branches(self):
        return self._get("br", lambda: {b.code: b for b in services.list_branches(False)})

    def branch_id(self, code):
        code = _s(code)
        if not code:
            return None
        b = self.branches().get(code)
        if not b:
            raise ValueError(f'كود الفرع "{code}" مش موجود')
        return b.id

    def cost_centers(self):
        return self._get("cc", lambda: {c.code: c for c in services.list_cost_centers(False)})

    def cost_center_id(self, code):
        code = _s(code)
        if not code:
            return None
        c = self.cost_centers().get(code)
        if not c:
            raise ValueError(f'كود مركز التكلفة "{code}" مش موجود')
        return c.id

    def products(self):
        return self._get("prod", lambda: {p.sku: p for p in services.list_products(active_only=False)})

    def product_id(self, sku, label="الصنف"):
        sku = _s(sku)
        p = self.products().get(sku)
        if not p:
            raise ValueError(f'{label}: كود الصنف "{sku}" مش موجود')
        return p.id

    def employees(self):
        return self._get("emp", lambda: list(services.list_employees(active_only=False)))

    def reset(self, *keys):
        for k in keys:
            self._c.pop(k, None)


# ----------------------------------------------------------------------
# محرك الكيانات العام
# ----------------------------------------------------------------------
def _register(key, label, group, columns, export_rows, import_row=None, notes=(), counts_updates=True):
    headers = [c[0] for c in columns]
    fields = dict(columns)

    def template(path):
        wb = openpyxl.Workbook(); ws = wb.active; ws.title = label[:31]
        ws.append(headers); _style_header(ws, len(headers))
        ws2 = wb.create_sheet("تعليمات"); ws2.column_dimensions["A"].width = 100
        for line in ["الحقول اللي عليها * إجبارية.", "احذف أي صف مثال قبل الاستيراد.", *notes]:
            ws2.append([line])
        wb.save(path)
        return path

    def export(path):
        wb = openpyxl.Workbook(); ws = wb.active; ws.title = label[:31]
        ws.append(headers); _style_header(ws, len(headers))
        for r in export_rows():
            ws.append(list(r))
        wb.save(path)
        return path

    def importer(path):
        wb = openpyxl.load_workbook(path, data_only=True)
        ws = wb.worksheets[0]
        header_row = [c.value for c in ws[1]]
        col_map = _row_values(ws, header_row, fields)
        required = [f for h, f in columns if h.endswith("*")]
        if not col_map or any(f not in col_map for f in required):
            raise ValueError("الملف مش مطابق للقالب - حمّل القالب الفارغ لهذا النوع واملأه (عناوين الأعمدة لازم تكون زي القالب)")
        ctx = _Ctx()
        created = updated = 0
        errors = []
        for row_idx in range(2, ws.max_row + 1):
            _tick(row_idx - 1, ws.max_row - 1)
            row = [c.value for c in ws[row_idx]]
            if all(v in (None, "") for v in row):
                continue
            values = {f: _get(row, col_map, f) for f in fields.values()}
            try:
                res = import_row(values, ctx)
                if res == "created":
                    created += 1
                elif res == "updated":
                    updated += 1
            except Exception as ex:
                errors.append((row_idx, str(ex)))
        return {"created": created, "updated": updated, "errors": errors}

    ENTITY_SPECS[key] = {"label": label, "group": group, "template": template if import_row else None, "export": export, "import": importer if import_row else None}


# ======================================================================
# فئات الأصناف
# ======================================================================
def _imp_category(v, ctx):
    name = _s(v["name"])
    if not name:
        raise ValueError("اسم الفئة مطلوب")
    if name in {c.name for c in services.list_categories()}:
        return None
    services.create_category(name)
    return "created"


_register("categories", "فئات الأصناف", "الأصناف والمخزون", [("اسم الفئة*", "name")], lambda: [[c.name] for c in services.list_categories()], _imp_category,
          ["الفئة الموجودة بنفس الاسم بتتجاهل."])


# ======================================================================
# الفروع
# ======================================================================
def _imp_branch(v, ctx):
    code, name = _s(v["code"]), _s(v["name"])
    if not code or not name:
        raise ValueError("كود واسم الفرع مطلوبان")
    active = _yes_no(v["is_active"], True); default = _yes_no(v["is_default"], False)
    ex = ctx.branches().get(code)
    if ex:
        services.update_branch(ex.id, name=name, is_active=active, **({"is_default": True} if default else {}))
        ctx.reset("br")
        return "updated"
    bid = services.create_branch(code, name, default)
    if not active:
        services.update_branch(bid, is_active=False)
    ctx.reset("br")
    return "created"


_register("branches", "الفروع", "المنظمة", [("الكود*", "code"), ("الاسم*", "name"), ("افتراضي (نعم/لا)", "is_default"), ("نشط (نعم/لا)", "is_active")],
          lambda: [[b.code, b.name, _yn(b.is_default), _yn(b.is_active)] for b in services.list_branches(False)], _imp_branch, ["كود الفرع الموجود بيتحدّث اسمه وحالته."])


# ======================================================================
# مراكز التكلفة
# ======================================================================
def _imp_cost_center(v, ctx):
    code, name = _s(v["code"]), _s(v["name"])
    if not code or not name:
        raise ValueError("كود واسم مركز التكلفة مطلوبان")
    parent = ctx.cost_center_id(v["parent"]) if _s(v["parent"]) else None
    branch = ctx.branch_id(v["branch"])
    active = _yes_no(v["is_active"], True)
    ex = ctx.cost_centers().get(code)
    if ex:
        services.update_cost_center(ex.id, name=name, parent_id=parent, branch_id=branch, is_active=active)
        ctx.reset("cc")
        return "updated"
    cid = services.create_cost_center(code, name, parent, branch)
    if not active:
        services.update_cost_center(cid, is_active=False)
    ctx.reset("cc")
    return "created"


def _exp_cost_centers():
    by_id = {c.id: c for c in services.list_cost_centers(False)}
    br = {b.id: b.code for b in services.list_branches(False)}
    return [[c.code, c.name, by_id[c.parent_id].code if c.parent_id in by_id else "", br.get(c.branch_id, ""), _yn(c.is_active)] for c in by_id.values()]


_register("cost_centers", "مراكز التكلفة", "المنظمة", [("الكود*", "code"), ("الاسم*", "name"), ("كود المركز الأب", "parent"), ("كود الفرع", "branch"), ("نشط (نعم/لا)", "is_active")],
          _exp_cost_centers, _imp_cost_center, ["المركز الأب لازم يكون مكتوب في صف قبل صف المركز الفرعي (أو موجود بالفعل)."])


# ======================================================================
# المخازن
# ======================================================================
def _imp_location(v, ctx):
    name = _s(v["name"])
    if not name:
        raise ValueError("اسم المخزن مطلوب")
    branch = ctx.branch_id(v["branch"]); cc = ctx.cost_center_id(v["cc"])
    active = _yes_no(v["is_active"], True); default = _yes_no(v["is_default"], False)
    ex = next((l for l in services.list_locations(active_only=False) if l.name == name), None)
    if ex:
        services.update_location(ex.id, branch_id=branch, cost_center_id=cc, is_default=default, is_active=active)
        return "updated"
    lid = services.create_location(name, branch, cc, default)
    if not active:
        services.update_location(lid, is_active=False)
    return "created"


def _exp_locations():
    br = {b.id: b.code for b in services.list_branches(False)}; cc = {c.id: c.code for c in services.list_cost_centers(False)}
    return [[l.name, br.get(l.branch_id, ""), cc.get(l.cost_center_id, ""), _yn(l.is_default), _yn(l.is_active)] for l in services.list_locations(active_only=False)]


_register("locations", "المخازن", "الأصناف والمخزون", [("اسم المخزن*", "name"), ("كود الفرع", "branch"), ("كود مركز التكلفة", "cc"), ("افتراضي (نعم/لا)", "is_default"), ("نشط (نعم/لا)", "is_active")],
          _exp_locations, _imp_location, ["المخزن الموجود بنفس الاسم بيتحدّث."])


# ======================================================================
# طرق الدفع
# ======================================================================
def _imp_payment_method(v, ctx):
    name = _s(v["name"])
    if not name:
        raise ValueError("اسم طريقة الدفع مطلوب")
    acc = ctx.account_id(v["account"], False, "حساب طريقة الدفع")
    credit = _yes_no(v["is_credit"], False); active = _yes_no(v["is_active"], True); order = _i(v["order"], 0)
    ex = next((m for m in services.list_payment_methods(active_only=False) if m.name == name), None)
    if ex:
        services.update_payment_method(ex.id, name, acc, credit, active, order)
        return "updated"
    mid = services.create_payment_method(name, acc, credit, order)
    if not active:
        services.update_payment_method(mid, name, acc, credit, False, order)
    return "created"


def _exp_payment_methods():
    ab = {a.id: a.code for a in services.list_accounts()}
    return [[m.name, ab.get(m.account_id, ""), _yn(m.is_credit), _yn(m.is_active), m.sort_order] for m in services.list_payment_methods(active_only=False)]


_register("payment_methods", "طرق الدفع", "المحاسبة", [("الاسم*", "name"), ("كود الحساب", "account"), ("آجل (نعم/لا)", "is_credit"), ("نشط (نعم/لا)", "is_active"), ("الترتيب", "order")],
          _exp_payment_methods, _imp_payment_method, ["كود الحساب = الخزينة أو البنك اللي بيستقبل المبلغ. الطريقة الآجلة (على حساب العميل) من غير حساب."])


# ======================================================================
# قوائم الأسعار وأسعارها
# ======================================================================
def _imp_price_list(v, ctx):
    name = _s(v["name"])
    if not name:
        raise ValueError("اسم القائمة مطلوب")
    if name in {p.name for p in services.list_price_lists(False)}:
        return None
    services.create_price_list(name, _s(v["notes"]))
    return "created"


_register("price_lists", "قوائم الأسعار", "المبيعات", [("اسم القائمة*", "name"), ("ملاحظات", "notes")],
          lambda: [[p.name, p.notes or ""] for p in services.list_price_lists(False)], _imp_price_list, ["القائمة الموجودة بنفس الاسم بتتجاهل."])


def _imp_price_item(v, ctx):
    lname, sku = _s(v["list"]), _s(v["sku"])
    if not lname or not sku:
        raise ValueError("اسم القائمة وكود الصنف مطلوبين")
    price = _f(v["price"], None) if v["price"] not in (None, "") else None
    if price is None or price < 0:
        raise ValueError("السعر مطلوب ولازم يكون صفر أو أكتر")
    lst = ctx._get("pl", lambda: {p.name: p for p in services.list_price_lists(False)}).get(lname)
    if not lst:
        lid = services.create_price_list(lname)          # القائمة مش موجودة: بنعملها
        ctx.reset("pl"); lst = type("L", (), {"id": lid})
    services.set_price_list_item(lst.id, ctx.product_id(sku), price)
    return "updated"


def _exp_price_items():
    rows = []
    for pl in services.list_price_lists(False):
        for it in services.list_price_list_items(pl.id):
            rows.append([pl.name, it.product.sku if it.product else "", it.product.name if it.product else "", it.price])
    return rows


_register("price_list_items", "أسعار الأصناف في قوائم الأسعار", "المبيعات", [("اسم القائمة*", "list"), ("كود الصنف*", "sku"), ("اسم الصنف (للعرض)", "pname"), ("السعر*", "price")],
          _exp_price_items, _imp_price_item, ["القائمة الجديدة بتتعمل تلقائي. السعر الموجود لنفس الصنف بيتحدّث. اسم الصنف للعرض بس ومش بيتستورد."])


# ======================================================================
# الموظفون
# ======================================================================
def _imp_employee(v, ctx):
    name = _s(v["name"])
    if not name:
        raise ValueError("اسم الموظف مطلوب")
    code = _s(v["code"]); nid = _s(v["national_id"])
    hire = _date(v["hire_date"], False, "تاريخ التعيين")
    branch = ctx.branch_id(v["branch"]); cc = ctx.cost_center_id(v["cc"]); active = _yes_no(v["is_active"], True)
    fields = dict(name=name, national_id=nid, phone=_s(v["phone"]), email=_s(v["email"]), position=_s(v["position"]), department=_s(v["department"]), bank_account=_s(v["bank_account"]),
                  is_active=active)
    if hire:
        fields["hire_date"] = hire
    if branch:
        fields["branch_id"] = branch
    if cc:
        fields["cost_center_id"] = cc
    emps = ctx.employees()
    ex = next((e for e in emps if code and e.code == code), None) or next((e for e in emps if nid and e.national_id == nid), None)
    if ex:
        services.update_employee(ex.id, **fields)
        return "updated"
    eid = services.create_employee(name, nid, fields["phone"], fields["email"], fields["position"], fields["department"], hire, fields["bank_account"], branch, cc)
    if not active:
        services.update_employee(eid, is_active=False)
    ctx.reset("emp")
    return "created"


def _exp_employees():
    br = {b.id: b.code for b in services.list_branches(False)}; cc = {c.id: c.code for c in services.list_cost_centers(False)}
    return [[e.code, e.name, e.national_id or "", e.phone or "", e.email or "", e.position or "", e.department or "", _fmt_date(e.hire_date), e.bank_account or "",
             br.get(e.branch_id, ""), cc.get(e.cost_center_id, ""), _yn(e.is_active)] for e in services.list_employees(active_only=False)]


_register("employees", "الموظفون", "الموارد البشرية",
          [("كود الموظف", "code"), ("الاسم*", "name"), ("رقم الهوية / الإقامة", "national_id"), ("الهاتف", "phone"), ("البريد الإلكتروني", "email"), ("المسمى الوظيفي", "position"),
           ("القسم", "department"), ("تاريخ التعيين", "hire_date"), ("الحساب البنكي (IBAN)", "bank_account"), ("كود الفرع", "branch"), ("كود مركز التكلفة", "cc"), ("نشط (نعم/لا)", "is_active")],
          _exp_employees, _imp_employee, ["الموظف بيتعرّف بكود الموظف (EMP-...) أو رقم الهوية، لو موجود بيتحدّث وإلا بيتعمل جديد بكود تلقائي.", "تاريخ التعيين: 2026-01-31 أو 31/01/2026."])


# ======================================================================
# فئات الأصول الثابتة
# ======================================================================
def _imp_asset_category(v, ctx):
    name = _s(v["name"])
    if not name:
        raise ValueError("اسم الفئة مطلوب")
    a1 = ctx.account_id(v["asset"], True, "حساب الأصل"); a2 = ctx.account_id(v["accum"], True, "حساب مجمع الإهلاك"); a3 = ctx.account_id(v["expense"], True, "حساب مصروف الإهلاك")
    life = _i(v["life"], 60); salv = _f(v["salvage"], 0.0); active = _yes_no(v["is_active"], True)
    ex = next((c for c in services.list_asset_categories(False) if c.name == name), None)
    if ex:
        services.update_asset_category(ex.id, name, a1, a2, a3, life, salv, active)
        return "updated"
    cid = services.create_asset_category(name, a1, a2, a3, life, salv)
    if not active:
        services.update_asset_category(cid, name, a1, a2, a3, life, salv, False)
    return "created"


def _exp_asset_categories():
    ab = {a.id: a.code for a in services.list_accounts()}
    return [[c.name, ab.get(c.asset_account_id, ""), ab.get(c.accumulated_depreciation_account_id, ""), ab.get(c.depreciation_expense_account_id, ""), c.useful_life_months, c.salvage_percent,
             _yn(c.is_active)] for c in services.list_asset_categories(False)]


_register("asset_categories", "فئات الأصول الثابتة", "الأصول الثابتة",
          [("اسم الفئة*", "name"), ("كود حساب الأصل*", "asset"), ("كود حساب مجمع الإهلاك*", "accum"), ("كود حساب مصروف الإهلاك*", "expense"), ("العمر الإنتاجي (شهور)", "life"),
           ("نسبة القيمة التخريدية %", "salvage"), ("نشط (نعم/لا)", "is_active")], _exp_asset_categories, _imp_asset_category, ["الحسابات الثلاثة بأكوادها من شجرة الحسابات."])


# ======================================================================
# الموازنات
# ======================================================================
def _imp_budget(v, ctx):
    name = _s(v["name"]) or "موازنة"
    d1 = _date(v["from"], True, "من تاريخ"); d2 = _date(v["to"], True, "إلى تاريخ")
    acc = ctx.account_id(v["account"], True, "الحساب"); amount = _f(v["amount"], None) if v["amount"] not in (None, "") else None
    if amount is None:
        raise ValueError("مبلغ الموازنة مطلوب")
    branch = ctx.branch_id(v["branch"]); cc = ctx.cost_center_id(v["cc"])
    for b in services.list_budget_lines():
        if b.name == name and b.account_id == acc and b.date_from.date() == d1.date() and b.date_to.date() == d2.date() and (b.branch_id or 0) == (branch or 0) and (b.cost_center_id or 0) == (cc or 0):
            return None            # نفس البند موجود
    services.create_budget_line(name, d1, d2, acc, amount, branch, cc, _s(v["notes"]))
    return "created"


def _exp_budgets():
    br = {b.id: b.code for b in services.list_branches(False)}; cc = {c.id: c.code for c in services.list_cost_centers(False)}
    return [[b.name, _fmt_date(b.date_from), _fmt_date(b.date_to), b.account.code if b.account else "", b.budget_amount, br.get(b.branch_id, ""), cc.get(b.cost_center_id, ""), b.notes or ""]
            for b in services.list_budget_lines()]


_register("budgets", "الموازنات", "المحاسبة", [("الاسم*", "name"), ("من تاريخ*", "from"), ("إلى تاريخ*", "to"), ("كود الحساب*", "account"), ("مبلغ الموازنة*", "amount"), ("كود الفرع", "branch"), ("كود مركز التكلفة", "cc"), ("ملاحظات", "notes")],
          _exp_budgets, _imp_budget, ["البند المطابق تمامًا (الاسم والحساب والفترة) بيتجاهل."])


# ======================================================================
# مكونات التصنيع (BOM)
# ======================================================================
def _imp_bom(v, ctx):
    qty = _f(v["qty"], None) if v["qty"] not in (None, "") else None
    if qty is None:
        raise ValueError("الكمية مطلوبة")
    services.add_bom_line(ctx.product_id(v["product"], "الصنف النهائي"), ctx.product_id(v["component"], "المكوّن"), qty)
    return "updated"


def _exp_bom():
    rows = []
    for p in services.list_products(active_only=False):
        for l in services.list_bom_lines(p.id):
            rows.append([p.sku, p.name, l.component.sku if l.component else "", l.component.name if l.component else "", l.quantity])
    return rows


_register("bom", "مكونات التصنيع (BOM)", "الأصناف والمخزون", [("كود الصنف النهائي*", "product"), ("اسم الصنف (للعرض)", "pname"), ("كود المكوّن*", "component"), ("اسم المكوّن (للعرض)", "cname"), ("الكمية*", "qty")],
          _exp_bom, _imp_bom, ["كل صف = مكوّن واحد للصنف النهائي. المكوّن الموجود بيتحدّث كميته. أسماء الأصناف للعرض بس."])


# ======================================================================
# الفترات المحاسبية
# ======================================================================
def _imp_fiscal(v, ctx):
    name = _s(v["name"])
    if not name:
        raise ValueError("اسم الفترة مطلوب")
    d1 = _date(v["from"], True, "من تاريخ"); d2 = _date(v["to"], True, "إلى تاريخ")
    if name in {p.name for p in services.list_fiscal_periods()}:
        return None
    services.create_fiscal_period(name, d1, d2)
    return "created"


_register("fiscal_periods", "الفترات المحاسبية", "المحاسبة", [("الاسم*", "name"), ("من تاريخ*", "from"), ("إلى تاريخ*", "to"), ("مقفلة (للعرض بس)", "closed")],
          lambda: [[p.name, _fmt_date(p.date_from), _fmt_date(p.date_to), _yn(p.is_closed)] for p in services.list_fiscal_periods()], _imp_fiscal,
          ["الفترة الموجودة بنفس الاسم بتتجاهل. القفل بيتم من شاشة الفترات المحاسبية."])


# ======================================================================
# مجموعات الكاشير
# ======================================================================
def _imp_pos_group(v, ctx):
    name = _s(v["name"])
    if not name:
        raise ValueError("اسم المجموعة مطلوب")
    skus = [x.strip() for x in _s(v["skus"]).replace("،", ",").split(",") if x.strip()]
    ids = [ctx.product_id(x) for x in skus]
    ex = next((g for g in services.list_pos_groups(False) if g.name == name), None)
    if ex:
        services.update_pos_group(ex.id, color=_s(v["color"]), sort_order=_i(v["order"], 0), is_active=_yes_no(v["is_active"], True))
        services.set_pos_group_products(ex.id, ids)
        return "updated"
    gid = services.create_pos_group(name, _s(v["color"]), _i(v["order"], 0))
    services.set_pos_group_products(gid, ids)
    if not _yes_no(v["is_active"], True):
        services.update_pos_group(gid, is_active=False)
    return "created"


def _exp_pos_groups():
    pr = {p.id: p.sku for p in services.list_products(active_only=False)}
    return [[g.name, g.color or "", g.sort_order, _yn(g.is_active), ", ".join(pr.get(i, "") for i in services.get_pos_group_product_ids(g.id))] for g in services.list_pos_groups(False)]


_register("pos_groups", "مجموعات الكاشير (POS)", "الكاشير والمطعم", [("اسم المجموعة*", "name"), ("اللون", "color"), ("الترتيب", "order"), ("نشط (نعم/لا)", "is_active"), ("أكواد الأصناف (مفصولة بفاصلة)", "skus")],
          _exp_pos_groups, _imp_pos_group, ["أكواد الأصناف في خلية واحدة مفصولة بفاصلة: P001, P002, P003."])


# ======================================================================
# أقسام وطاولات المطعم
# ======================================================================
def _imp_rest_section(v, ctx):
    name = _s(v["name"])
    if not name:
        raise ValueError("اسم القسم مطلوب")
    if name in {s.name for s in services.list_restaurant_sections(False)}:
        return None
    services.create_restaurant_section(name, _i(v["order"], 0))
    return "created"


_register("restaurant_sections", "أقسام المطعم", "الكاشير والمطعم", [("اسم القسم*", "name"), ("الترتيب", "order")],
          lambda: [[s.name, s.sort_order] for s in services.list_restaurant_sections(False)], _imp_rest_section, ["القسم الموجود بنفس الاسم بيتجاهل."])


def _imp_rest_table(v, ctx):
    number = _s(v["number"])
    if not number:
        raise ValueError("رقم الطاولة مطلوب")
    if number in {t.number for t in services.list_restaurant_tables(active_only=False)}:
        return None
    sec = None
    if _s(v["section"]):
        sec = next((s.id for s in services.list_restaurant_sections(False) if s.name == _s(v["section"])), None)
        if sec is None:
            raise ValueError(f'القسم "{_s(v["section"])}" مش موجود')
    services.create_restaurant_table(number, _s(v["name"]), _i(v["seats"], 4), sec, ctx.branch_id(v["branch"]))
    return "created"


def _exp_rest_tables():
    sec = {s.id: s.name for s in services.list_restaurant_sections(False)}; br = {b.id: b.code for b in services.list_branches(False)}
    return [[t.number, t.name or "", t.seats, sec.get(t.section_id, ""), br.get(t.branch_id, "")] for t in services.list_restaurant_tables(active_only=False)]


_register("restaurant_tables", "طاولات المطعم", "الكاشير والمطعم", [("رقم الطاولة*", "number"), ("الاسم", "name"), ("عدد المقاعد", "seats"), ("القسم", "section"), ("كود الفرع", "branch")],
          _exp_rest_tables, _imp_rest_table, ["الطاولة الموجودة بنفس الرقم بتتجاهل. القسم لازم يكون موجود (استورد الأقسام الأول)."])


# ======================================================================
# أنواع المستندات / المستندات المتابَعة / قوالب المصروفات المقدمة
# ======================================================================
_KIND_AR = {"employee": "موظف", "company": "الشركة", "other": "أخرى"}
_KIND_FROM = {"موظف": "employee", "الشركة": "company", "شركة": "company", "أخرى": "other", "اخرى": "other", "employee": "employee", "company": "company", "other": "other"}


def _tiers(text):
    return services._parse_remind_days(_s(text))


def _imp_doc_type(v, ctx):
    name = _s(v["name"])
    if not name:
        raise ValueError("اسم النوع مطلوب")
    kind = _KIND_FROM.get(_s(v["kind"]) or "أخرى")
    if not kind:
        raise ValueError("نوع صاحب المستند: موظف أو الشركة أو أخرى")
    tiers = _tiers(v["tiers"]); active = _yes_no(v["is_active"], True)
    ex = next((t for t in services.list_document_types(False) if t.name == name), None)
    if ex:
        services.update_document_type(ex.id, name, kind, tiers, active)
        return "updated"
    tid = services.create_document_type(name, kind, tiers)
    if not active:
        services.update_document_type(tid, name, kind, tiers, False)
    return "created"


_register("document_types", "أنواع المستندات المتابَعة", "الموارد البشرية", [("الاسم*", "name"), ("خاص بـ (موظف/الشركة/أخرى)", "kind"), ("أيام التنبيه (مفصولة بفاصلة)", "tiers"), ("نشط (نعم/لا)", "is_active")],
          lambda: [[t.name, _KIND_AR.get(t.holder_kind, "أخرى"), t.remind_days, _yn(t.is_active)] for t in services.list_document_types(False)], _imp_doc_type,
          ["أيام التنبيه من القيم: 180,120,90,60,45,30,15,7,3,1,0 (0 = يوم الانتهاء)."])


def _imp_tracked_doc(v, ctx):
    tname = _s(v["type"])
    t = next((x for x in services.list_document_types(False) if x.name == tname), None)
    if not t:
        raise ValueError(f'نوع المستند "{tname}" مش موجود (استورد أنواع المستندات الأول)')
    expiry = _date(v["expiry"], True, "تاريخ الانتهاء")
    kind = _KIND_FROM.get(_s(v["kind"]) or _KIND_AR.get(t.holder_kind, "أخرى"))
    emp_id = None
    if kind == "employee":
        code = _s(v["emp_code"])
        e = next((x for x in ctx.employees() if x.code == code), None)
        if not e:
            raise ValueError(f'كود الموظف "{code}" مش موجود')
        emp_id = e.id
    number = _s(v["number"])
    for d in services.list_tracked_documents(include_archived=True):
        if d["type_id"] == t.id and (d["doc_number"] or "") == number and d["expiry_date"].date() == expiry.date() and d["employee_id"] == emp_id and d["holder_name"] == _s(v["holder"]):
            return None
    tiers = _tiers(v["tiers"]) or None
    services.create_tracked_document(t.id, expiry, _s(v["title"]), kind, emp_id, _s(v["holder"]), number, _s(v["issuer"]), _date(v["issue"], False, "تاريخ الإصدار"), tiers, _f(v["cost"], 0.0), _s(v["notes"]))
    return "created"


def _exp_tracked_docs():
    emp = {e.id: e.code for e in services.list_employees(active_only=False)}
    return [[d["type_name"], d["title"], _KIND_AR.get(d["holder_kind"], "أخرى"), emp.get(d["employee_id"], ""), d["holder_name"], d["doc_number"], d["issuer"], _fmt_date(d["issue_date"]),
             _fmt_date(d["expiry_date"]), ",".join(str(x) for x in d["remind_days"]), d["cost"], d["notes"]] for d in services.list_tracked_documents(include_archived=True)]


_register("tracked_documents", "المستندات والرخص المتابَعة", "الموارد البشرية",
          [("نوع المستند*", "type"), ("العنوان", "title"), ("خاص بـ (موظف/الشركة/أخرى)", "kind"), ("كود الموظف", "emp_code"), ("اسم صاحب المستند (لغير الموظف)", "holder"), ("رقم المستند", "number"),
           ("جهة الإصدار", "issuer"), ("تاريخ الإصدار", "issue"), ("تاريخ الانتهاء*", "expiry"), ("أيام التنبيه (فاضي = الافتراضي)", "tiers"), ("التكلفة", "cost"), ("ملاحظات", "notes")],
          _exp_tracked_docs, _imp_tracked_doc, ["نوع المستند لازم يكون موجود بنفس الاسم. للموظف: اكتب كود الموظف (EMP-...). المستند المطابق تمامًا بيتجاهل."])


def _imp_prepaid_tpl(v, ctx):
    name = _s(v["name"])
    if not name:
        raise ValueError("اسم القالب مطلوب")
    acc = ctx.account_id(v["account"], False, "حساب المصروف")
    months = _i(v["months"], 12); method = {"متساوية": "equal", "بالأيام": "daily", "equal": "equal", "daily": "daily"}.get(_s(v["method"]) or "متساوية")
    if not method:
        raise ValueError("طريقة التقسيم: متساوية أو بالأيام")
    active = _yes_no(v["is_active"], True)
    ex = next((t for t in services.list_prepaid_templates(False) if t.name == name), None)
    if ex:
        services.update_prepaid_template(ex.id, name, acc, months, method, ex.notes or "", active)
        return "updated"
    tid = services.create_prepaid_template(name, acc, months, method)
    if not active:
        services.update_prepaid_template(tid, name, acc, months, method, "", False)
    return "created"


_register("prepaid_templates", "قوالب المصروفات المدفوعة مقدمًا", "المحاسبة", [("الاسم*", "name"), ("كود حساب المصروف", "account"), ("المدة (شهور)", "months"), ("طريقة التقسيم (متساوية/بالأيام)", "method"), ("نشط (نعم/لا)", "is_active")],
          lambda: [[t.name, t.expense_account.code if t.expense_account else "", t.months, "بالأيام" if t.method == "daily" else "متساوية", _yn(t.is_active)] for t in services.list_prepaid_templates(False)],
          _imp_prepaid_tpl, ["القالب الموجود بنفس الاسم بيتحدّث."])


for _k, _g in GROUPS.items():
    if _k in ENTITY_SPECS:
        ENTITY_SPECS[_k]["group"] = _g
for _k, _sp in ENTITY_SPECS.items():
    _sp.setdefault("group", "أخرى")


# ======================================================================
# تصدير أي جدول من قاعدة البيانات (كل الجداول)
# ======================================================================
TABLE_LABELS = {
    "accounts": "شجرة الحسابات", "accounting_settings": "إعدادات المحاسبة", "agenda_items": "جدول الأعمال", "allowance_insurance_items": "بنود المخصصات والتأمينات",
    "allowance_insurance_runs": "دورات المخصصات والتأمينات", "annual_leave_settlements": "تسويات الإجازات السنوية", "asset_categories": "فئات الأصول الثابتة",
    "asset_inventories": "جرد الأصول الثابتة", "asset_inventory_lines": "بنود جرد الأصول", "bom_lines": "مكونات التصنيع", "branches": "الفروع", "budget_lines": "الموازنات",
    "categories": "فئات الأصناف", "chat_conversations": "المحادثات", "chat_members": "أعضاء المحادثات", "chat_messages": "الرسائل", "company_settings": "إعدادات الشركة",
    "contracts": "عقود الموظفين", "cost_centers": "مراكز التكلفة", "custody_expenses": "مصروفات العهد", "custody_fundings": "سندات صرف العهد", "custody_settlements": "تسويات العهد",
    "customers": "العملاء", "damage_items": "أصناف التوالف", "damage_lines": "بنود التوالف", "damage_records": "سجلات التوالف", "depreciation_run_items": "بنود دورات الإهلاك",
    "depreciation_runs": "دورات الإهلاك", "document_alerts": "تنبيهات المستندات", "document_renewals": "تجديدات المستندات", "document_types": "أنواع المستندات", "employee_custodies": "عهد الموظفين",
    "employees": "الموظفون", "eos_settlements": "تسويات نهاية الخدمة", "fiscal_periods": "الفترات المحاسبية", "fixed_assets": "الأصول الثابتة", "goods_receipt_note_items": "بنود سندات الاستلام",
    "goods_receipt_notes": "سندات استلام المشتريات", "inventory_adjustment_lines": "بنود التسويات الجردية", "inventory_adjustments": "التسويات الجردية", "journal_entries": "القيود اليومية",
    "journal_entry_lines": "بنود القيود", "leave_requests": "طلبات الإجازات", "locations": "المخازن", "payments": "سندات القبض والصرف", "payment_methods": "طرق الدفع",
    "payroll_runs": "دورات الرواتب", "payslips": "قسائم الرواتب", "permission_requests": "طلبات الاستئذان", "pos_group_items": "أصناف مجموعات الكاشير", "pos_groups": "مجموعات الكاشير",
    "prepaid_expense_lines": "أقساط المصروفات المقدمة", "prepaid_expenses": "المصروفات المدفوعة مقدمًا", "prepaid_templates": "قوالب المصروفات المقدمة", "price_list_items": "أسعار قوائم الأسعار",
    "price_lists": "قوائم الأسعار", "product_unit_lines": "وحدات وأسعار الأصناف", "product_units": "وحدات القياس", "production_orders": "أوامر التصنيع", "products": "الأصناف",
    "purchase_invoice_attachments": "مرفقات فواتير المشتريات", "purchase_invoice_items": "بنود فواتير المشتريات", "purchase_invoices": "فواتير المشتريات",
    "resignation_requests": "طلبات الاستقالة", "restaurant_order_items": "بنود طلبات المطعم", "restaurant_orders": "طلبات المطعم", "restaurant_sections": "أقسام المطعم",
    "restaurant_tables": "طاولات المطعم", "sales_invoice_items": "بنود فواتير المبيعات", "sales_invoices": "فواتير المبيعات", "sales_quotation_items": "بنود عروض الأسعار",
    "sales_quotations": "عروض الأسعار", "stock_moves": "حركات المخزون", "suppliers": "الموردون", "tracked_documents": "المستندات والرخص المتابَعة", "users": "المستخدمون",
}
ADMIN_ONLY_TABLES = {"users", "chat_conversations", "chat_members", "chat_messages", "agenda_items", "document_alerts"}
HIDDEN_COLUMNS = {"password_hash", "smtp_password", "zatca_password"}
COL_LABELS = {"id": "المعرّف", "code": "الكود", "name": "الاسم", "number": "الرقم", "date": "التاريخ", "notes": "ملاحظات", "description": "البيان", "reference": "المرجع", "is_active": "نشط",
               "created_at": "تاريخ الإنشاء", "created_by": "أنشأه", "state": "الحالة", "total": "الإجمالي", "subtotal": "قبل الضريبة", "tax_total": "الضريبة", "tax_rate": "نسبة الضريبة %",
               "tax_amount": "مبلغ الضريبة", "discount": "الخصم", "quantity": "الكمية", "unit_price": "سعر الوحدة", "unit_cost": "تكلفة الوحدة", "cost_price": "سعر التكلفة", "sale_price": "سعر البيع",
               "amount": "المبلغ", "debit": "مدين", "credit": "دائن", "voided": "ملغاة", "posted": "مرحّل", "line_total": "إجمالي السطر", "phone": "الهاتف", "email": "البريد",
               "address": "العنوان", "vat_number": "الرقم الضريبي", "sku": "كود الصنف", "unit": "الوحدة", "username": "اسم المستخدم", "full_name": "الاسم الكامل", "role": "الدور",
               "sort_order": "الترتيب", "account_type": "نوع الحساب", "cash_flow_type": "نوع التدفق النقدي", "title": "العنوان", "expiry_date": "تاريخ الانتهاء", "start_date": "تاريخ البداية",
               "date_from": "من تاريخ", "date_to": "إلى تاريخ", "balance": "الرصيد", "is_pos": "كاشير", "payment_method": "طريقة الدفع", "kind": "النوع"}
FK_LABELS = {"customer": "العميل", "supplier": "المورد", "product": "الصنف", "component": "المكوّن", "account": "الحساب", "branch": "الفرع", "cost_center": "مركز التكلفة", "employee": "الموظف",
             "location": "المخزن", "entry": "القيد", "invoice": "الفاتورة", "category": "الفئة", "parent": "الأب", "unit": "الوحدة", "payment_method": "طريقة الدفع", "price_list": "قائمة الأسعار",
             "section": "القسم", "template": "القالب", "type": "النوع", "document": "المستند", "prepaid": "المصروف المقدم", "custody": "العهدة", "run": "الدورة", "conversation": "المحادثة",
             "quotation": "عرض السعر", "receipt": "سند الاستلام", "order": "الطلب", "group": "المجموعة", "asset": "الأصل", "inventory": "الجرد", "adjustment": "التسوية", "damage": "سجل التوالف",
             "contract": "العقد", "journal_entry": "القيد", "initial_entry": "قيد الدفع", "closing_journal": "قيد الإقفال", "sales_invoice": "فاتورة المبيعات", "purchase_invoice": "فاتورة المشتريات",
             "return_funding": "سند الإرجاع", "table": "الطاولة", "pos_group": "مجموعة الكاشير", "default_branch": "الفرع الافتراضي"}

COL_LABELS.update({
    "accrued_days": "أيام مستحقة", "accumulated_depreciation": "مجمع الإهلاك", "alerted_at": "وقت التنبيه", "all_day": "طوال اليوم", "anchor_day": "يوم التكرار", "annual_leave_days": "أيام الإجازة السنوية",
    "archived": "مؤرشف", "assigned_to": "مكلَّف", "attachment_kind": "نوع المرفق", "attachment_name": "اسم المرفق", "attachment_path": "مسار المرفق", "attachment_size": "حجم المرفق", "auto_post": "ترحيل آلي",
    "bank_account": "الحساب البنكي", "barcode": "الباركود", "basic_salary": "الراتب الأساسي", "body": "نص الرسالة", "book_value": "القيمة الدفترية", "branch_ids": "الفروع المسموحة", "budget_amount": "مبلغ الموازنة",
    "can_be_purchased": "يُشترى", "can_be_sold": "يُباع", "category": "الفئة", "closed_at": "تاريخ الإقفال", "closing_balance": "الرصيد الختامي", "color": "اللون", "contract_number": "رقم العقد",
    "contract_type": "نوع العقد", "cost": "التكلفة", "cost_center_ids": "مراكز التكلفة المسموحة", "counted_qty": "الكمية المعدودة", "course": "الدورة", "customer_name": "اسم العميل", "daily_rate": "الأجر اليومي",
    "days": "الأيام", "decided_at": "تاريخ القرار", "decided_by": "القرار بواسطة", "decision_note": "ملاحظة القرار", "default_tax_rate": "نسبة الضريبة الافتراضية", "deleted": "محذوفة", "department": "القسم",
    "difference_qty": "فرق الكمية", "direction": "الاتجاه", "disposal_amount": "مبلغ الاستبعاد", "disposal_date": "تاريخ الاستبعاد", "disposal_loss": "خسارة الاستبعاد", "disposed_at": "تاريخ الاستبعاد",
    "doc_number": "رقم المستند", "done_at": "تاريخ الإنجاز", "done_count": "عدد مرات الإنجاز", "employer_gosi_amount": "تأمينات صاحب العمل", "encashed_amount": "مبلغ التعويض النقدي", "encashed_days": "أيام التعويض",
    "end_date": "تاريخ النهاية", "entry_date": "تاريخ القيد", "eos_gratuity": "مكافأة نهاية الخدمة", "eos_gross": "إجمالي نهاية الخدمة", "eos_multiplier": "معامل نهاية الخدمة", "eos_provision_amount": "مخصص نهاية الخدمة",
    "eos_provision_rate": "نسبة مخصص نهاية الخدمة", "eos_provisioned": "مخصص نهاية الخدمة المحمَّل", "expense_type": "نوع المصروف", "extra_allowance": "بدل إضافي", "extra_allowance_note": "ملاحظة البدل الإضافي",
    "extra_reason": "سبب الإضافي", "factor_to_base": "معامل التحويل للوحدة الأساسية", "file_name": "اسم الملف", "file_path": "مسار الملف", "funded": "المموَّل", "gosi_deduction": "خصم التأمينات",
    "gosi_rate_employee": "نسبة تأمينات الموظف", "gosi_rate_employer": "نسبة تأمينات صاحب العمل", "guests": "عدد الضيوف", "hire_date": "تاريخ التعيين", "holder_kind": "صاحب المستند", "holder_name": "اسم صاحب المستند",
    "hours": "الساعات", "housing_allowance": "بدل السكن", "image_path": "مسار الصورة", "invoice_date": "تاريخ الفاتورة", "invoice_number": "رقم الفاتورة", "invoice_prefix_purchase": "بادئة فواتير المشتريات",
    "invoice_prefix_sales": "بادئة فواتير المبيعات", "invoiced": "اتحوّل لفاتورة", "is_closed": "مقفلة", "is_credit": "آجل", "is_default": "افتراضي", "is_system": "رسالة نظام", "issue_date": "تاريخ الإصدار",
    "issuer": "جهة الإصدار", "joined_at": "تاريخ الانضمام", "joined_message_id": "أول رسالة مرئية", "kitchen_status": "حالة المطبخ", "last_leave_date": "آخر إجازة", "last_message_at": "آخر رسالة",
    "last_read_message_id": "آخر رسالة مقروءة", "last_working_date": "آخر يوم عمل", "leave_amount": "مبلغ الإجازة", "leave_provision_amount": "مخصص الإجازة", "leave_provision_rate": "نسبة مخصص الإجازة",
    "leave_provisioned": "مخصص الإجازة المحمَّل", "leave_type": "نوع الإجازة", "left": "غادر", "logo_path": "مسار الشعار", "method": "الطريقة", "month": "الشهر", "months": "عدد الشهور", "monthly_wage": "الأجر الشهري",
    "national_id": "رقم الهوية", "net_salary": "صافي الراتب", "new_expiry": "الانتهاء الجديد", "next_alert_at": "التنبيه القادم", "note": "ملاحظة", "old_expiry": "الانتهاء القديم", "opened_at": "وقت الفتح",
    "opened_by": "فتحه", "opening_balance": "الرصيد الافتتاحي", "order_type": "نوع الطلب", "other_allowances": "بدلات أخرى", "other_deductions": "خصومات أخرى", "other_deductions_note": "ملاحظة الخصومات",
    "other_dues": "مستحقات أخرى", "other_dues_note": "ملاحظة المستحقات", "paid_at": "تاريخ الدفع", "partner_id": "الطرف (المعرّف)", "partner_type": "نوع الطرف", "party_type": "نوع الطرف", "payment_kind": "طريقة السداد",
    "payment_method_name": "طريقة الدفع", "payment_type": "نوع السند", "period_from": "الفترة من", "period_start": "بداية الفترة", "period_to": "الفترة إلى", "permissions": "الصلاحيات", "position": "المسمى الوظيفي",
    "posted_at": "تاريخ الترحيل", "previous_gross_salary": "الراتب السابق", "price": "السعر", "print_color": "لون الطباعة", "print_font_size": "حجم خط الطباعة", "print_footer_note": "تذييل الطباعة",
    "print_paper_size": "حجم ورق الطباعة", "print_show_logo": "إظهار الشعار", "print_show_qr": "إظهار QR", "priority": "الأولوية", "product_type": "نوع الصنف", "purchase_date": "تاريخ الشراء", "purpose": "الغرض",
    "raise_amount": "مبلغ الزيادة", "raise_note": "سبب الزيادة", "reason": "السبب", "remind_before_minutes": "التنبيه قبل (دقيقة)", "remind_days": "أيام التنبيه", "renewal_count": "عدد التجديدات",
    "renewed_at": "تاريخ التجديد", "reorder_point": "حد الطلب", "repeat": "التكرار", "requested_at": "تاريخ الطلب", "returned_before": "مرتجع سابقًا", "returned_now": "مرتجع الآن", "salvage_percent": "نسبة الخردة %",
    "salvage_value": "القيمة التخريدية", "scan_count": "عدد مرات المسح", "scanned_at": "وقت المسح", "seats": "المقاعد", "sender": "المرسل", "sent_at": "وقت الإرسال", "seq": "التسلسل", "settlement_date": "تاريخ التسوية",
    "snoozed_until": "مؤجّل حتى", "source_id": "مصدر القيد (المعرّف)", "source_type": "نوع المصدر", "spent": "المصروف", "start_at": "موعد البداية", "status": "الحالة", "submission_date": "تاريخ التقديم",
    "supplier_name": "اسم المورد", "supplier_vat": "الرقم الضريبي للمورد", "termination_date": "تاريخ الإنهاء", "termination_reason": "سبب الإنهاء", "termination_type": "نوع الإنهاء", "terms": "الشروط",
    "theoretical_qty": "الكمية الدفترية", "tier": "مرحلة التنبيه", "time_from": "من الساعة", "time_to": "إلى الساعة", "total_amount": "إجمالي المبلغ", "total_cost": "إجمالي التكلفة", "transport_allowance": "بدل الانتقال",
    "unit_factor": "معامل الوحدة", "unused_leave_days": "أيام إجازة غير مستخدمة", "uploaded_at": "تاريخ الرفع", "used_days": "أيام مستخدمة", "useful_life_months": "العمر الإنتاجي (شهور)", "valid_until": "ساري حتى",
    "year": "السنة", "years_of_service": "سنوات الخدمة", "zatca_api_url": "رابط الزكاة", "zatca_enabled": "تفعيل الزكاة", "zatca_environment": "بيئة الزكاة", "zatca_response": "رد الزكاة", "zatca_status": "حالة الزكاة",
    "zatca_username": "مستخدم الزكاة", "smtp_host": "خادم البريد", "smtp_port": "منفذ البريد", "smtp_username": "مستخدم البريد", "smtp_from_email": "بريد المرسل", "smtp_use_tls": "تشفير TLS", "code_prefix": "بادئة الكود",
    "code_start": "بداية الترقيم", "code_next": "الرقم التالي", "code_width": "عرض الرقم", "code_key": "مفتاح الترقيم", "scope_type": "نطاق (النوع)", "scope_id": "نطاق (المعرّف)", "scope_label": "النطاق", "direct_key": "مفتاح المحادثة",
    "closing_journal_id": "قيد الإقفال (المعرّف)", "journal_entry_id": "القيد (المعرّف)", "disposal_journal_id": "قيد الاستبعاد (المعرّف)", "initial_entry_id": "قيد الدفع (المعرّف)"})
_ACC_LABELS = {"accumulated": "مجمع الإهلاك", "accumulated_depreciation": "مجمع الإهلاك", "asset": "الأصل", "asset_loss": "خسائر الأصول", "basic_salary": "الراتب الأساسي", "cash": "النقدية", "cogs": "تكلفة المبيعات",
               "damage": "التوالف", "depreciation_expense": "مصروف الإهلاك", "employee_advances": "سلف الموظفين", "employee_custody": "عهد الموظفين", "employer_gosi_expense": "مصروف تأمينات صاحب العمل",
               "employer_gosi_payable": "تأمينات صاحب العمل المستحقة", "eos_provision_expense": "مصروف مخصص نهاية الخدمة", "eos_provision_payable": "مخصص نهاية الخدمة", "expense": "المصروف",
               "extra_allowance": "البدل الإضافي", "gosi_deduction": "خصم التأمينات", "gosi_employer_expense": "مصروف تأمينات صاحب العمل", "gosi_employer_payable": "تأمينات صاحب العمل المستحقة",
               "gosi_payable": "التأمينات المستحقة", "housing_allowance": "بدل السكن", "inventory": "المخزون", "inventory_gain": "أرباح الجرد", "inventory_loss": "خسائر الجرد",
               "leave_provision_expense": "مصروف مخصص الإجازات", "leave_provision_payable": "مخصص الإجازات", "other_allowances": "البدلات الأخرى", "other_deductions": "الخصومات الأخرى", "payable": "الموردين",
               "prepaid": "المدفوع مقدمًا", "prepaid_expenses": "المصروفات المدفوعة مقدمًا", "purchase_discount": "خصم المشتريات", "purchases": "المشتريات", "receivable": "العملاء", "retained_earnings": "الأرباح المرحلة",
               "salaries_payable": "الرواتب المستحقة", "salary_expense": "مصروف الرواتب", "sales": "المبيعات", "sales_discount": "خصم المبيعات", "transport_allowance": "بدل الانتقال", "vat_input": "ضريبة المدخلات",
               "vat_output": "ضريبة المخرجات"}
FK_LABELS.update({"item": "البند", "payment": "السند", "record": "السجل", "renewed_from": "مجدَّد من", "resignation_request": "طلب الاستقالة"})
FK_LABELS.update({k + "_account": "حساب " + v for k, v in _ACC_LABELS.items()})
_DISPLAY_PREF = ("code", "name", "number", "sku", "username", "title", "full_name")


def raw_tables(role_admin=False):
    """[(اسم الجدول، الوصف العربي، عدد الصفوف)] لكل جداول القاعدة (الحساسة للأدمن بس)."""
    out = []
    session = SessionLocal()
    try:
        from sqlalchemy import func
        for t in Base.metadata.sorted_tables:
            if t.name in ADMIN_ONLY_TABLES and not role_admin:
                continue
            n = session.execute(select(func.count()).select_from(t)).scalar() or 0
            out.append((t.name, TABLE_LABELS.get(t.name, t.name), int(n)))
    finally:
        session.close()
    return sorted(out, key=lambda x: x[1])


def _col_label(col):
    if col.name in COL_LABELS:
        return COL_LABELS[col.name]
    return col.name.replace("_", " ")


def _display_cols(ref_table):
    cols = [c for c in _DISPLAY_PREF if c in ref_table.columns]
    return cols[:2] if len(cols) >= 2 and cols[0] == "code" and "name" in cols else cols[:1]


def _write_table_sheet(ws, session, table):
    cols = [c for c in table.columns if c.name not in HIDDEN_COLUMNS]
    fk_maps = {}
    for c in cols:
        for fk in c.foreign_keys:
            ref = fk.column.table
            dc = _display_cols(ref)
            if dc:
                data = session.execute(select(ref.c[fk.column.name], *[ref.c[x] for x in dc])).all()
                fk_maps[c.name] = {r[0]: " - ".join(str(x) for x in r[1:] if x not in (None, "")) for r in data}
    headers = []
    for c in cols:
        if c.name in fk_maps:
            base = c.name[:-3] if c.name.endswith("_id") else c.name
            headers.append(FK_LABELS.get(base, base.replace("_", " ")) + " (المعرّف)")
            headers.append(FK_LABELS.get(base, base.replace("_", " ")))
        else:
            headers.append(_col_label(c))
    ws.append(headers)
    _style_header(ws, len(headers))
    for row in session.execute(select(*cols)).all():
        out = []
        for c, v in zip(cols, row):
            if isinstance(v, bool):
                v = _yn(v)
            elif isinstance(v, datetime.datetime) and v.tzinfo:
                v = v.replace(tzinfo=None)
            out.append(v if v is not None else "")
            if c.name in fk_maps:
                out.append(fk_maps[c.name].get(v, "") if v is not None else "")
        ws.append(out)
    for i in range(1, len(headers) + 1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = 18
    ws.freeze_panes = "A2"
    return len(ws["A"]) - 1


def _check_raw_access(table_name):
    if table_name not in Base.metadata.tables:
        raise ValueError(f"الجدول {table_name} غير موجود")
    if table_name in ADMIN_ONLY_TABLES and (services.current_user() or {}).get("role") != "admin":
        raise PermissionError("تصدير هذا الجدول لمدير النظام فقط")


def export_raw_table(table_name, output_path):
    """تصدير جدول واحد (أي جدول) بأعمدة عربية، والأعمدة المرتبطة بتتوسّع لاسم مقروء (العميل، الصنف، الحساب...)."""
    _check_raw_access(table_name)
    wb = openpyxl.Workbook(); ws = wb.active
    ws.title = TABLE_LABELS.get(table_name, table_name)[:31]
    session = SessionLocal()
    try:
        n = _write_table_sheet(ws, session, Base.metadata.tables[table_name])
    finally:
        session.close()
    wb.save(output_path)
    return {"path": output_path, "rows": n}


def export_all_tables(output_path):
    """كل جداول القاعدة في ملف Excel واحد (ورقة لكل جدول + ورقة فهرس). جداول الأدمن (مستخدمين/رسائل...) لمدير النظام فقط."""
    is_admin = (services.current_user() or {}).get("role") == "admin"
    wb = openpyxl.Workbook(); idx = wb.active; idx.title = "فهرس"
    idx.append(["الجدول", "الوصف", "عدد الصفوف", "اسم الورقة"]); _style_header(idx, 4)
    used = {"فهرس"}; total = 0; tables = 0
    session = SessionLocal()
    try:
        for t in Base.metadata.sorted_tables:
            if t.name in ADMIN_ONLY_TABLES and not is_admin:
                continue
            title = (TABLE_LABELS.get(t.name, t.name))[:31]
            base, k = title, 2
            while title in used:
                title = f"{base[:28]}_{k}"; k += 1
            used.add(title)
            ws = wb.create_sheet(title)
            n = _write_table_sheet(ws, session, t)
            idx.append([t.name, TABLE_LABELS.get(t.name, t.name), n, title]); total += n; tables += 1
    finally:
        session.close()
    for col, w in (("A", 28), ("B", 34), ("C", 14), ("D", 34)):
        idx.column_dimensions[col].width = w
    wb.save(output_path)
    return {"path": output_path, "tables": tables, "rows": total}
