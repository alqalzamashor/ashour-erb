"""تحويل شجرة الحسابات للشجرة الجديدة (erp_app/chart_data.py أو ملف Excel بنفس التنسيق):
 1) إضافة/تحديث حسابات الشجرة الجديدة.
 2) تحويل كل ربط كان على حساب افتراضي قديم (إعدادات المحاسبة، طرق الدفع، الأصناف، العقود، العملاء/الموردين،
    الأصول، الموازنات...) للحساب المقابل في الشجرة الجديدة.
 3) اختياريًا: نقل القيود والحركات القديمة للحسابات الجديدة المقابلة.
 4) حذف الحسابات القديمة اللي بقت من غير أي استخدام.
 كل ده في معاملة واحدة (لو حصل أي خطأ مفيش حاجة بتتغيّر) مع نسخة احتياطية من قاعدة البيانات قبل التنفيذ."""
import datetime
import os
import re
import shutil

from sqlalchemy import update, select, func

from . import services
from .chart_data import (NEW_CHART, OLD_TO_NEW_CODES, OLD_TO_NEW_GROUPS, NEW_SETTINGS_CODES, PAYROLL_FIELD_NEW_CODES, payment_method_new_code)
from .db import Base, SessionLocal, engine, Account, AccountingSettings, PaymentMethod

_TYPE = {"أصول": "asset", "التزامات": "liability", "حقوق ملكية": "equity", "إيرادات": "income", "مصروفات": "expense",
         "اهلاك الاصول": "expense", "إهلاك الأصول": "expense",
         "asset": "asset", "liability": "liability", "equity": "equity", "income": "income", "expense": "expense"}
_CF = {"نقدية": "cash", "تشغيلي": "operating", "استثماري": "investing", "تمويلي": "financing",
       "cash": "cash", "operating": "operating", "investing": "investing", "financing": "financing"}
_PAYROLL_OLD_CODES = ("5400", "2400", "2300")
_CASH_NEW_CODES = {"121001", "121002", "121003", "121004", "121005", "121006"}


def load_chart_rows(path=None):
    """بيرجّع (الصفوف، التحذيرات). path=None = الشجرة الجديدة المدمجة في البرنامج. أو ملف Excel بنفس تنسيق التصدير
    (الأنواع الناقصة بتاخد نوع الأب، ونوع 'اهلاك الاصول' بيتحوّل لمصروفات)."""
    if not path:
        return list(NEW_CHART), []
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.worksheets[0]
    rows, by_code, warnings = [], {}, []
    for n, r in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if not r or r[0] in (None, ""):
            continue
        code = str(r[0]).strip()
        name = re.sub(r"\s+", " ", str(r[1] or "")).strip()
        typ = _TYPE.get(str(r[2]).strip()) if len(r) > 2 and r[2] else None
        parent = str(r[3]).strip() if len(r) > 3 and r[3] not in (None, "") else None
        cf = _CF.get(str(r[4]).strip()) if len(r) > 4 and r[4] else None
        if not name:
            raise ValueError(f"الصف {n}: اسم الحساب فاضي")
        if code in by_code:
            raise ValueError(f"الصف {n}: الكود {code} مكرر في الملف")
        if parent and parent not in by_code:
            raise ValueError(f"الصف {n}: الحساب الأب {parent} مش موجود قبل الصف ده في الملف")
        if typ is None and parent:
            typ = by_code[parent][2]
        if typ is None:
            raise ValueError(f"الصف {n}: نوع الحساب {code} غير معروف ومفيش أب يتورّث منه")
        if code in _CASH_NEW_CODES:
            cf = "cash"
        by_code[code] = (code, name, typ, parent, cf)
        rows.append(by_code[code])
    if not rows:
        raise ValueError("الملف فاضي")
    return rows, warnings


def _fk_columns():
    """كل أعمدة قاعدة البيانات اللي بتشاور على accounts.id (جدول، عمود)."""
    out = []
    for t in Base.metadata.sorted_tables:
        for c in t.columns:
            for fk in c.foreign_keys:
                if fk.column.table.name == "accounts" and fk.column.name == "id":
                    out.append((t, c))
    return out


def _count_refs(session, account_id, cols, skip=()):
    total, per = 0, {}
    for t, c in cols:
        if (t.name, c.name) in skip:
            continue
        n = session.execute(select(func.count()).select_from(t).where(c == account_id)).scalar() or 0
        if n:
            per[f"{t.name}.{c.name}"] = n; total += n
    return total, per


def _old_default_accounts(session, new_codes):
    """الحسابات القديمة الافتراضية الموجودة فعلًا (كود قديم مش موجود في الشجرة الجديدة) + مقابلها."""
    pairs = {}
    for old_code, new_code in OLD_TO_NEW_CODES.items():
        if old_code in new_codes:
            continue
        old = session.query(Account).filter(Account.code == old_code).first()
        if not old:
            continue
        pairs[old.id] = (old, new_code)
    return pairs


def plan_chart_import(rows=None, move_history=True, remove_old=True):
    """معاينة من غير أي تغيير: إيه اللي هيتضاف/يتحدّث، وكل حساب قديم هيتحوّل لمين وكام ربط/قيد عليه."""
    rows = rows or list(NEW_CHART)
    new_codes = {r[0] for r in rows}
    session = SessionLocal()
    try:
        existing = {a.code for a in session.query(Account).all()}
        cols = _fk_columns()
        pairs = _old_default_accounts(session, new_codes)
        mapping = []
        for old_id, (old, new_code) in sorted(pairs.items(), key=lambda kv: kv[1][0].code):
            target = next((r for r in rows if r[0] == new_code), None)
            skip = () if move_history else (("journal_entry_lines", "account_id"),)
            je = session.execute(select(func.count()).select_from(Base.metadata.tables["journal_entry_lines"]).where(Base.metadata.tables["journal_entry_lines"].c.account_id == old_id)).scalar() or 0
            refs, per = _count_refs(session, old_id, [x for x in cols if (x[0].name, x[1].name) != ("accounts", "parent_id")])
            kids = session.query(Account).filter(Account.parent_id == old_id).count()
            mapping.append({"old_code": old.code, "old_name": old.name, "new_code": new_code if target else None, "new_name": target[1] if target else "(غير موجود في الشجرة)",
                            "journal_lines": je, "other_refs": refs - je, "refs_detail": per, "children": kids})
        custom = []
        for a in session.query(Account).all():
            if a.code in new_codes or a.id in pairs:
                continue
            parent = session.get(Account, a.parent_id) if a.parent_id else None
            custom.append({"code": a.code, "name": a.name, "parent_code": parent.code if parent else None,
                           "new_parent_code": OLD_TO_NEW_GROUPS.get(parent.code) if parent and parent.id in pairs else (parent.code if parent else None)})
        return {"to_create": len([r for r in rows if r[0] not in existing]), "to_update": len([r for r in rows if r[0] in existing]), "total_rows": len(rows),
                "mapping": mapping, "custom_accounts": custom, "journal_lines_to_move": sum(m["journal_lines"] for m in mapping) if move_history else 0,
                "has_old_defaults": bool(pairs), "move_history": move_history, "remove_old": remove_old}
    finally:
        session.close()


def _backup_database():
    path = engine.url.database
    if not path or not os.path.isfile(path):
        return None
    folder = os.path.join(os.path.dirname(os.path.abspath(path)), "backups")
    os.makedirs(folder, exist_ok=True)
    dest = os.path.join(folder, f"erp_before_chart_import_{datetime.datetime.now():%Y%m%d_%H%M%S}.db")
    shutil.copy2(path, dest)
    return dest


def apply_chart_import(rows=None, move_history=True, remove_old=True, backup=True):
    """تنفيذ التحويل (معاملة واحدة). بيرجّع تقرير بكل اللي اتعمل."""
    services.require_action("accounts", "delete")
    rows = rows or list(NEW_CHART)
    new_codes = {r[0] for r in rows}
    backup_path = _backup_database() if backup else None
    session = SessionLocal()
    report = {"backup": backup_path, "created": 0, "updated": 0, "relinked": {}, "journal_lines_moved": 0, "removed": [], "kept": [], "reparented": [],
              "settings_filled": [], "payment_methods": [], "payroll_overrides": 0, "warnings": []}
    try:
        def totals():
            jl = Base.metadata.tables["journal_entry_lines"]
            d, c = session.execute(select(func.coalesce(func.sum(jl.c.debit), 0), func.coalesce(func.sum(jl.c.credit), 0))).one()
            return round(float(d), 2), round(float(c), 2)
        before = totals()

        # ---- 1) الشجرة الجديدة
        by_code = {a.code: a for a in session.query(Account).all()}
        for code, name, typ, parent_code, cf in rows:
            a = by_code.get(code)
            if a:
                a.name = name; a.account_type = typ; a.cash_flow_type = cf; a.is_active = True
                report["updated"] += 1
            else:
                a = Account(code=code, name=name, account_type=typ, cash_flow_type=cf, is_active=True)
                session.add(a); by_code[code] = a
                report["created"] += 1
        session.flush()
        for code, name, typ, parent_code, cf in rows:
            by_code[code].parent_id = by_code[parent_code].id if parent_code else None
        session.flush()

        # ---- 2) خريطة القديم -> الجديد
        pairs = _old_default_accounts(session, new_codes)
        old_map, group_map = {}, {}
        for old_id, (old, new_code) in pairs.items():
            tgt = by_code.get(new_code)
            if not tgt:
                report["warnings"].append(f"مقابل الحساب {old.code} ({new_code}) مش موجود في الشجرة - اتساب من غير تحويل")
                continue
            old_map[old_id] = tgt.id
            grp = by_code.get(OLD_TO_NEW_GROUPS.get(old.code, ""))
            group_map[old_id] = grp.id if grp else tgt.id
        old_ids_by_code = {pairs[i][0].code: i for i in pairs}
        salary_old_ids = {old_ids_by_code[c] for c in _PAYROLL_OLD_CODES if c in old_ids_by_code}
        cols = _fk_columns()
        tables = Base.metadata.tables

        # ---- 3) استثناءات: طرق الدفع بالاسم، وحقول الرواتب المتخصصة
        pm = tables["payment_methods"]
        for m in session.query(PaymentMethod).all():
            if m.account_id in old_map:
                new_acc = by_code.get(payment_method_new_code(m.name))
                if new_acc:
                    report["payment_methods"].append((m.name, pairs[m.account_id][0].code, new_acc.code))
                    m.account_id = new_acc.id
        session.flush()
        for t, c in cols:
            new_code = PAYROLL_FIELD_NEW_CODES.get(c.name)
            if not new_code or not salary_old_ids:
                continue
            tgt = by_code.get(new_code)
            if not tgt:
                continue
            res = session.execute(update(t).where(c.in_(list(salary_old_ids))).values({c.name: tgt.id}))
            if res.rowcount:
                report["payroll_overrides"] += res.rowcount
                report["relinked"][f"{t.name}.{c.name}"] = report["relinked"].get(f"{t.name}.{c.name}", 0) + res.rowcount

        # ---- 4) إعادة الربط العامة
        for t, c in cols:
            key = f"{t.name}.{c.name}"
            if (t.name, c.name) == ("accounts", "parent_id"):
                for old_id, grp_id in group_map.items():
                    res = session.execute(update(t).where(c == old_id, t.c.id != grp_id).values({c.name: grp_id}))
                    if res.rowcount:
                        for a in session.query(Account).filter(Account.parent_id == grp_id).all():
                            if a.code not in new_codes and a.id not in old_map:
                                report["reparented"].append((a.code, a.name, by_code[[k for k, v in by_code.items() if v.id == grp_id][0]].code))
                continue
            if (t.name, c.name) == ("journal_entry_lines", "account_id") and not move_history:
                continue
            for old_id, new_id in old_map.items():
                res = session.execute(update(t).where(c == old_id).values({c.name: new_id}))
                if res.rowcount:
                    report["relinked"][key] = report["relinked"].get(key, 0) + res.rowcount
                    if key == "journal_entry_lines.account_id":
                        report["journal_lines_moved"] += res.rowcount
        session.flush()

        # ---- 5) إعدادات لسه فاضية
        cfg = session.get(AccountingSettings, 1)
        if cfg:
            for field, code in NEW_SETTINGS_CODES.items():
                if getattr(cfg, field) is None and code in by_code:
                    setattr(cfg, field, by_code[code].id); report["settings_filled"].append(field)
        session.flush()

        # ---- 6) حذف القديم اللي مبقاش مستخدم (الأبناء الأول)
        if remove_old:
            remaining = dict(pairs)
            progress = True
            while remaining and progress:
                progress = False
                for old_id, (old, _nc) in list(remaining.items()):
                    refs, per = _count_refs(session, old_id, cols)
                    if refs == 0:
                        session.delete(session.get(Account, old_id)); session.flush()
                        report["removed"].append((old.code, old.name)); del remaining[old_id]; progress = True
            for old_id, (old, _nc) in remaining.items():
                refs, per = _count_refs(session, old_id, cols)
                report["kept"].append({"code": old.code, "name": old.name, "refs": refs, "detail": per})
        after = totals()
        if before != after:
            raise RuntimeError(f"إجمالي القيود اتغيّر أثناء التحويل ({before} -> {after}) - اتلغى التنفيذ")
        session.commit()
        report["totals"] = after
        return report
    except Exception:
        session.rollback()
        raise
    finally:
        session.close()
