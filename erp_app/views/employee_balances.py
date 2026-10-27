import datetime
import os
import flet as ft
from .. import services, theme

COLS = ['الكود', 'الموظف', 'تاريخ التعيين', 'سنوات الخدمة', 'الأجر الشهري', 'الإجازة المستحقة', 'المستخدمة', 'الرصيد (يوم)',
        'قيمة الرصيد', 'نهاية الخدمة (استقالة)', 'نهاية الخدمة (إنهاء)', 'إجمالي عند الاستقالة', 'إجمالي عند الإنهاء']


KEYS = ['code', 'name', 'hire', 'years', 'wage', 'earned', 'used', 'days', 'amount', 'eos_res', 'eos_term', 'tot_res', 'tot_term']


def _vals(r):
    return [r['code'], r['name'], r['hire_date'].strftime('%Y-%m-%d') if r['hire_date'] else '-', r['years'], r['monthly_wage'],
            r['leave_earned'], r['leave_used'], r['leave_days'], r['leave_amount'], r['eos_resignation'], r['eos_termination'],
            round(r['leave_amount'] + r['eos_resignation'], 2), round(r['leave_amount'] + r['eos_termination'], 2)]


def export_balances_xlsx(rows, as_of, path):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    wb = Workbook(); ws = wb.active; ws.title = 'أرصدة الموظفين'; ws.sheet_view.rightToLeft = True
    ws.append([f'أرصدة الموظفين (إجازات ونهاية خدمة) حتى {as_of:%Y-%m-%d}']); ws['A1'].font = Font(bold=True, size=13)
    ws.append(COLS)
    for c in ws[2]:
        c.font = Font(bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor='1F6F5C'); c.alignment = Alignment(horizontal='center', wrap_text=True)
    for r in rows:
        ws.append(_vals(r))
    n = len(rows) + 2
    ws.append(['', 'الإجمالي'] + [''] * 6 + [f'=SUM({chr(64 + i)}3:{chr(64 + i)}{n})' for i in range(9, 14)])
    for c in ws[n + 1]:
        c.font = Font(bold=True)
    for i, w in enumerate([10, 26, 14, 12, 14, 14, 12, 12, 14, 16, 16, 16, 16], 1):
        ws.column_dimensions[chr(64 + i)].width = w
    wb.save(path); return path


def build_employee_balances(page: ft.Page):
    """أرصدة الموظفين: رصيد الإجازة (مستحق/مستخدم/متبقي وقيمته) ومكافأة نهاية الخدمة
    لو استقال أو أُنهيت خدمته، لحد تاريخ تختاره. للقراءة بس - التسوية الفعلية من شاشة تسويات نهاية الخدمة."""
    as_of_f = theme.DateField('الأرصدة حتى تاريخ', page, width=160, value=datetime.date.today().strftime('%Y-%m-%d'))
    search_f = ft.TextField(label='بحث (كود / اسم)', width=220, prefix_icon=ft.Icons.SEARCH)
    inactive_cb = ft.Checkbox(label='يشمل غير النشطين', value=False)
    cards = ft.Row(wrap=True, spacing=12)
    from .. import report_engine as RE
    from ..report_viewer import ReportViewer, Mode
    M = RE.MONEY
    viewer = ReportViewer(page, 'employee_balances', 'أرصدة الموظفين', [
        RE.Col('code', 'الكود'), RE.Col('name', 'الموظف', width=2, bold=True), RE.Col('hire', 'تاريخ التعيين'), RE.Col('years', 'سنوات الخدمة', RE.QTY),
        RE.Col('wage', 'الأجر الشهري', M, 'sum'), RE.Col('earned', 'الإجازة المستحقة', RE.QTY, 'sum'), RE.Col('used', 'المستخدمة', RE.QTY, 'sum'), RE.Col('days', 'الرصيد (يوم)', RE.QTY, 'sum', bold=True),
        RE.Col('amount', 'قيمة الرصيد', M, 'sum'), RE.Col('eos_res', 'نهاية الخدمة (استقالة)', M, 'sum'), RE.Col('eos_term', 'نهاية الخدمة (إنهاء)', M, 'sum'),
        RE.Col('tot_res', 'إجمالي عند الاستقالة', M, 'sum', bold=True), RE.Col('tot_term', 'إجمالي عند الإنهاء', M, 'sum', bold=True)],
        show_export=False, chart_value='amount')
    state = {'rows': []}

    def snack(msg):
        page.overlay.append(ft.SnackBar(ft.Text(msg), open=True)); page.update()

    def current_date():
        return datetime.datetime.strptime(as_of_f.value, '%Y-%m-%d')

    def refresh(e=None):
        try:
            rows = services.employee_balances_report(current_date(), active_only=not inactive_cb.value, search=search_f.value)
        except Exception as ex:
            snack(str(ex)); return
        state['rows'] = rows
        viewer.set_rows([dict(zip(KEYS, _vals(r))) for r in rows])

        def card(label, value, sub=''):
            return theme.stat_card(label, value, theme.PRIMARY, sub)
        cards.controls = [card('عدد الموظفين', str(len(rows))),
                          card('إجمالي أيام الإجازة', f"{sum(r['leave_days'] for r in rows):,.2f}"),
                          card('قيمة أرصدة الإجازات', f"{sum(r['leave_amount'] for r in rows):,.2f}"),
                          card('مكافآت نهاية الخدمة (استقالة)', f"{sum(r['eos_resignation'] for r in rows):,.2f}"),
                          card('مكافآت نهاية الخدمة (إنهاء)', f"{sum(r['eos_termination'] for r in rows):,.2f}")]
        page.update()

    def export(e=None):
        try:
            out_dir = os.path.join(services.APP_DIR, 'exports'); os.makedirs(out_dir, exist_ok=True)
            d = current_date()
            path = export_balances_xlsx(state['rows'], d, os.path.join(out_dir, f'employee_balances_{d:%Y%m%d}.xlsx'))
            theme.open_or_download_file(page, path); snack('تم إنشاء الملف')
        except Exception as ex:
            snack(f'تعذر التصدير: {ex}')

    as_of_f.on_change = refresh; search_f.on_change = refresh; inactive_cb.on_change = refresh
    refresh()
    return ft.Column([
        ft.Row([theme.page_title('أرصدة الموظفين', 'رصيد الإجازات ومكافأة نهاية الخدمة المستحقة لكل موظف حتى تاريخ معيّن'),
                ft.Container(expand=True), ft.OutlinedButton('Excel', icon=ft.Icons.TABLE_VIEW, on_click=export)], spacing=8),
        ft.Container(height=4), cards, ft.Row([as_of_f.control, search_f, inactive_cb], wrap=True, spacing=10),
        ft.Container(height=2), viewer.control,
    ], expand=True, scroll=ft.ScrollMode.AUTO)
