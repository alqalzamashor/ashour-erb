import datetime
import os
import flet as ft
from .. import services, theme

COLS = ['الشهر', 'رقم الدورة', 'الكود', 'الموظف', 'الفرع', 'القسم', 'الصافي', 'الحالة', 'تاريخ الصرف', 'سند الصرف']


def _status(r):
    if r['paid']:
        return 'مدفوع'
    return 'غير مدفوع' if r['run_posted'] else 'غير مدفوع (الدورة غير مرحّلة)'


def _vals(r):
    return [f"{r['month']}/{r['year']}", r['run_number'], r['code'], r['name'], r['branch'], r['department'], r['net'], _status(r),
            r['paid_at'].strftime('%Y-%m-%d') if r['paid_at'] else '-', r['payment_number']]


def export_payroll_report_xlsx(data, path, title='تقرير الرواتب'):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    wb = Workbook(); ws = wb.active; ws.title = 'الرواتب'; ws.sheet_view.rightToLeft = True
    ws.append([title]); ws['A1'].font = Font(bold=True, size=13)
    ws.append(COLS)
    for c in ws[2]:
        c.font = Font(bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor='1F6F5C'); c.alignment = Alignment(horizontal='center', wrap_text=True)
    for r in data['rows']:
        ws.append(_vals(r))
    t = data['totals']
    ws.append(['', '', '', 'الإجمالي', '', '', t['total'], f"مدفوع {t['paid']:,.2f}", f"غير مدفوع {t['unpaid']:,.2f}", ''])
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)
    for i in range(1, len(COLS) + 1):
        ws.column_dimensions[ws.cell(row=2, column=i).column_letter].width = 18
    ws2 = wb.create_sheet('ملخص الشهور'); ws2.sheet_view.rightToLeft = True
    ws2.append(['الشهر', 'عدد القسائم', 'الإجمالي', 'المدفوع', 'غير المدفوع'])
    for m in data['months']:
        ws2.append([m['month'], m['count'], m['total'], m['paid'], m['unpaid']])
    wb.save(path); return path


def export_payroll_report_pdf(data, path, title='تقرير الرواتب', rtl=True, filters=None):
    from .. import print_screen
    t = data['totals']
    rows = [[_fmt_cell(v) for v in _vals(r)] for r in data['rows']]
    rows.append(['', '', '', 'الإجمالي', '', '', f"{t['total']:,.2f}", f"مدفوع {t['paid']:,.2f}", f"غير مدفوع {t['unpaid']:,.2f}", ''])
    tables = [{'columns': list(COLS), 'rows': rows, 'caption': ''}]
    if data.get('months'):
        tables.append({'columns': ['الشهر', 'عدد القسائم', 'الإجمالي', 'المدفوع', 'غير المدفوع'],
                       'rows': [[str(m['month']), str(m['count']), f"{m['total']:,.2f}", f"{m['paid']:,.2f}", f"{m['unpaid']:,.2f}"] for m in data['months']], 'caption': 'ملخص الشهور'})
    return print_screen.build_pdf({'filters': filters or [], 'stats': [], 'tables': tables}, title, path, rtl=rtl)


def _fmt_cell(v):
    return f'{v:,.2f}' if isinstance(v, float) else str(v)


def build_payroll_report(page: ft.Page):
    """تقرير الرواتب بالشهور: مدفوع / غير مدفوع لكل الموظفين، وتقدر تصرف غير المدفوع مباشرة من هنا."""
    today = datetime.date.today()
    year_f = ft.TextField(label='السنة', width=100, value=str(today.year), keyboard_type=ft.KeyboardType.NUMBER)
    month_dd = ft.Dropdown(label='الشهر', width=140, value='0', options=[ft.dropdown.Option('0', 'كل الشهور')] + [ft.dropdown.Option(str(m), str(m)) for m in range(1, 13)])
    status_dd = ft.Dropdown(label='حالة الصرف', width=160, value='all', options=[ft.dropdown.Option('all', 'الكل'), ft.dropdown.Option('unpaid', 'غير مدفوع'), ft.dropdown.Option('paid', 'مدفوع')])
    branch_dd = ft.Dropdown(label='الفرع', width=170, value='0', options=[ft.dropdown.Option('0', 'كل الفروع')] + [ft.dropdown.Option(str(b.id), b.name) for b in services.list_branches()])
    cc_dd = ft.Dropdown(label='القسم', width=170, value='0', options=[ft.dropdown.Option('0', 'كل الأقسام')] + [ft.dropdown.Option(str(c.id), c.name) for c in services.list_cost_centers()])
    search_f = ft.TextField(label='بحث (كود / اسم)', width=200, prefix_icon=ft.Icons.SEARCH)
    cards = ft.Row(wrap=True, spacing=12)
    months_host = ft.Container()
    table_host = ft.Container()
    sel_text = ft.Text('', weight=ft.FontWeight.BOLD)
    pay_btn = ft.FilledButton('صرف المختار', icon=ft.Icons.PAYMENTS_OUTLINED, bgcolor=theme.PRIMARY)
    msg = ft.Text('', color=theme.DANGER, visible=False)
    state = {'data': None, 'sel': {}}      # sel: payslip_id -> net

    def snack(m):
        page.overlay.append(ft.SnackBar(ft.Text(m), open=True)); page.update()

    def val(dd):
        return None if dd.value in (None, '', '0') else dd.value

    def update_sel():
        n = len(state['sel']); tot = sum(state['sel'].values())
        sel_text.value = f'المختار: {n} موظف بمبلغ {tot:,.2f}' if n else 'اختار القسائم غير المدفوعة اللي عايز تصرفها'
        pay_btn.disabled = n == 0

    def toggle(pid, net, on):
        if on:
            state['sel'][pid] = net
        else:
            state['sel'].pop(pid, None)
        update_sel(); page.update()

    def payable(r):
        return (not r['paid']) and r['run_posted']

    def preview_voucher(pid):
        try:
            theme.open_or_download_file(page, services.generate_payment_pdf_file(pid))
        except Exception as ex:
            snack(f'تعذرت معاينة السند: {ex}')

    def render_table():
        data = state['data']
        rows = []
        for r in data['rows']:
            v = _vals(r)
            if payable(r):
                first = ft.Checkbox(value=r['payslip_id'] in state['sel'], on_change=lambda e, pid=r['payslip_id'], net=r['net']: toggle(pid, net, e.control.value))
            else:
                first = ft.Icon(ft.Icons.CHECK_CIRCLE if r['paid'] else ft.Icons.LOCK_OUTLINE, size=18, color=theme.PRIMARY if r['paid'] else theme.TEXT_MUTED)
            color = theme.PRIMARY if r['paid'] else theme.DANGER
            cells = [ft.DataCell(first)] + [ft.DataCell(ft.Text(f'{x:,.2f}' if isinstance(x, float) else str(x),
                                                                   weight=ft.FontWeight.BOLD if i in (3, 7) else None, color=color if i == 7 else None))
                                            for i, x in enumerate(v)]
            if r['paid'] and r.get('payment_id'):
                cells.append(ft.DataCell(ft.IconButton(ft.Icons.VISIBILITY_OUTLINED, tooltip='معاينة سند الصرف', icon_size=20, icon_color=theme.PRIMARY,
                                                       on_click=lambda e, pid=r['payment_id']: preview_voucher(pid))))
            else:
                cells.append(ft.DataCell(ft.Text('')))
            rows.append(ft.DataRow(cells=cells))
        t = data['totals']
        rows.append(ft.DataRow(cells=[ft.DataCell(ft.Text(''))] + [ft.DataCell(ft.Text(x, weight=ft.FontWeight.BOLD, color=theme.PRIMARY)) for x in
                                                                    ['', '', '', 'الإجمالي', '', '', f"{t['total']:,.2f}", '', '', '', '']]))
        table_host.content = theme.h_scroll(ft.DataTable(columns=[ft.DataColumn(ft.Text(''))] + [ft.DataColumn(ft.Text(c)) for c in COLS] + [ft.DataColumn(ft.Text(''))], rows=rows), vscroll=True)

    adv_sw = ft.Switch(label='عرض متقدم (فرز وفلترة وتجميع ورسم)', value=False)
    adv = {'viewer': None}
    adv_host = ft.Container(visible=False, expand=True)

    def render_adv():
        from .. import report_engine as RE
        from ..report_viewer import ReportViewer, Mode
        if adv['viewer'] is None:
            adv['viewer'] = ReportViewer(page, 'payroll_adv', 'تقرير الرواتب', [
                RE.Col('month', 'الشهر'), RE.Col('run_number', 'رقم الدورة'), RE.Col('code', 'الكود'), RE.Col('name', 'الموظف', width=2, bold=True), RE.Col('branch', 'الفرع'),
                RE.Col('department', 'القسم'), RE.Col('net', 'الصافي', RE.MONEY, 'sum'), RE.Col('status', 'الحالة'), RE.Col('paid_at', 'تاريخ الصرف', RE.DATE), RE.Col('payment_number', 'سند الصرف')],
                modes=[Mode('by_month', 'حسب الشهر', ['month']), Mode('by_status', 'حسب الحالة', ['status']), Mode('by_branch', 'حسب الفرع', ['branch']),
                       Mode('by_dept', 'حسب القسم', ['department']), Mode('month_status', 'شهر ← حالة', ['month', 'status'])],
                on_row_click=lambda r: preview_voucher(r['payment_id']) if r.get('payment_id') else None, chart_value='net')
            adv_host.content = adv['viewer'].control
        out = []
        for r in state['data']['rows']:
            out.append(dict(r, month=f"{r['year']}-{int(r['month']):02d}", status=_status(r), paid_at=r['paid_at'], payment_number=r.get('payment_number') or ''))
        adv['viewer'].set_rows(out)

    def on_adv(e=None):
        on = bool(adv_sw.value)
        adv_host.visible = on; table_wrap.visible = not on
        if on and state['data']:
            render_adv()
        page.update()
    adv_sw.on_change = on_adv

    def refresh(e=None, keep_sel=False):
        try:
            data = services.payroll_payments_report(year_f.value or today.year, val(month_dd), status_dd.value, val(branch_dd), val(cc_dd), search_f.value)
        except Exception as ex:
            msg.value = str(ex); msg.visible = True; table_host.content = None; page.update(); return
        msg.visible = False; state['data'] = data
        ids = {r['payslip_id'] for r in data['rows'] if payable(r)}
        state['sel'] = {k: v for k, v in state['sel'].items() if k in ids} if keep_sel else {}
        t = data['totals']

        def card(label, value, color=theme.PRIMARY):
            return theme.card(ft.Column([ft.Text(label, size=12, color=theme.TEXT_MUTED), ft.Text(value, size=20, weight=ft.FontWeight.BOLD, color=color)], spacing=0, tight=True), padding=14)
        cards.controls = [card('عدد القسائم', str(t['count'])), card('إجمالي صافي الرواتب', f"{t['total']:,.2f}"),
                          card('المدفوع', f"{t['paid']:,.2f}"), card('غير المدفوع', f"{t['unpaid']:,.2f}", theme.DANGER)]
        chips = [theme.card(ft.Column([ft.Text(f"شهر {m['month']}  •  {m['count']} قسيمة", size=12, weight=ft.FontWeight.BOLD),
                                       ft.Text(f"الإجمالي {m['total']:,.2f}", size=11, color=theme.TEXT_MUTED),
                                       ft.Text(f"مدفوع {m['paid']:,.2f}", size=11, color=theme.PRIMARY),
                                       ft.Text(f"غير مدفوع {m['unpaid']:,.2f}", size=11, color=theme.DANGER if m['unpaid'] else theme.TEXT_MUTED)], spacing=0, tight=True), padding=10)
                 for m in data['months']]
        months_host.content = ft.Row(chips, wrap=True, spacing=8) if chips else None
        render_table(); update_sel()
        if adv_sw.value:
            render_adv()
        page.update()

    def select_all_unpaid(e=None):
        state['sel'] = {r['payslip_id']: r['net'] for r in state['data']['rows'] if payable(r)} if state['data'] else {}
        render_table(); update_sel(); page.update()

    def clear_sel(e=None):
        state['sel'] = {}; render_table(); update_sel(); page.update()

    def open_pay_dialog(e=None):
        if not state['sel']:
            return
        methods = [m for m in services.list_payment_methods(active_only=True) if not m.is_credit]
        pm_dd = ft.Dropdown(label='طريقة الدفع *', width=280, options=[ft.dropdown.Option(str(m.id), m.name) for m in methods], value=str(methods[0].id) if methods else None)
        pdate_f = theme.DateField('تاريخ الصرف', page, width=160, value=today.strftime('%Y-%m-%d'))
        ref_f = ft.TextField(label='مرجع (اختياري)', width=200)
        perr = ft.Text('', color=theme.DANGER, visible=False)
        if not methods:
            perr.value = 'مفيش طريقة دفع فورية متاحة (غير آجلة) - ضيف واحدة من شاشة "طرق الدفع" الأول'; perr.visible = True

        def confirm(ev):
            try:
                if not pm_dd.value:
                    raise ValueError('اختر طريقة الدفع')
                d = datetime.datetime.strptime(pdate_f.value, '%Y-%m-%d')
                res = services.pay_payslips_multi(list(state['sel']), int(pm_dd.value), payment_date=d, reference=ref_f.value or '')
                dlg.open = False; page.update()
                refresh()
                snack(f"تم صرف {res['employee_count']} موظف بمبلغ {res['amount']:,.2f} - سندات: {'، '.join(res['payments'])}")
            except Exception as ex:
                perr.value = str(ex); perr.visible = True; page.update()

        dlg = ft.AlertDialog(modal=True, title=ft.Text('صرف الرواتب غير المدفوعة'),
                             content=ft.Column([ft.Text(f"{len(state['sel'])} موظف بمبلغ {sum(state['sel'].values()):,.2f}", weight=ft.FontWeight.BOLD, size=16),
                                                ft.Text('هيتعمل سند صرف وقيد (مدين الرواتب المستحقة / دائن حساب طريقة الدفع) لكل شهر مختار.', color=theme.TEXT_MUTED, size=12),
                                                pm_dd, pdate_f.control, ref_f, perr], tight=True, spacing=10, width=420),
                             actions=[ft.TextButton('تراجع', on_click=lambda ev: (setattr(dlg, 'open', False), page.update())),
                                      ft.FilledButton('تأكيد الصرف', icon=ft.Icons.PAYMENTS_OUTLINED, bgcolor=theme.PRIMARY, on_click=confirm)])
        page.overlay.append(dlg); dlg.open = True; page.update()

    def export(e=None):
        if not state['data']:
            snack('مفيش بيانات لتصديرها'); return
        try:
            out_dir = os.path.join(services.APP_DIR, 'exports'); os.makedirs(out_dir, exist_ok=True)
            path = export_payroll_report_xlsx(state['data'], os.path.join(out_dir, f"payroll_report_{year_f.value}_{today:%Y%m%d}.xlsx"), f'تقرير الرواتب {year_f.value}')
            theme.open_or_download_file(page, path); snack('تم إنشاء الملف')
        except Exception as ex:
            snack(f'تعذر التصدير: {ex}')

    def export_pdf(e=None):
        if not state['data']:
            snack('مفيش بيانات لتصديرها'); return
        try:
            out_dir = os.path.join(services.APP_DIR, 'exports'); os.makedirs(out_dir, exist_ok=True)
            from .. import print_screen
            flt = print_screen.extract(ft.Row([year_f, month_dd, status_dd, branch_dd, cc_dd, search_f]))['filters']
            path = export_payroll_report_pdf(state['data'], os.path.join(out_dir, f"payroll_report_{year_f.value}_{today:%Y%m%d}.pdf"), f'تقرير الرواتب {year_f.value}', filters=flt)
            theme.open_or_download_file(page, path); snack('تم إنشاء الملف')
        except Exception as ex:
            snack(f'تعذر التصدير: {ex}')

    for f in (month_dd, status_dd, branch_dd, cc_dd):
        f.on_select = refresh
    year_f.on_change = refresh; search_f.on_change = refresh
    pay_btn.on_click = open_pay_dialog
    table_wrap = ft.Column([ft.Row([sel_text, ft.Container(expand=True), ft.OutlinedButton('تحديد كل غير المدفوع', on_click=select_all_unpaid),
                                    ft.TextButton('إلغاء التحديد', on_click=clear_sel), pay_btn], spacing=8),
                            theme.card(table_host, expand=True)], expand=True)
    refresh()
    return ft.Column([
        ft.Row([theme.page_title('تقرير الرواتب', 'المدفوع وغير المدفوع لكل الموظفين بالشهور - وصرف غير المدفوع مباشرة من هنا'),
                ft.Container(expand=True), ft.OutlinedButton('PDF', icon=ft.Icons.PICTURE_AS_PDF, on_click=export_pdf), ft.OutlinedButton('Excel', icon=ft.Icons.TABLE_VIEW, on_click=export)], spacing=8),
        ft.Container(height=4), cards, ft.Row([year_f, month_dd, status_dd, branch_dd, cc_dd, search_f], wrap=True, spacing=10), msg,
        months_host, adv_sw,
        table_wrap, adv_host,
    ], expand=True)
