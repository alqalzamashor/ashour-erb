"""القوائم المالية وفق IFRS: شجرة (رئيسي ← فرعي ← فرعي)، أعمدة لكل فرع + موحد، مقارنة بالفترة السابقة، تصدير Excel/PDF.

بتُستخدم في شاشتين: «القوائم المالية الموحدة والفروع» (أعمدة الفروع شغالة افتراضيًا) وتبويب «القوائم المالية (IFRS)» في التقارير المالية.
"""
import datetime
import os

import flet as ft

from .. import ifrs, services, theme

STATEMENTS = [
    ('position', 'قائمة المركز المالي'),
    ('pl', 'قائمة الربح أو الخسارة والدخل الشامل'),
    ('cf', 'قائمة التدفقات النقدية'),
    ('equity', 'قائمة التغيرات في حقوق الملكية'),
    ('cc', 'ربحية مراكز التكلفة'),
]
TITLES = dict(STATEMENTS)
LEVELS = [('1', 'مستوى 1'), ('2', 'مستوى 2'), ('3', 'مستوى 3'), ('4', 'مستوى 4'), ('5', 'مستوى 5'), ('all', 'كل المستويات')]


def _row(rid, parent, level, kind, name, values, code='', has_children=False):
    return ifrs._row(rid, parent, level, kind, name, values, code, has_children)


def _cf_result(cf):
    def sec(sid, title, items, net):
        rows = [_row(sid, None, 0, 'section', title, None, has_children=True)]
        for i, it in enumerate(items):
            rows.append(_row(f'{sid}:{i}', sid, 1, 'account', it['description'], [it['amount']]))
        rows.append(_row(sid + ':net', None, 0, 'subtotal', f'صافي التدفق النقدي من {title}', [net]))
        return rows
    rows = [_row('open', None, 0, 'subtotal', 'النقدية وما في حكمها - أول الفترة', [cf['opening_cash']])]
    rows += sec('op', 'الأنشطة التشغيلية', cf['operating_in'] + cf['operating_out'], cf['net_operating'])
    rows += sec('inv', 'الأنشطة الاستثمارية', cf['investing_in'] + cf['investing_out'], cf['net_investing'])
    rows += sec('fin', 'الأنشطة التمويلية', cf['financing_in'] + cf['financing_out'], cf['net_financing'])
    rows.append(_row('net', None, 0, 'total', 'صافي التغير في النقدية وما في حكمها', [cf['net']]))
    rows.append(_row('close', None, 0, 'total', 'النقدية وما في حكمها - آخر الفترة', [cf['closing_cash']]))
    if abs(cf.get('unexplained') or 0) > 0.005:
        rows.append(_row('unx', None, 0, 'warn', 'فرق غير مفسَّر (راجع تصنيف التدفق النقدي للحسابات)', [cf['unexplained']]))
    return {'rows': rows, 'columns': [{'label': 'المبلغ', 'kind': 'total'}]}


def _equity_result(eq):
    rows = [_row('open', None, 0, 'subtotal', 'حقوق الملكية - أول الفترة', [eq['opening_equity']]),
            _row('mv', None, 0, 'section', 'الحركة خلال الفترة', None, has_children=True)]
    for i, m in enumerate(eq['movements']):
        rows.append(_row(f'mv:{i}', 'mv', 1, 'account', m['name'], [m['amount']], m.get('code', '')))
    rows.append(_row('np', None, 0, 'subtotal', 'صافي ربح / (خسارة) الفترة', [eq['net_profit']]))
    rows.append(_row('close', None, 0, 'total', 'حقوق الملكية - آخر الفترة', [eq['closing_equity']]))
    return {'rows': rows, 'columns': [{'label': 'المبلغ', 'kind': 'total'}]}


def _cc_result(items):
    rows = [_row(f'cc:{i}', None, 0, 'account', c['name'], [c['income'], c['expense'], c['profit']], c['code']) for i, c in enumerate(items)]
    if items:
        rows.append(_row('cc:total', None, 0, 'total', 'الإجمالي', [sum(c['income'] for c in items), sum(c['expense'] for c in items), sum(c['profit'] for c in items)]))
    return {'rows': rows, 'columns': [{'label': 'الإيرادات', 'kind': 'total'}, {'label': 'المصروفات', 'kind': 'total'}, {'label': 'صافي الربح', 'kind': 'total'}]}


def build_ifrs_statements(page: ft.Page, by_branch=False, title='القوائم المالية (IFRS)', subtitle='', with_title=True, statements=None):
    statements = statements or [k for k, _ in STATEMENTS]
    branches_all = services.list_branches(True)
    allowed = services.user_branch_ids()
    branches = [b for b in branches_all if not allowed or b.id in allowed]
    centers = services.list_cost_centers(True)
    today = datetime.date.today()

    stmt = ft.Dropdown(label='القائمة', width=300, value=statements[0], options=[ft.dropdown.Option(k, TITLES[k]) for k in statements])
    rng = theme.DateRangeField('الفترة / كما في', page, width=290, start=f'{today.year}-01-01', end=today.isoformat())
    branch = ft.Dropdown(label='الفرع', width=210, value='', options=[ft.dropdown.Option('', 'كل الفروع - موحد')] + [ft.dropdown.Option(str(b.id), f'{b.code} - {b.name}') for b in branches])
    center = ft.Dropdown(label='مركز التكلفة', width=210, value='', options=[ft.dropdown.Option('', 'كل مراكز التكلفة')] + [ft.dropdown.Option(str(c.id), f'{c.code} - {c.name}') for c in centers])
    sw_branch = ft.Switch(label='عمود لكل فرع + الإجمالي الموحد', value=by_branch)
    sw_cmp = ft.Switch(label='مقارنة بالفترة السابقة', value=False)
    sw_zero = ft.Switch(label='إظهار الأرصدة الصفرية', value=False)
    level = ft.Dropdown(label='مستوى التفصيل', width=150, value='2', options=[ft.dropdown.Option(k, v) for k, v in LEVELS])
    err = ft.Text('', color=theme.DANGER, visible=False)
    head = ft.Column([], spacing=2, horizontal_alignment=ft.CrossAxisAlignment.CENTER)
    holder = ft.Column([], spacing=0)
    state = {'res': None, 'collapsed': set(), 'expanded': set()}

    def d(v):
        v = (v or '').strip()
        return datetime.date.fromisoformat(v) if v else None

    def lim():
        return None if level.value == 'all' else int(level.value or 2)

    def compute():
        k = stmt.value
        df, dt = d(rng.start.value), d(rng.end.value)
        bid = int(branch.value) if branch.value else None
        cid = int(center.value) if center.value else None
        if k in ('position', 'pl'):
            res = ifrs.build(k, date_from=df, date_to=dt, branches=[(b.id, b.name) for b in branches],
                             by_branch=sw_branch.value and not bid, compare=sw_cmp.value, cost_center_id=cid)
            if bid:        # فرع محدد: نفس القائمة لكنها بحركات الفرع ده بس
                res = _single_branch(k, df, dt, bid, cid, sw_cmp.value)
        elif k == 'cf':
            res = _cf_result(services.cash_flow_statement(df, dt, bid, cid))
        elif k == 'equity':
            res = _equity_result(services.statement_of_changes_in_equity(df, dt))
        else:
            res = _cc_result(services.cost_center_statement(df, dt, bid))
        return res

    def _single_branch(k, df, dt, bid, cid, compare):
        chart = ifrs.load_chart()
        cur = ifrs.movements(df if k == 'pl' else None, dt, cid)
        vals = [ifrs._col_values(cur, bid)]
        cols = [{'label': 'المبلغ' if not compare else 'الفترة الحالية', 'kind': 'total'}]
        if compare and dt:
            pf, pt = ifrs._prior_range(k, df, dt)
            vals.append(ifrs._col_values(ifrs.movements(pf if k == 'pl' else None, pt, cid), bid))
            cols.append({'label': 'الفترة السابقة', 'kind': 'prior'})
        rows = (ifrs.financial_position if k == 'position' else ifrs.profit_or_loss)(chart, vals)
        if compare and dt:
            for r in rows:
                if r['values'] is not None:
                    a, b = r['values'][0], r['values'][1]
                    r['values'] += [a - b, ((a - b) / abs(b) * 100.0) if abs(b) > 0.004 else None]
            cols += [{'label': 'التغير', 'kind': 'delta'}, {'label': '%', 'kind': 'pct'}]
        return {'rows': rows, 'columns': cols, 'chart': chart}

    def subtitle_text():
        df, dt = rng.start.value, rng.end.value
        if stmt.value == 'position':
            return f'كما في {dt}'
        return f'للفترة من {df} إلى {dt}'

    def scope_text():
        parts = []
        if branch.value:
            parts.append('الفرع: ' + next((f'{b.code} - {b.name}' for b in branches if str(b.id) == branch.value), ''))
        elif sw_branch.value and stmt.value in ('position', 'pl'):
            parts.append('موحّد - تفصيل حسب الفرع')
        else:
            parts.append('كل الفروع (موحّد)')
        if center.value:
            parts.append('مركز التكلفة: ' + next((f'{c.code} - {c.name}' for c in centers if str(c.id) == center.value), ''))
        return '  |  '.join(parts)

    def is_open(r):
        if r['id'] in state['collapsed']:
            return False
        if r['id'] in state['expanded']:
            return True
        l = lim()
        return l is None or r['level'] < l

    def toggle(r):
        if is_open(r):
            state['expanded'].discard(r['id']); state['collapsed'].add(r['id'])
        else:
            state['collapsed'].discard(r['id']); state['expanded'].add(r['id'])
        render()
        page.update()

    def zero(r):
        return r['values'] is not None and all(abs(v or 0) < 0.005 for v in r['values'])

    def render():
        res = state['res']
        if not res:
            return
        cols = res['columns']
        rows = ifrs.visible_rows(res['rows'], lim(), state['collapsed'], state['expanded'])
        if not sw_zero.value:
            rows = [r for r in rows if not (r['kind'] in ('account', 'group') and zero(r))]
        co = None
        try:
            co = services.get_company_settings()
        except Exception:       # noqa: BLE001
            pass
        head.controls = [ft.Text((co.name if co and getattr(co, 'name', '') else ''), size=16, weight=ft.FontWeight.BOLD),
                         ft.Text(TITLES[stmt.value], size=15, weight=ft.FontWeight.BOLD, color=theme.PRIMARY),
                         ft.Text(subtitle_text(), size=12, color=theme.TEXT_MUTED),
                         ft.Text(scope_text(), size=11, color=theme.TEXT_MUTED)]
        NUM_W = 150 if len(cols) <= 4 else 130
        wide = len(cols) > 5           # أعمدة كتير (فروع كتير): عرض ثابت + تمرير أفقي بالعجلة؛ غير كده بيملا العرض
        hdr = ft.Container(
            ft.Row([ft.Container(ft.Text('البند', weight=ft.FontWeight.BOLD, size=13, color='#FFFFFF'), expand=True)] +
                   [ft.Container(ft.Text(c['label'], weight=ft.FontWeight.BOLD, size=13, color='#FFFFFF', text_align=ft.TextAlign.CENTER), width=NUM_W, alignment=ft.Alignment(0, 0)) for c in cols],
                   spacing=0),
            bgcolor=theme.PRIMARY, padding=ft.Padding(left=12, right=12, top=10, bottom=10), border_radius=ft.BorderRadius(8, 8, 0, 0))
        out = [hdr]
        for r in rows:
            k = r['kind']
            if k == 'spacer':
                out.append(ft.Container(height=14))
                continue
            bold = ft.FontWeight.BOLD if k in ('title', 'section', 'subtotal', 'total', 'warn') or r['has_children'] else None
            color = theme.DANGER if k == 'warn' else (theme.PRIMARY if k == 'title' else None)
            label = (f"{r['code']}  " if r.get('code') and k in ('account', 'group') else '') + r['name']
            collapsible = r['has_children'] and k in ('section', 'group')
            indent = 20 * r['level'] if k in ('account', 'group') else 0
            icon = ft.Icon(ft.Icons.KEYBOARD_ARROW_DOWN if is_open(r) else ft.Icons.KEYBOARD_ARROW_LEFT, size=18, color=theme.PRIMARY) if collapsible else ft.Container(width=18)
            name = ft.Container(ft.Row([ft.Container(width=indent), icon, ft.Text(label, weight=bold, color=color, size=14 if k in ('title', 'total') else 13, expand=True)], spacing=4), expand=True)
            cells = []
            for i, c in enumerate(cols):
                v = None if r['values'] is None or i >= len(r['values']) else r['values'][i]
                neg = v is not None and c['kind'] != 'pct' and v < -0.004
                cells.append(ft.Container(ft.Text('\u200e' + ifrs.fmt(v, c['kind']) + '\u200e', weight=bold, color=theme.DANGER if (neg and k != 'warn') else color,
                                                  size=13, text_align=ft.TextAlign.RIGHT), width=NUM_W, alignment=ft.Alignment(1, 0)))
            bg = {'title': theme.PRIMARY_LIGHT, 'section': theme.MAIN_ROW_BG, 'total': theme.PRIMARY_LIGHT, 'warn': '#FEE2E2'}.get(k)
            border = ft.Border(top=ft.BorderSide(1, theme.PRIMARY)) if k in ('subtotal', 'total') else ft.Border(bottom=ft.BorderSide(1, theme.BORDER_LIGHT))
            out.append(ft.Container(ft.Row([name] + cells, spacing=0), bgcolor=bg, border=border, ink=collapsible,
                                    on_click=(lambda e, row=r: toggle(row)) if collapsible else None,
                                    padding=ft.Padding(left=12, right=12, top=7 if k not in ('total',) else 9, bottom=7 if k not in ('total',) else 9)))
        grid = ft.Column(out, spacing=0, width=(380 + NUM_W * len(cols)) if wide else None, horizontal_alignment=ft.CrossAxisAlignment.STRETCH)
        holder.controls = [ft.Row([grid], scroll=ft.ScrollMode.AUTO) if wide else grid] if len(out) > 1 else [ft.Text('لا توجد بيانات في الفترة المحددة', color=theme.TEXT_MUTED)]

    def refresh(e=None):
        try:
            state['res'] = compute()
            state['collapsed'], state['expanded'] = set(), set()
            err.visible = False
            render()
        except Exception as ex:       # noqa: BLE001
            err.value = f'تعذر تحميل القائمة: {ex}'; err.visible = True
        if e is not None:
            page.update()

    def rerender(e=None):
        render()
        page.update()

    def set_all(open_):
        res = state['res']
        if not res:
            return
        state['collapsed'], state['expanded'] = set(), set()
        for r in res['rows']:
            if r['has_children']:
                (state['expanded'] if open_ else state['collapsed']).add(r['id'])
        render(); page.update()

    def apply_ui(e=None):
        on_pl = stmt.value in ('position', 'pl')
        sw_branch.visible = sw_cmp.visible = on_pl and True
        level.visible = on_pl or stmt.value in ('cf', 'equity')
        branch.visible = stmt.value != 'equity'
        center.visible = stmt.value != 'equity'
        refresh(e)

    # ------------------------------------------------------------- تصدير
    def export_rows():
        res = state['res']
        rows = ifrs.visible_rows(res['rows'], lim(), state['collapsed'], state['expanded'])
        out = []
        for r in rows:
            if r['kind'] == 'spacer' or (not sw_zero.value and r['kind'] in ('account', 'group') and zero(r)):
                continue
            vals = None if r['values'] is None else [ifrs.fmt(v, c['kind']) for v, c in zip(r['values'], res['columns'])]
            name = (f"{r['code']} - " if r.get('code') and r['kind'] in ('account', 'group') else '') + r['name']
            out.append((r['kind'], r['level'] if r['kind'] in ('account', 'group') else 0, name, vals))
        return out

    def do_export(kind):
        try:
            if not state['res']:
                refresh()
            exp = export_rows()
            cols = state['res']['columns']
            folder = services.downloads_dir()
            os.makedirs(folder, exist_ok=True)
            stamp = datetime.datetime.now().strftime('%Y%m%d_%H%M%S')
            title_ = TITLES[stmt.value]
            sub = subtitle_text() + '  |  ' + scope_text()
            if kind == 'xlsx':
                path = os.path.join(folder, f'ifrs_{stmt.value}_{stamp}.xlsx')
                _xlsx(path, title_, sub, cols, exp)
            else:
                from .. import print_screen
                path = os.path.join(folder, f'ifrs_{stmt.value}_{stamp}.pdf')
                body = [[('    ' * lv) + nm] + (v or ['' for _ in cols]) for _, lv, nm, v in exp]
                print_screen.build_pdf({'filters': [], 'stats': [], 'tables': [{'columns': ['البند'] + [c['label'] for c in cols], 'rows': body, 'caption': ''}]},
                                       title_, path, subtitle=sub)
            theme.open_or_download_file(page, path)
            page.overlay.append(ft.SnackBar(ft.Text(f'تم إنشاء الملف: {path}'), open=True)); page.update()
        except Exception as ex:       # noqa: BLE001
            page.overlay.append(ft.SnackBar(ft.Text(f'تعذر التصدير: {ex}'), open=True)); page.update()

    rng.start.on_change = refresh
    rng.end.on_change = refresh
    for c in (stmt,):
        c.on_select = apply_ui
    for c in (branch, center):
        c.on_select = refresh
    for c in (sw_branch, sw_cmp):
        c.on_change = refresh
    sw_zero.on_change = rerender
    level.on_select = lambda e: (state['collapsed'].clear(), state['expanded'].clear(), rerender())

    apply_ui()
    controls = []
    if with_title:
        controls.append(theme.page_title(title, subtitle or 'قوائم مالية بصيغة IFRS: الحساب الرئيسي ثم الفرعي ثم الفرعي، مع أعمدة الفروع والإجمالي الموحد'))
    controls.append(theme.card(ft.Column([
        ft.Row([stmt, rng.control, branch, center], wrap=True, spacing=10),
        ft.Row([sw_branch, sw_cmp, sw_zero], wrap=True, spacing=16),
        ft.Row([level,
                ft.OutlinedButton('توسيع الكل', icon=ft.Icons.UNFOLD_MORE, on_click=lambda e: set_all(True)),
                ft.OutlinedButton('طي الكل', icon=ft.Icons.UNFOLD_LESS, on_click=lambda e: set_all(False)),
                ft.Container(expand=True),
                ft.IconButton(ft.Icons.REFRESH, tooltip='تحديث', on_click=refresh),
                ft.ElevatedButton('Excel', icon=ft.Icons.TABLE_VIEW, on_click=lambda e: do_export('xlsx')),
                ft.ElevatedButton('PDF', icon=ft.Icons.PICTURE_AS_PDF, on_click=lambda e: do_export('pdf'))], wrap=True, spacing=8),
        err], spacing=10)))
    controls.append(theme.card(ft.Column([head, ft.Divider(height=8), holder], spacing=6)))
    return ft.Column(controls, expand=True, scroll=ft.ScrollMode.AUTO, spacing=12)


def _xlsx(path, title, sub, cols, exp):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    wb = Workbook(); ws = wb.active; ws.title = title[:31]; ws.sheet_view.rightToLeft = True
    try:
        co = services.get_company_settings()
        co_name = co.name if co and getattr(co, 'name', '') else ''
    except Exception:       # noqa: BLE001
        co_name = ''
    n = len(cols) + 1
    for i, (txt, sz) in enumerate(((co_name, 16), (title, 14), (sub, 10)), 1):
        ws.merge_cells(start_row=i, start_column=1, end_row=i, end_column=n)
        c = ws.cell(i, 1, txt); c.font = Font(size=sz, bold=i < 3); c.alignment = Alignment(horizontal='center')
    hdr = 5
    green = PatternFill('solid', fgColor='1A7A5E'); light = PatternFill('solid', fgColor='EAF5F0')
    for j, h in enumerate(['البند'] + [c['label'] for c in cols], 1):
        c = ws.cell(hdr, j, h); c.font = Font(bold=True, color='FFFFFF'); c.fill = green; c.alignment = Alignment(horizontal='center')
    thin, dbl = Side(style='thin'), Side(style='double')
    for i, (kind, lvl, name, vals) in enumerate(exp, hdr + 1):
        bold = kind in ('title', 'section', 'subtotal', 'total', 'warn')
        a = ws.cell(i, 1, name); a.font = Font(bold=bold); a.alignment = Alignment(horizontal='right', indent=lvl)
        for j, v in enumerate(vals or [], 2):
            c = ws.cell(i, j, v); c.font = Font(bold=bold); c.alignment = Alignment(horizontal='left')
        for j in range(1, n + 1):
            cell = ws.cell(i, j)
            if kind in ('title', 'section'):
                cell.fill = light
            if kind == 'subtotal':
                cell.border = Border(top=thin)
            if kind == 'total':
                cell.fill = light; cell.border = Border(top=thin, bottom=dbl)
    ws.column_dimensions['A'].width = 58
    for j in range(2, n + 1):
        ws.column_dimensions[chr(64 + j)].width = 20
    ws.freeze_panes = ws.cell(hdr + 1, 2)
    wb.save(path)
