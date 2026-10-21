"""الحضور والانصراف وربط جهاز البصمة بالرواتب.

مصادر البصمات:
  1) ADMS push: جهاز ZKTeco (وأشباهه) بيبعت البصمات لوحده للسيرفر على البورت الداخلي (/iclock/cdata). الجهاز الجديد بيظهر "بانتظار الموافقة".
  2) سحب مباشر بالـ IP من الجهاز (محتاج مكتبة pyzk: pip install pyzk) - اختياري.
  3) استيراد ملف CSV / Excel (تصدير الجهاز أو برنامجه).
الغياب في الرواتب = أيام العمل المفروضة (أيام العقد − الراحة الأسبوعية − العطلات الرسمية − الإجازات المعتمدة) اللي مفيهاش أي بصمة،
وبيتحسب بس للموظف اللي ليه رقم على الجهاز وعنده بصمات في الشهر ده (غير كده الغياب بيفضل يدوي)."""
import csv
import datetime
import os
import re

from . import services as S
from .db import WorkShift, AttendanceDevice, AttendanceLog, Contract, Employee, Holiday, LeaveRequest, Payslip, CompanySettings, get_session

_AR_DIGITS = str.maketrans("٠١٢٣٤٥٦٧٨٩۰۱۲۳۴۵۶۷۸۹", "01234567890123456789")
WEEKDAY_NAMES = {5: "السبت", 6: "الأحد", 0: "الإثنين", 1: "الثلاثاء", 2: "الأربعاء", 3: "الخميس", 4: "الجمعة"}
MAX_PENDING_DEVICES = 20


def norm_id(x):
    """رقم الموظف على الجهاز بعد التنظيف (أصفار البداية والمسافات والأرقام العربية)."""
    t = str(x or "").translate(_AR_DIGITS).strip()
    return t.lstrip("0") or ("0" if t else "")


# ------------------------------------------------------------------ الإعدادات والعطلات
def _parse_days(raw):
    return {int(x) for x in str(raw).split(",") if x.strip().isdigit() and 0 <= int(x) <= 6}


def get_weekly_off(employee_id=None):
    """أيام الراحة الأسبوعية: راحة وردية الموظف لو محددة، وإلا راحة الشركة."""
    with get_session() as s:
        if employee_id:
            e = s.get(Employee, int(employee_id))
            sh = s.get(WorkShift, e.shift_id) if e and e.shift_id else None
            if sh and sh.weekly_off_days is not None and str(sh.weekly_off_days).strip() != "":
                return _parse_days(sh.weekly_off_days)
        c = s.get(CompanySettings, 1)
        raw = (c.weekly_off_days if c and c.weekly_off_days is not None else "4,5")
    return _parse_days(raw)


def set_weekly_off(days):
    S.require_action("attendance", "edit")
    days = sorted({int(d) for d in days if 0 <= int(d) <= 6})
    with get_session() as s:
        c = s.get(CompanySettings, 1)
        if not c:
            c = CompanySettings(id=1); s.add(c)
        c.weekly_off_days = ",".join(str(d) for d in days)
        s.commit()


def list_holidays(year=None):
    with get_session() as s:
        q = s.query(Holiday)
        if year:
            q = q.filter(Holiday.date >= datetime.datetime(year, 1, 1), Holiday.date < datetime.datetime(year + 1, 1, 1))
        rows = q.order_by(Holiday.date).all()
        s.expunge_all()
        return rows


def add_holiday(date, name=""):
    S.require_action("attendance", "create")
    d = datetime.datetime(date.year, date.month, date.day)
    with get_session() as s:
        if s.query(Holiday).filter(Holiday.date == d).first():
            raise ValueError("اليوم ده مسجّل كعطلة بالفعل")
        s.add(Holiday(date=d, name=(name or "").strip()))
        s.commit()


def delete_holiday(holiday_id):
    S.require_action("attendance", "delete")
    with get_session() as s:
        h = s.get(Holiday, int(holiday_id))
        if h:
            s.delete(h); s.commit()


# ------------------------------------------------------------------ ربط الموظفين بأرقام الجهاز
def set_employee_device_id(employee_id, device_user_id):
    S.require_action("attendance", "edit")
    new = (device_user_id or "").translate(_AR_DIGITS).strip()
    with get_session() as s:
        e = s.get(Employee, int(employee_id))
        if not e:
            raise ValueError("الموظف غير موجود")
        if new:
            for o in s.query(Employee).filter(Employee.id != e.id, Employee.device_user_id != "").all():
                if norm_id(o.device_user_id) == norm_id(new):
                    raise ValueError(f'الرقم {new} مربوط بالفعل بالموظف "{o.name}"')
        e.device_user_id = new
        s.commit()
    relink_logs()


def relink_logs():
    """بيربط البصمات القديمة (اللي رقمها ماكانش مربوط بموظف) بالموظفين بعد ربط الأرقام."""
    with get_session() as s:
        by_norm = {norm_id(e.device_user_id): e.id for e in s.query(Employee).filter(Employee.device_user_id != "").all()}
        n = 0
        for lg in s.query(AttendanceLog).filter(AttendanceLog.employee_id.is_(None)).all():
            eid = by_norm.get(lg.norm_id)
            if eid:
                lg.employee_id = eid; n += 1
        s.commit()
        return n


# ------------------------------------------------------------------ تسجيل البصمات
def import_punches(punches, source="file", device_serial=""):
    """punches: قايمة (رقم الجهاز، datetime). بيرجّع {inserted, duplicates, unmatched:[أرقام مش مربوطة]}."""
    inserted = dups = 0
    unmatched = set()
    with get_session() as s:
        by_norm = {norm_id(e.device_user_id): e.id for e in s.query(Employee).filter(Employee.device_user_id != "").all()}
        seen = set()
        for uid, t in punches:
            nid = norm_id(uid)
            if not nid or not isinstance(t, datetime.datetime):
                continue
            t = t.replace(microsecond=0)
            key = (nid, t)
            if key in seen or s.query(AttendanceLog.id).filter(AttendanceLog.norm_id == nid, AttendanceLog.punch_time == t).first():
                dups += 1
                continue
            seen.add(key)
            eid = by_norm.get(nid)
            if not eid:
                unmatched.add(str(uid).strip())
            s.add(AttendanceLog(employee_id=eid, device_user_id=str(uid).strip()[:30], norm_id=nid[:30], punch_time=t, source=source, device_serial=device_serial or ""))
            inserted += 1
        s.commit()
    return {"inserted": inserted, "duplicates": dups, "unmatched": sorted(unmatched)}


_DT_FORMATS = ["%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%Y/%m/%d %H:%M:%S", "%Y/%m/%d %H:%M", "%d/%m/%Y %H:%M:%S", "%d/%m/%Y %H:%M",
               "%d-%m-%Y %H:%M:%S", "%d-%m-%Y %H:%M", "%m/%d/%Y %H:%M:%S", "%m/%d/%Y %I:%M:%S %p", "%m/%d/%Y %I:%M %p", "%d/%m/%Y %I:%M:%S %p"]
_DATE_FORMATS = ["%Y-%m-%d", "%Y/%m/%d", "%d/%m/%Y", "%d-%m-%Y", "%m/%d/%Y"]
_TIME_FORMATS = ["%H:%M:%S", "%H:%M", "%I:%M:%S %p", "%I:%M %p"]


def _parse_dt(v):
    if isinstance(v, datetime.datetime):
        return v
    if isinstance(v, datetime.date):
        return datetime.datetime(v.year, v.month, v.day)
    t = str(v or "").translate(_AR_DIGITS).strip()
    for f in _DT_FORMATS:
        try:
            return datetime.datetime.strptime(t, f)
        except ValueError:
            pass
    return None


def _parse_date_time(d, t):
    if isinstance(d, datetime.datetime):
        base = d
    else:
        base = None
        ds = str(d or "").translate(_AR_DIGITS).strip()
        for f in _DATE_FORMATS:
            try:
                base = datetime.datetime.strptime(ds, f); break
            except ValueError:
                pass
    if base is None:
        return None
    if isinstance(t, datetime.datetime):
        t = t.time()
    if isinstance(t, datetime.time):
        return datetime.datetime.combine(base.date(), t)
    ts = str(t or "").translate(_AR_DIGITS).strip()
    for f in _TIME_FORMATS:
        try:
            return datetime.datetime.combine(base.date(), datetime.datetime.strptime(ts, f).time())
        except ValueError:
            pass
    return None


_ID_HEADS = ("userid", "user id", "user_id", "pin", "badge", "badgenumber", "employee id", "employeeid", "emp id", "empid", "id", "no", "رقم الموظف", "كود الموظف", "الرقم", "رقم", "كود", "رقم البصمة")
_DT_HEADS = ("datetime", "date time", "date/time", "time", "punch", "checktime", "check time", "timestamp", "التاريخ والوقت", "تاريخ ووقت", "الوقت", "وقت البصمة", "التاريخ")
_DATE_HEADS = ("date", "التاريخ", "تاريخ", "اليوم")
_TIME_HEADS = ("time", "الوقت", "وقت", "الساعة", "check in", "checkin", "check out", "checkout", "دخول", "خروج", "حضور", "انصراف")


def _head(h):
    return re.sub(r"\s+", " ", str(h or "").translate(_AR_DIGITS).strip().lower())


def _read_rows(path):
    ext = os.path.splitext(path)[1].lower()
    if ext in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb.active
        return [list(r) for r in ws.iter_rows(values_only=True)]
    with open(path, "rb") as f:
        raw = f.read()
    for enc in ("utf-8-sig", "utf-16", "cp1256", "latin-1"):
        try:
            text = raw.decode(enc); break
        except (UnicodeDecodeError, UnicodeError):
            continue
    first = text.splitlines()[0] if text.strip() else ""
    delim = max([",", ";", "\t", "|"], key=lambda d: first.count(d))
    return [r for r in csv.reader(text.splitlines(), delimiter=delim)]


def parse_attendance_file(path):
    """يقرا ملف الجهاز (CSV أو Excel). الأعمدة بتتعرّف تلقائيًا: رقم الموظف + (تاريخ ووقت في عمود واحد) أو (تاريخ وعمود/أعمدة وقت).
    ملفات ZK الخام من غير عناوين (رقم <tab> تاريخ ووقت ...) مدعومة كمان. بيرجّع (بصمات، عدد صفوف اتخطت)."""
    rows = [r for r in _read_rows(path) if any(str(c or "").strip() for c in r)]
    if not rows:
        raise ValueError("الملف فاضي")
    heads = [_head(c) for c in rows[0]]
    def find(cands, exclude=()):
        for i, h in enumerate(heads):
            if i not in exclude and h in cands:
                return i
        return None
    id_i = find(_ID_HEADS)
    dt_i = find(_DT_HEADS, exclude={id_i} if id_i is not None else ())
    date_i = find(_DATE_HEADS)
    time_is = [i for i, h in enumerate(heads) if h in _TIME_HEADS and i != date_i]
    has_header = id_i is not None and (dt_i is not None or date_i is not None)
    punches, skipped = [], 0
    if has_header:
        body = rows[1:]
        for r in body:
            r = list(r) + [""] * (len(heads) - len(r))
            uid = r[id_i]
            ts = []
            if dt_i is not None and dt_i != date_i:
                t = _parse_dt(r[dt_i])
                if t:
                    ts.append(t)
            if not ts and date_i is not None:
                cand = [i for i in time_is] or ([dt_i] if dt_i not in (None, date_i) else [])
                for i in cand:
                    t = _parse_date_time(r[date_i], r[i])
                    if t:
                        ts.append(t)
                if not cand:
                    t = _parse_dt(r[date_i])
                    if t:
                        ts.append(t)
            if not str(uid or "").strip() or not ts:
                skipped += 1
                continue
            punches += [(uid, t) for t in ts]
        if not punches and body:
            raise ValueError("ما قدرتش أقرا التواريخ - الصيغة المدعومة مثلًا 2026-03-15 08:02 أو 15/03/2026 08:02")
        return punches, skipped
    # من غير عناوين: رقم ثم تاريخ/وقت
    for r in rows:
        cells = [str(c if c is not None else "").strip() for c in r]
        t = _parse_dt(cells[1]) if len(cells) > 1 else None
        if t is None and len(cells) > 2:
            t = _parse_date_time(cells[1], cells[2])
        if cells and cells[0] and t:
            punches.append((cells[0], t))
        else:
            skipped += 1
    if not punches:
        raise ValueError("مش لاقي أعمدة (رقم الموظف + التاريخ والوقت) - اتأكد من عناوين الأعمدة")
    return punches, skipped


def import_file(path):
    S.require_action("attendance", "create")
    punches, skipped = parse_attendance_file(path)
    res = import_punches(punches, source="file")
    res["skipped_rows"] = skipped
    return res


# ------------------------------------------------------------------ الأجهزة (ADMS push + سحب)
def list_devices():
    with get_session() as s:
        rows = s.query(AttendanceDevice).order_by(AttendanceDevice.id).all()
        s.expunge_all()
        return rows


def save_device(serial, name="", host="", port=4370, enabled=True, device_id=None):
    S.require_action("attendance", "edit" if device_id else "create")
    serial = (serial or "").strip()
    if not serial:
        raise ValueError("السيريال نمبر مطلوب (مكتوب على الجهاز أو في قائمة معلومات الجهاز)")
    with get_session() as s:
        dup = s.query(AttendanceDevice).filter(AttendanceDevice.serial == serial)
        if device_id:
            dup = dup.filter(AttendanceDevice.id != int(device_id))
        if dup.first():
            raise ValueError("الجهاز ده مسجّل بالفعل")
        d = s.get(AttendanceDevice, int(device_id)) if device_id else AttendanceDevice(serial=serial)
        if not d:
            raise ValueError("الجهاز غير موجود")
        d.serial = serial; d.name = (name or "").strip(); d.host = (host or "").strip(); d.port = int(port or 4370); d.enabled = bool(enabled)
        if not device_id:
            s.add(d)
        s.commit()
        return d.id


def delete_device(device_id):
    S.require_action("attendance", "delete")
    with get_session() as s:
        d = s.get(AttendanceDevice, int(device_id))
        if d:
            s.delete(d); s.commit()


def adms_handle(method, path, query, body="", ip=""):
    """بيرد على بروتوكول ADMS (push) بتاع ZKTeco. بيرجّع (status, text). الجهاز غير المسجّل بيتسجّل 'بانتظار الموافقة' ومفيش بصمات بتتحفظ منه."""
    sn = (query.get("SN") or query.get("sn") or [""])[0].strip() if isinstance(query.get("SN") or query.get("sn") or [""], list) else str(query.get("SN") or "")
    if not sn or not re.fullmatch(r"[A-Za-z0-9_\-\.]{3,60}", sn):
        return 400, "ERROR: bad SN"
    with get_session() as s:
        dev = s.query(AttendanceDevice).filter(AttendanceDevice.serial == sn).first()
        if dev is None:
            if s.query(AttendanceDevice).filter(AttendanceDevice.enabled.is_(False)).count() >= MAX_PENDING_DEVICES:
                return 403, "ERROR: too many pending devices"
            dev = AttendanceDevice(serial=sn, name="", enabled=False)
            s.add(dev)
        dev.last_seen = datetime.datetime.now(); dev.last_ip = (ip or "")[:60]
        enabled = bool(dev.enabled)
        s.commit()
    p = path.rstrip("/")
    if method == "GET" and p.endswith("/cdata"):
        return 200, (f"GET OPTION FROM: {sn}\nStamp=9999\nOpStamp=9999\nErrorDelay=60\nDelay=30\nTransTimes=00:00;14:00\nTransInterval=1\n"
                     "TransFlag=TransData AttLog\nRealtime=1\nEncrypt=0\n")
    if method == "GET":
        return 200, "OK"
    if p.endswith("/cdata") and method == "POST":
        table = ((query.get("table") or [""])[0] if isinstance(query.get("table"), list) else str(query.get("table") or "")).upper()
        if table != "ATTLOG":
            return 200, "OK: 0"
        if not enabled:
            return 200, "OK: 0"       # الجهاز لسه مش مفعّل - مش بنحفظ حاجة (والجهاز مش هيعيد الإرسال للأبد)
        punches = []
        for line in (body or "").splitlines():
            f = line.split("\t")
            if len(f) >= 2:
                t = _parse_dt(f[1])
                if t and f[0].strip():
                    punches.append((f[0], t))
        res = import_punches(punches, source="device", device_serial=sn)
        with get_session() as s:
            d = s.query(AttendanceDevice).filter(AttendanceDevice.serial == sn).first()
            if d:
                d.last_count = (d.last_count or 0) + res["inserted"]; s.commit()
        return 200, f"OK: {len(punches)}"
    return 200, "OK"


def pull_device(device_id):
    """سحب البصمات مباشرة من الجهاز بالـ IP (يحتاج: pip install pyzk وجهاز يسمح بالاتصال على بورت 4370)."""
    S.require_action("attendance", "create")
    with get_session() as s:
        d = s.get(AttendanceDevice, int(device_id))
        if not d:
            raise ValueError("الجهاز غير موجود")
        host, port, serial = d.host, d.port or 4370, d.serial
    if not host:
        raise ValueError("اكتب عنوان IP الجهاز الأول")
    try:
        from zk import ZK
    except ImportError:
        raise ValueError("مكتبة السحب المباشر غير مثبتة - شغّل: pip install pyzk")
    conn = None
    try:
        conn = ZK(host, port=int(port), timeout=10).connect()
        recs = conn.get_attendance() or []
    except Exception as ex:      # noqa: BLE001
        raise ValueError(f"تعذر الاتصال بالجهاز {host}:{port} - {ex}")
    finally:
        try:
            conn and conn.disconnect()
        except Exception:      # noqa: BLE001
            pass
    res = import_punches([(r.user_id, r.timestamp) for r in recs], source="pull", device_serial=serial)
    with get_session() as s:
        d = s.get(AttendanceDevice, int(device_id))
        d.last_seen = datetime.datetime.now(); s.commit()
    return res



# ------------------------------------------------------------------ قواعد الدوام (تأخير / إضافي)
def _hm(txt, default):
    try:
        h, m = str(txt or default).strip().split(":")[:2]
        h, m = int(h), int(m)
        if 0 <= h < 24 and 0 <= m < 60:
            return h * 60 + m
    except Exception:
        pass
    return _hm(default, "08:00") if default != "08:00" else 8 * 60


def get_shift_rules(employee_id=None):
    """قواعد الدوام: دوام الشركة الافتراضي، أو وردية الموظف (بداية/نهاية/سماحية) لو متربط بوردية. التفعيل والإضافي من إعدادات الشركة."""
    with get_session() as s:
        c = s.get(CompanySettings, 1)
        g = lambda k, d: (getattr(c, k, None) if c is not None and getattr(c, k, None) is not None else d)
        start, end, grace, shift_name = _hm(g("shift_start", "08:00"), "08:00"), _hm(g("shift_end", "16:00"), "16:00"), int(g("late_grace_minutes", 15)), ""
        if employee_id:
            e = s.get(Employee, int(employee_id))
            sh = s.get(WorkShift, e.shift_id) if e and e.shift_id else None
            if sh:
                start, end, shift_name = _hm(sh.start_time, "08:00"), _hm(sh.end_time, "16:00"), sh.name
                if sh.grace_minutes is not None:
                    grace = int(sh.grace_minutes)
        return {"start": start, "end": end, "start_txt": f"{start // 60:02d}:{start % 60:02d}", "end_txt": f"{end // 60:02d}:{end % 60:02d}",
                "shift_name": shift_name, "grace": grace, "late_on": bool(g("late_deduction_on", False)), "ot_on": bool(g("overtime_on", False)),
                "ot_min": int(g("overtime_min_minutes", 30)), "ot_mult": float(g("overtime_multiplier", 1.5)),
                "shift_minutes": max(60, (end - start) if end > start else (end + 1440 - start)),
                **_overnight_fields(start, end)}


def _overnight_fields(start, end):
    """وردية بتعدّي منتصف الليل (النهاية <= البداية؛ تساويهم = 24 ساعة): النهاية بتتحسب في اليوم التالي، والبصمة اللي قبل منتصف
    الفترة اللي بين نهاية الوردية وبداية اللي بعدها (cut) بتتنسب ليوم العمل اللي قبله."""
    over = end <= start
    return {"overnight": over, "end_abs": end + 1440 if over else end, "cut": (end + (start - end) // 2) if over else 0}


def _work_date(dt, rules):
    """يوم العمل اللي البصمة بتتنسب له (للوردية الليلية: بصمة الفجر بتتنسب ليوم الوردية اللي بدأت امبارح)."""
    if rules.get("overnight") and dt.hour * 60 + dt.minute < rules.get("cut", 0):
        return dt.date() - datetime.timedelta(days=1)
    return dt.date()


def _abs_min(dt, work_date):
    return (dt.date() - work_date).days * 1440 + dt.hour * 60 + dt.minute


def set_shift_rules(start, end, grace, late_on, ot_on, ot_min, ot_mult):
    S.require_action("attendance", "edit")
    def chk(t, label):
        try:
            h, m = str(t).strip().split(":"); h, m = int(h), int(m)
            assert 0 <= h < 24 and 0 <= m < 60
        except Exception:
            raise ValueError(f"{label} لازم تكون بالشكل HH:MM (مثال 08:00)")
        return f"{h:02d}:{m:02d}"
    start, end = chk(start, "بداية الدوام"), chk(end, "نهاية الدوام")
    try:
        grace, ot_min, ot_mult = int(grace or 0), int(ot_min or 0), float(ot_mult or 0)
    except (TypeError, ValueError):
        raise ValueError("السماحية والحد الأدنى للإضافي أرقام صحيحة، والمعامل رقم")
    if grace < 0 or ot_min < 0 or ot_mult < 1:
        raise ValueError("القيم غير صحيحة (السماحية والحد الأدنى ≥ 0، ومعامل الإضافي ≥ 1)")
    with get_session() as s:
        c = s.get(CompanySettings, 1)
        if not c:
            c = CompanySettings(id=1); s.add(c)
        c.shift_start, c.shift_end, c.late_grace_minutes = start, end, grace
        c.late_deduction_on, c.overtime_on, c.overtime_min_minutes, c.overtime_multiplier = bool(late_on), bool(ot_on), ot_min, ot_mult
        s.commit()


def day_metrics(first, last, count, rules, off_day=False, work_date=None):
    """دقائق التأخير / الخروج المبكر / الإضافي ليوم عمل واحد من أول وآخر بصمة.
    يوم دوام عادي: تأخير = وقت أول بصمة بعد (بداية الدوام + السماحية) محسوب من بداية الدوام،
    خروج مبكر = قبل النهاية، إضافي = بعد النهاية (لو ≥ الحد الأدنى). يوم راحة/عطلة: كل مدة التواجد إضافي (لو ≥ الحد الأدنى).
    الخروج المبكر والإضافي محتاجين بصمتين على الأقل (بصمة واحدة = مش عارفين وقت الخروج).
    الوردية الليلية: الأوقات بتتحسب من بداية يوم العمل (بصمة بعد منتصف الليل = دقائق + 1440)."""
    wd = work_date or _work_date(first, rules)
    f = _abs_min(first, wd)
    l = _abs_min(last, wd)
    end = rules.get("end_abs", rules["end"])
    late = early = ot = 0
    if off_day:
        if count >= 2 and l - f >= max(1, rules["ot_min"]):
            ot = l - f
        return late, early, ot
    if f > rules["start"] + rules["grace"]:
        late = f - rules["start"]
    if count >= 2:
        if l < end:
            early = end - l
        elif l - end >= max(1, rules["ot_min"]):
            ot = l - end
    return late, early, ot



# ------------------------------------------------------------------ الورديات
def list_shifts(active_only=False):
    with get_session() as s:
        q = s.query(WorkShift).order_by(WorkShift.name)
        if active_only:
            q = q.filter(WorkShift.is_active.is_(True))
        return q.all()


def _check_hm(t, label):
    try:
        h, m = str(t).strip().split(":"); h, m = int(h), int(m)
        assert 0 <= h < 24 and 0 <= m < 60
    except Exception:
        raise ValueError(f"{label} لازم تكون بالشكل HH:MM (مثال 08:00)")
    return f"{h:02d}:{m:02d}"


def save_shift(name, start, end, grace=None, weekly_off=None, is_active=True, shift_id=None):
    """وردية جديدة أو تعديل. grace فارغ = سماحية الشركة. weekly_off: قايمة أيام أو None = راحة الشركة.
    الوردية اللي بتعدّي منتصف الليل مدعومة (مثلًا 22:00 → 06:00، أو 16:00 → 00:00)، ولو البداية = النهاية يبقى 24 ساعة."""
    S.require_action("attendance", "edit")
    name = (name or "").strip()
    if not name:
        raise ValueError("اسم الوردية مطلوب")
    start, end = _check_hm(start, "بداية الوردية"), _check_hm(end, "نهاية الوردية")
    if grace in (None, ""):
        grace = None
    else:
        try:
            grace = int(grace)
        except (TypeError, ValueError):
            raise ValueError("السماحية رقم صحيح")
        if grace < 0:
            raise ValueError("السماحية لا تقل عن صفر")
    off = None if weekly_off is None else ",".join(str(d) for d in sorted({int(d) for d in weekly_off if 0 <= int(d) <= 6}))
    with get_session() as s:
        dup = s.query(WorkShift).filter(WorkShift.name == name)
        if shift_id:
            dup = dup.filter(WorkShift.id != int(shift_id))
        if dup.first():
            raise ValueError("اسم الوردية موجود بالفعل")
        sh = s.get(WorkShift, int(shift_id)) if shift_id else WorkShift()
        if shift_id and not sh:
            raise ValueError("الوردية غير موجودة")
        sh.name, sh.start_time, sh.end_time, sh.grace_minutes, sh.weekly_off_days, sh.is_active = name, start, end, grace, off, bool(is_active)
        if not shift_id:
            s.add(sh)
        s.commit()
        return sh.id


def delete_shift(shift_id):
    S.require_action("attendance", "edit")
    with get_session() as s:
        if s.query(Employee).filter(Employee.shift_id == int(shift_id)).count():
            raise ValueError("في موظفين مربوطين بالوردية دي - فُكّ ربطهم أو غيّر ورديتهم الأول")
        sh = s.get(WorkShift, int(shift_id))
        if sh:
            s.delete(sh); s.commit()


def set_employee_shift(employee_id, shift_id):
    """ربط موظف بوردية (None = دوام الشركة الافتراضي). الوردية بتنطبق على كل البصمات عند الحساب (مفيش تاريخ سريان)."""
    S.require_action("attendance", "edit")
    with get_session() as s:
        e = s.get(Employee, int(employee_id))
        if not e:
            raise ValueError("الموظف غير موجود")
        if shift_id:
            if not s.get(WorkShift, int(shift_id)):
                raise ValueError("الوردية غير موجودة")
            e.shift_id = int(shift_id)
        else:
            e.shift_id = None
        s.commit()


# ------------------------------------------------------------------ الاستعلام والغياب
def daily_summary(date_from, date_to, employee_id=None, limit=1000):
    """صف لكل (موظف، يوم عمل): أول بصمة وآخر بصمة وعدد البصمات. الوردية الليلية: بصمات بعد منتصف الليل بتتنسب ليوم بدء الوردية."""
    t0 = datetime.datetime(date_from.year, date_from.month, date_from.day)
    t1 = datetime.datetime(date_to.year, date_to.month, date_to.day) + datetime.timedelta(days=1)
    out = {}
    rules_cache = {}

    def rules_of(eid):
        if eid not in rules_cache:
            rules_cache[eid] = (get_shift_rules(eid), get_weekly_off(eid))
        return rules_cache[eid][0]
    with get_session() as s:
        q = s.query(AttendanceLog).filter(AttendanceLog.punch_time >= t0 - datetime.timedelta(days=1), AttendanceLog.punch_time < t1 + datetime.timedelta(days=1))
        if employee_id:
            q = q.filter(AttendanceLog.employee_id == int(employee_id))
        names = {e.id: e.name for e in s.query(Employee).all()}
        for lg in q.order_by(AttendanceLog.punch_time).all():
            wd = _work_date(lg.punch_time, rules_of(lg.employee_id))
            if not (date_from <= wd <= date_to):
                continue
            key = (lg.employee_id or ("?" + lg.norm_id), wd)
            r = out.setdefault(key, {"employee_id": lg.employee_id, "name": names.get(lg.employee_id) or f"غير مربوط ({lg.device_user_id})",
                                     "date": wd, "first": lg.punch_time, "last": lg.punch_time, "count": 0})
            r["last"] = lg.punch_time; r["count"] += 1
    with get_session() as s:
        hol = {h.date.date() for h in s.query(Holiday).filter(Holiday.date >= t0, Holiday.date < t1).all()}
    for r in out.values():
        eid = r["employee_id"]
        rules = rules_of(eid); off = rules_cache[eid][1]
        r["shift"] = rules["shift_name"]
        r["late"], r["early"], r["overtime"] = day_metrics(r["first"], r["last"], r["count"], rules, r["date"].weekday() in off or r["date"] in hol, work_date=r["date"])
    rows = sorted(out.values(), key=lambda r: (r["date"], r["name"]), reverse=True)
    return rows[:limit]


def _employed_dates(s, employee_id, first, last):
    days = set()
    for c in s.query(Contract).filter(Contract.employee_id == int(employee_id), S._contract_state_filter()).all():
        e = c.termination_date or c.end_date
        a, b = max(c.start_date.date(), first), min(e.date() if e else last, last)
        d = a
        while d <= b:
            days.add(d); d += datetime.timedelta(days=1)
    return days


def attendance_absent_days(s, employee_id, month, year, today=None):
    """أيام الغياب من البصمة، أو None لو مفيش بيانات بصمة للموظف في الشهر (فالغياب يدوي)."""
    emp = s.get(Employee, int(employee_id))
    if not emp or not (emp.device_user_id or "").strip():
        return None
    first, last = S._month_bounds(month, year)
    t0 = datetime.datetime(first.year, first.month, first.day)
    t1 = datetime.datetime(last.year, last.month, last.day) + datetime.timedelta(days=1)
    nid = norm_id(emp.device_user_id)
    _r = get_shift_rules(emp.id)
    one = datetime.timedelta(days=1)
    punch_days = {_work_date(lg.punch_time, _r) for lg in s.query(AttendanceLog).filter(
        AttendanceLog.punch_time >= t0 - one, AttendanceLog.punch_time < t1 + one,
        (AttendanceLog.employee_id == emp.id) | (AttendanceLog.norm_id == nid)).all()}
    punch_days = {d for d in punch_days if first <= d <= last}
    if not punch_days:
        return None
    today = today or datetime.date.today()
    cutoff = min(last, today - datetime.timedelta(days=1))       # اليوم الجاري لسه ما خلصش
    expected = {d for d in _employed_dates(s, emp.id, first, cutoff)}
    off = get_weekly_off(emp.id)
    expected = {d for d in expected if d.weekday() not in off}
    expected -= {h.date.date() for h in s.query(Holiday).filter(Holiday.date >= t0, Holiday.date < t1).all()}
    for r in s.query(LeaveRequest).filter(LeaveRequest.employee_id == emp.id, LeaveRequest.state == "approved",
                                          LeaveRequest.start_date < t1, LeaveRequest.end_date >= t0).all():
        d = max(r.start_date.date(), first)
        while d <= min(r.end_date.date(), last):
            expected.discard(d); d += datetime.timedelta(days=1)
    return float(len([d for d in expected if d not in punch_days]))


def attendance_time_minutes(s, employee_id, month, year, today=None):
    """(دقائق تأخير، دقائق خروج مبكر، دقائق إضافي) للموظف في الشهر من البصمة، أو None لو مفيش بصمات/رقم جهاز.
    الأيام اللي الموظف فيها غير موظف (قبل التعيين/بعد الإنهاء) أو في إجازة معتمدة مش بتتحسب تأخير؛ اليوم الجاري مش بيتحسب (لسه ما خلصش)."""
    emp = s.get(Employee, int(employee_id))
    if not emp or not (emp.device_user_id or "").strip():
        return None
    first, last = S._month_bounds(month, year)
    t0 = datetime.datetime(first.year, first.month, first.day)
    t1 = datetime.datetime(last.year, last.month, last.day) + datetime.timedelta(days=1)
    nid = norm_id(emp.device_user_id)
    days = {}
    rules = get_shift_rules(emp.id)
    one = datetime.timedelta(days=1)
    for lg in s.query(AttendanceLog).filter(AttendanceLog.punch_time >= t0 - one, AttendanceLog.punch_time < t1 + one,
                                            (AttendanceLog.employee_id == emp.id) | (AttendanceLog.norm_id == nid)).order_by(AttendanceLog.punch_time).all():
        wd = _work_date(lg.punch_time, rules)
        if not (first <= wd <= last):
            continue
        r = days.setdefault(wd, [lg.punch_time, lg.punch_time, 0])
        r[1] = lg.punch_time; r[2] += 1
    if not days:
        return None
    today = today or datetime.date.today()
    cutoff = min(last, today - datetime.timedelta(days=1))
    employed = _employed_dates(s, emp.id, first, cutoff)
    on_leave = set()
    for r in s.query(LeaveRequest).filter(LeaveRequest.employee_id == emp.id, LeaveRequest.state == "approved",
                                          LeaveRequest.start_date < t1, LeaveRequest.end_date >= t0).all():
        d = max(r.start_date.date(), first)
        while d <= min(r.end_date.date(), last):
            on_leave.add(d); d += datetime.timedelta(days=1)
    off = get_weekly_off(emp.id)
    hol = {h.date.date() for h in s.query(Holiday).filter(Holiday.date >= t0, Holiday.date < t1).all()}
    late = early = ot = 0
    for d, (f, l, c) in days.items():
        if d not in employed or d in on_leave:
            continue
        a, b, o = day_metrics(f, l, c, rules, d.weekday() in off or d in hol, work_date=d)
        late += a; early += b; ot += o
    return float(late), float(early), float(ot)


def apply_attendance_to_run(s, run):
    """بيحسب أيام الغياب من البصمة لكل قسيمة في الدورة (اللي ليها بيانات بصمة) ويعيد حساب الصافي. بيرجّع عدد القسايم اللي اتحدّثت."""
    n = 0
    for p in s.query(Payslip).filter(Payslip.run_id == run.id).all():
        a = attendance_absent_days(s, p.employee_id, run.month, run.year)
        tm = attendance_time_minutes(s, p.employee_id, run.month, run.year)
        if a is not None or tm is not None:
            if a is not None:
                p.absent_days = a
            if tm is not None:
                p.late_minutes = tm[0] + tm[1]          # التأخير + الخروج المبكر
                p.overtime_minutes = tm[2]
            S.recalc_payslip(s, p, run)
            n += 1
    return n


def sync_run_absence(run_id):
    """زرار 'احسب الغياب من البصمة' في شاشة الرواتب (دورة مسودة بس)."""
    S.require_action("payroll", "edit")
    from .db import PayrollRun
    with get_session() as s:
        run = s.get(PayrollRun, int(run_id))
        if not run:
            raise ValueError("الدورة غير موجودة")
        if run.state != "draft":
            raise ValueError("مينفعش تعدّل على دورة مرحّلة - رجّعها مسودة أولاً")
        n = apply_attendance_to_run(s, run)
        s.commit()
        return n
