"""الربط مع البنوك المحلية (هيكل عام بدون مزوّد محدد):
- حسابات بنكية مربوطة بشجرة الحسابات
- مزوّدو دفع: جهاز شبكة (إدخال رقم الموافقة) / تجريبي / API عام بقوالب (جهاز أو بوابة أونلاين)
- عمليات دفع (بطاقة على الكاشير، روابط دفع أونلاين) وربطها بالفواتير والسندات
- استيراد كشف البنك ومطابقته مع القيود
- ملفات تحويل الرواتب والموردين بصيغة عامة (CSV/Excel) قابلة للتعديل حسب صيغة البنك
"""
import csv
import datetime
import hashlib
import json
import os
import re
import uuid

from sqlalchemy.orm import joinedload

from . import services
from .db import (BankAccount, BankStatementLine, Customer, Employee, JournalEntry, JournalEntryLine, PaymentMethod, PaymentProvider,
                 PaymentTransaction, Supplier, get_session)

PROVIDER_KINDS = {
    "terminal_manual": "جهاز شبكة (مدى/فيزا) - إدخال رقم الموافقة",
    "mock": "تجريبي (بدون بنك فعلي)",
    "terminal_http": "جهاز/وسيط عبر API (قوالب)",
    "online_http": "بوابة دفع أونلاين عبر API (قوالب)",
}


# ---------------------------------------------------------------- مساعدات
def validate_iban(iban):
    """تحقق mod-97 من الآيبان (السعودي: SA + 22 رقم)."""
    v = re.sub(r"\s+", "", iban or "").upper()
    if not re.fullmatch(r"[A-Z]{2}\d{2}[A-Z0-9]{10,30}", v):
        return False
    if v.startswith("SA") and len(v) != 24:
        return False
    num = "".join(str(int(c, 36)) for c in v[4:] + v[:4])
    return int(num) % 97 == 1


def _user():
    return (services.current_user() or {}).get("username", "")


def _cfg(p):
    try:
        return json.loads(p.config_json or "{}")
    except ValueError:
        return {}


def _get_path(data, path):
    cur = data
    for part in [x for x in (path or "").split(".") if x]:
        if isinstance(cur, list):
            try:
                cur = cur[int(part)]
            except (ValueError, IndexError):
                return None
        elif isinstance(cur, dict):
            cur = cur.get(part)
        else:
            return None
    return cur


def _fill(obj, ctx):
    if isinstance(obj, str):
        for k, v in ctx.items():
            obj = obj.replace("{" + k + "}", str(v))
        return obj
    if isinstance(obj, dict):
        return {k: _fill(v, ctx) for k, v in obj.items()}
    if isinstance(obj, list):
        return [_fill(v, ctx) for v in obj]
    return obj


# ---------------------------------------------------------------- الحسابات البنكية
def list_bank_accounts(active_only=False):
    with get_session() as s:
        q = s.query(BankAccount).options(joinedload(BankAccount.account))
        if active_only:
            q = q.filter(BankAccount.is_active.is_(True))
        rows = q.order_by(BankAccount.code).all()
        s.expunge_all()
        return rows


def save_bank_account(code, name, account_id, bank_name="", iban="", swift="", establishment_id="", is_active=True, bank_account_id=None):
    services.require_action("banking", "edit" if bank_account_id else "create")
    code = (code or "").strip(); name = (name or "").strip()
    if not code or not name:
        raise ValueError("كود واسم الحساب البنكي مطلوبين")
    if not account_id:
        raise ValueError("اربط الحساب البنكي بحساب من شجرة الحسابات")
    if iban and not validate_iban(iban):
        raise ValueError("الآيبان غير صحيح (السعودي: SA + 22 رقم)")
    with get_session() as s:
        q = s.query(BankAccount).filter(BankAccount.code == code)
        if bank_account_id:
            q = q.filter(BankAccount.id != int(bank_account_id))
        if q.first():
            raise ValueError("كود الحساب البنكي موجود قبل كده")
        b = s.get(BankAccount, int(bank_account_id)) if bank_account_id else BankAccount()
        b.code, b.name, b.bank_name, b.iban, b.swift = code, name, (bank_name or "").strip(), re.sub(r"\s+", "", iban or "").upper(), (swift or "").strip()
        b.establishment_id = (establishment_id or "").strip(); b.account_id = int(account_id); b.is_active = bool(is_active)
        if not bank_account_id:
            s.add(b)
        s.commit()
        return b.id


def delete_bank_account(bank_account_id):
    services.require_action("banking", "delete")
    with get_session() as s:
        if s.query(PaymentProvider).filter(PaymentProvider.bank_account_id == int(bank_account_id)).count() or \
           s.query(BankStatementLine).filter(BankStatementLine.bank_account_id == int(bank_account_id)).count():
            raise ValueError("الحساب عليه مزوّدين أو كشوف مستوردة - أوقفه بدل الحذف")
        s.delete(s.get(BankAccount, int(bank_account_id))); s.commit()


# ---------------------------------------------------------------- المزوّدون
def list_providers(active_only=False):
    with get_session() as s:
        q = s.query(PaymentProvider).options(joinedload(PaymentProvider.bank_account))
        if active_only:
            q = q.filter(PaymentProvider.is_active.is_(True))
        rows = q.order_by(PaymentProvider.name).all()
        s.expunge_all()
        return rows


def get_provider(provider_id):
    with get_session() as s:
        p = s.query(PaymentProvider).options(joinedload(PaymentProvider.bank_account)).filter(PaymentProvider.id == int(provider_id)).first()
        if p:
            s.expunge_all()
        return p


def save_provider(name, kind, bank_account_id=None, config_json="{}", is_active=True, provider_id=None):
    services.require_action("banking", "edit" if provider_id else "create")
    name = (name or "").strip()
    if not name:
        raise ValueError("اسم المزوّد مطلوب")
    if kind not in PROVIDER_KINDS:
        raise ValueError("نوع المزوّد غير معروف")
    try:
        cfg = json.loads(config_json or "{}")
        if not isinstance(cfg, dict):
            raise ValueError
    except ValueError:
        raise ValueError("إعدادات الـ API لازم تكون JSON صحيح")
    if kind.endswith("_http") and not cfg.get("base_url"):
        raise ValueError("حدد base_url في إعدادات الـ API")
    with get_session() as s:
        q = s.query(PaymentProvider).filter(PaymentProvider.name == name)
        if provider_id:
            q = q.filter(PaymentProvider.id != int(provider_id))
        if q.first():
            raise ValueError("اسم المزوّد موجود قبل كده")
        p = s.get(PaymentProvider, int(provider_id)) if provider_id else PaymentProvider()
        p.name, p.kind, p.bank_account_id, p.config_json, p.is_active = name, kind, int(bank_account_id) if bank_account_id else None, json.dumps(cfg, ensure_ascii=False), bool(is_active)
        if not provider_id:
            s.add(p)
        s.commit()
        return p.id


def delete_provider(provider_id):
    services.require_action("banking", "delete")
    with get_session() as s:
        if s.query(PaymentMethod).filter(PaymentMethod.provider_id == int(provider_id)).count() or \
           s.query(PaymentTransaction).filter(PaymentTransaction.provider_id == int(provider_id)).count():
            raise ValueError("المزوّد مستخدم في طرق دفع أو عمليات - أوقفه بدل الحذف")
        s.delete(s.get(PaymentProvider, int(provider_id))); s.commit()


def provider_needs_reference(provider_id):
    """جهاز الشبكة اليدوي: الكاشير لازم يكتب رقم الموافقة/المرجع من الجهاز."""
    p = get_provider(provider_id) if provider_id else None
    return bool(p and p.is_active and p.kind == "terminal_manual")


# ---------------------------------------------------------------- تنفيذ الدفع
def _http(cfg, section, ctx):
    import httpx
    sec = cfg.get(section) or {}
    if not sec.get("path"):
        raise ValueError(f"إعدادات «{section}» ناقصة (path) في المزوّد")
    url = cfg["base_url"].rstrip("/") + "/" + sec["path"].lstrip("/")
    headers = _fill(cfg.get("headers") or {}, ctx)
    body = _fill(sec.get("body"), ctx) if sec.get("body") is not None else None
    method = (sec.get("method") or "POST").upper()
    r = httpx.request(method, _fill(url, ctx), headers=headers, json=body if method != "GET" else None, timeout=float(cfg.get("timeout", 60)))
    try:
        data = r.json()
    except ValueError:
        data = {"raw": r.text}
    return r.status_code, data, sec


def _new_txn(p, kind, amount, **kw):
    with get_session() as s:
        t = PaymentTransaction(provider_id=p.id, kind=kind, amount=round(float(amount), 2), created_by=_user(), **kw)
        s.add(t); s.commit()
        return t.id


def _update_txn(txn_id, **kw):
    with get_session() as s:
        t = s.get(PaymentTransaction, int(txn_id))
        for k, v in kw.items():
            setattr(t, k, v)
        s.commit()


def charge_card(provider_id, amount, manual_reference="", payment_method_id=None, customer_id=None, description=""):
    """خصم بطاقة على جهاز/مزوّد. بيرجّع dict فيه status=approved وreference وcard_last4؛ وبيرفع ValueError لو مرفوضة."""
    p = get_provider(provider_id)
    if not p:
        raise ValueError("مزوّد الدفع غير موجود")
    if not p.is_active:
        raise ValueError(f"مزوّد الدفع «{p.name}» موقوف")
    amount = round(float(amount), 2)
    if amount <= 0:
        raise ValueError("مبلغ العملية لازم يكون أكبر من صفر")
    tid = _new_txn(p, "charge", amount, payment_method_id=int(payment_method_id) if payment_method_id else None, customer_id=int(customer_id) if customer_id else None)
    try:
        if p.kind == "terminal_manual":
            ref = (manual_reference or "").strip()
            if not ref:
                raise ValueError(f"اكتب رقم الموافقة/المرجع من جهاز الشبكة ({p.name})")
            res = {"status": "approved", "reference": ref, "card_last4": "", "message": "تمت الموافقة على الجهاز (إدخال يدوي)"}
        elif p.kind == "mock":
            if int(round(amount * 100)) % 100 == 13:
                res = {"status": "declined", "reference": "", "card_last4": "", "message": "مرفوضة (تجريبي: المبلغ بينتهي بـ .13)"}
            else:
                res = {"status": "approved", "reference": "MOCK" + uuid.uuid4().hex[:8].upper(), "card_last4": "4242", "message": "موافقة تجريبية"}
        elif p.kind in ("terminal_http", "online_http"):
            cfg = _cfg(p)
            ctx = {"amount": f"{amount:.2f}", "amount_minor": int(round(amount * 100)), "currency": cfg.get("currency", "SAR"), "reference": f"TX{tid}",
                   "description": description or "", "terminal": cfg.get("terminal_id", ""), "merchant": cfg.get("merchant_id", ""), "txn_id": tid}
            code, data, sec = _http(cfg, "charge", ctx)
            ok_values = [str(x).lower() for x in (sec.get("ok_values") or ["true", "approved", "success", "paid", "00", "0"])]
            okv = str(_get_path(data, sec.get("ok_path", "status"))).lower()
            ok = 200 <= code < 300 and okv in ok_values
            res = {"status": "approved" if ok else "declined", "reference": str(_get_path(data, sec.get("ref_path", "reference")) or ""),
                   "card_last4": str(_get_path(data, sec.get("last4_path", "")) or ""), "external_id": str(_get_path(data, sec.get("id_path", "")) or ""),
                   "message": str(_get_path(data, sec.get("message_path", "message")) or f"HTTP {code}")}
        else:
            raise ValueError("نوع المزوّد لا يدعم الخصم المباشر")
    except ValueError as ex:
        _update_txn(tid, status="failed", message=str(ex))
        raise
    except Exception as ex:   # أخطاء الشبكة
        _update_txn(tid, status="failed", message=f"تعذر الاتصال بالمزوّد: {ex}")
        raise ValueError(f"تعذر الاتصال بمزوّد الدفع «{p.name}»: {ex}")
    _update_txn(tid, status=res["status"], reference=res.get("reference", ""), card_last4=res.get("card_last4", ""), external_id=res.get("external_id", ""), message=res.get("message", ""))
    if res["status"] != "approved":
        raise ValueError(f"العملية مرفوضة من البنك: {res.get('message', '')}")
    return {**res, "transaction_id": tid, "provider": p.name}


def attach_invoice(transaction_id, invoice_id):
    if transaction_id:
        _update_txn(transaction_id, invoice_id=int(invoice_id))


def mark_orphaned(transaction_id, message):
    """العملية اتخصمت من العميل لكن الفاتورة فشلت - بتتعلّم عشان تتراجع من الجهاز/البنك."""
    if transaction_id:
        _update_txn(transaction_id, status="orphaned", message=message)


def create_payment_link(provider_id, amount, customer_id=None, description=""):
    """رابط دفع أونلاين: بيتبعت للعميل، وبعد السداد (التحقق من الحالة) بيتعمل سند قبض تلقائي على حساب البنك."""
    services.require_action("banking", "create")
    p = get_provider(provider_id)
    if not p or not p.is_active:
        raise ValueError("اختر مزوّد دفع شغّال")
    amount = round(float(amount or 0), 2)
    if amount <= 0:
        raise ValueError("اكتب مبلغ أكبر من صفر")
    if not customer_id:
        raise ValueError("اختر العميل")
    tid = _new_txn(p, "link", amount, customer_id=int(customer_id))
    try:
        if p.kind == "mock":
            ext = uuid.uuid4().hex[:12]; url = f"https://pay.example.test/mock/{ext}"
        elif p.kind == "online_http":
            cfg = _cfg(p)
            ctx = {"amount": f"{amount:.2f}", "amount_minor": int(round(amount * 100)), "currency": cfg.get("currency", "SAR"), "reference": f"TX{tid}", "description": description or "", "merchant": cfg.get("merchant_id", ""), "txn_id": tid}
            code, data, sec = _http(cfg, "link", ctx)
            ext = str(_get_path(data, sec.get("id_path", "id")) or ""); url = str(_get_path(data, sec.get("url_path", "url")) or "")
            if not (200 <= code < 300 and url):
                raise ValueError(f"المزوّد رفض إنشاء الرابط: {_get_path(data, sec.get('message_path', 'message')) or code}")
        else:
            raise ValueError("المزوّد ده مش بوابة أونلاين - اختر مزوّد من نوع بوابة أو تجريبي")
    except Exception as ex:
        _update_txn(tid, status="failed", message=str(ex))
        raise ValueError(str(ex))
    _update_txn(tid, external_id=ext, link_url=url, message=description or "")
    return get_transaction(tid)


def check_payment_link(transaction_id):
    """يتحقق من حالة رابط الدفع؛ لو اتسدد بيعمل سند قبض على حساب البنك (مرة واحدة بس)."""
    t = get_transaction(transaction_id)
    if not t or t.kind != "link":
        raise ValueError("العملية غير موجودة")
    if t.status == "paid":
        return t
    p = get_provider(t.provider_id)
    paid, ref = False, t.reference
    if p.kind == "mock":
        paid, ref = True, "MOCKPAID" + (t.external_id or "")[:6].upper()
    elif p.kind == "online_http":
        cfg = _cfg(p)
        code, data, sec = _http(cfg, "status", {"external_id": t.external_id, "reference": f"TX{t.id}", "merchant": cfg.get("merchant_id", ""), "txn_id": t.id})
        paid_values = [str(x).lower() for x in (sec.get("paid_values") or ["paid", "success", "captured", "true"])]
        paid = 200 <= code < 300 and str(_get_path(data, sec.get("paid_path", "status"))).lower() in paid_values
        ref = str(_get_path(data, sec.get("ref_path", "")) or ref or "")
    if not paid:
        return t
    bank_gl = p.bank_account.account_id if p.bank_account else None
    if not bank_gl:
        raise ValueError("اربط المزوّد بحساب بنكي عشان يتسجل سند القبض")
    pid = services.create_receipt(t.customer_id, t.amount, bank_gl, reference=ref or f"LINK{t.id}", notes=f"سداد أونلاين عبر {p.name}", partner_type="customer")
    _update_txn(t.id, status="paid", reference=ref or "", receipt_payment_id=pid)
    return get_transaction(t.id)


def get_transaction(transaction_id):
    with get_session() as s:
        t = s.query(PaymentTransaction).options(joinedload(PaymentTransaction.provider), joinedload(PaymentTransaction.customer)).filter(PaymentTransaction.id == int(transaction_id)).first()
        if t:
            s.expunge_all()
        return t


def list_transactions(kind=None, status=None, limit=300):
    with get_session() as s:
        q = s.query(PaymentTransaction).options(joinedload(PaymentTransaction.provider), joinedload(PaymentTransaction.customer))
        if kind:
            q = q.filter(PaymentTransaction.kind == kind)
        if status:
            q = q.filter(PaymentTransaction.status == status)
        rows = q.order_by(PaymentTransaction.id.desc()).limit(limit).all()
        s.expunge_all()
        return rows


# ---------------------------------------------------------------- كشف البنك والمطابقة
_HDR = {
    "date": ("date", "تاريخ", "التاريخ", "value date", "posting date"),
    "description": ("description", "details", "narrative", "البيان", "التفاصيل", "الوصف", "بيان"),
    "reference": ("reference", "ref", "المرجع", "رقم المرجع", "رقم العملية", "cheque", "الشيك"),
    "debit": ("debit", "withdrawal", "withdrawals", "مدين", "سحب", "مسحوبات", "المدين"),
    "credit": ("credit", "deposit", "deposits", "دائن", "إيداع", "ايداع", "مودعات", "الدائن"),
    "amount": ("amount", "المبلغ", "قيمة"),
    "balance": ("balance", "الرصيد", "رصيد"),
}


def _num(v):
    if v in (None, ""):
        return 0.0
    if isinstance(v, (int, float)):
        return float(v)
    t = str(v).replace(",", "").replace("٬", "").replace("٫", ".").strip()
    neg = t.startswith("(") and t.endswith(")") or t.startswith("-") or t.endswith("-") or t.upper().endswith("DR")
    t = re.sub(r"[^0-9.]", "", t)
    try:
        x = float(t) if t else 0.0
    except ValueError:
        return 0.0
    return -x if neg else x


def _date(v):
    if isinstance(v, datetime.datetime):
        return v
    if isinstance(v, datetime.date):
        return datetime.datetime.combine(v, datetime.time())
    t = str(v or "").strip()[:19]
    for f in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%Y/%m/%d", "%m/%d/%Y", "%Y-%m-%d %H:%M:%S", "%d.%m.%Y"):
        try:
            return datetime.datetime.strptime(t, f)
        except ValueError:
            pass
    return None


def _read_rows(path):
    ext = os.path.splitext(path)[1].lower()
    if ext in (".xlsx", ".xlsm"):
        import openpyxl
        wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
        return [list(r) for r in wb.active.iter_rows(values_only=True)]
    with open(path, "r", encoding="utf-8-sig", errors="replace", newline="") as f:
        text = f.read()
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    return [row for row in csv.reader(text.splitlines(), dialect)]


def import_statement(bank_account_id, path):
    """يستورد كشف البنك (CSV/Excel): بيعرّف الأعمدة من أسماء العناوين (عربي/إنجليزي)، ويتجاهل السطور المكررة المستوردة قبل كده. بيرجّع (عدد جديد، عدد مكرر)."""
    services.require_action("banking", "create")
    rows = _read_rows(path)
    hdr_i, cols = None, {}
    for i, r in enumerate(rows[:30]):
        found = {}
        for j, c in enumerate(r):
            t = str(c or "").strip().lower()
            for key, names in _HDR.items():
                if key not in found and t and any(t == n.lower() or t.startswith(n.lower()) for n in names):
                    found[key] = j
        if "date" in found and ("debit" in found or "credit" in found or "amount" in found):
            hdr_i, cols = i, found
            break
    if hdr_i is None:
        raise ValueError("مش لاقي عناوين الأعمدة (التاريخ + مدين/دائن أو المبلغ) - راجع ملف الكشف")
    batch = datetime.datetime.now().strftime("%Y%m%d%H%M%S")
    new = dup = 0
    with get_session() as s:
        seen = {h for (h,) in s.query(BankStatementLine.row_hash).filter(BankStatementLine.bank_account_id == int(bank_account_id)).all()}
        for r in rows[hdr_i + 1:]:
            g = lambda k: r[cols[k]] if k in cols and cols[k] < len(r) else None
            d = _date(g("date"))
            if not d:
                continue
            if "amount" in cols and "debit" not in cols and "credit" not in cols:
                a = _num(g("amount")); debit, credit = (-a, 0.0) if a < 0 else (0.0, a)
            else:
                debit, credit = abs(_num(g("debit"))), abs(_num(g("credit")))
            if not debit and not credit:
                continue
            desc = str(g("description") or "").strip()[:300]; ref = str(g("reference") or "").strip()[:100]
            h = hashlib.sha1(f"{bank_account_id}|{d:%Y-%m-%d}|{desc}|{ref}|{debit:.2f}|{credit:.2f}".encode()).hexdigest()
            if h in seen:
                dup += 1
                continue
            seen.add(h)
            s.add(BankStatementLine(bank_account_id=int(bank_account_id), batch=batch, date=d, description=desc, reference=ref, debit=debit, credit=credit,
                                    balance=_num(g("balance")) if g("balance") not in (None, "") else None, row_hash=h))
            new += 1
        s.commit()
    return new, dup


def auto_match(bank_account_id, days=5):
    """يطابق سطور الكشف غير المطابقة مع سطور قيود حساب البنك بنفس المبلغ (إيداع=مدين، سحب=دائن) وتاريخ متقارب. بيرجّع عدد المطابق."""
    services.require_action("banking", "edit")
    n = 0
    with get_session() as s:
        b = s.get(BankAccount, int(bank_account_id))
        used = {j for (j,) in s.query(BankStatementLine.journal_line_id).filter(BankStatementLine.journal_line_id.isnot(None)).all()}
        gl = (s.query(JournalEntryLine, JournalEntry).join(JournalEntry, JournalEntry.id == JournalEntryLine.entry_id)
              .filter(JournalEntryLine.account_id == b.account_id).all())
        pool = [(l, e) for l, e in gl if l.id not in used]
        for line in s.query(BankStatementLine).filter(BankStatementLine.bank_account_id == b.id, BankStatementLine.status == "unmatched").order_by(BankStatementLine.date).all():
            amt = line.credit or line.debit
            best = None
            for l, e in pool:
                gl_amt = l.debit if line.credit else l.credit
                if abs((gl_amt or 0) - amt) > 0.005:
                    continue
                gap = abs((e.date - line.date).days)
                if gap > days:
                    continue
                score = gap - (10 if line.reference and line.reference in f"{e.reference} {l.description}" else 0)
                if best is None or score < best[0]:
                    best = (score, l)
            if best:
                line.journal_line_id = best[1].id; line.status = "matched"
                pool = [(l, e) for l, e in pool if l.id != best[1].id]
                n += 1
        s.commit()
    return n


def list_statement_lines(bank_account_id, status=None):
    with get_session() as s:
        q = s.query(BankStatementLine).filter(BankStatementLine.bank_account_id == int(bank_account_id))
        if status:
            q = q.filter(BankStatementLine.status == status)
        rows = q.order_by(BankStatementLine.date, BankStatementLine.id).all()
        s.expunge_all()
        return rows


def unmatched_gl_lines(bank_account_id):
    """سطور دفتر الأستاذ لحساب البنك اللي لسه مالهاش مقابل في الكشف (شيكات/تحويلات متأخرة...)."""
    with get_session() as s:
        b = s.get(BankAccount, int(bank_account_id))
        used = {j for (j,) in s.query(BankStatementLine.journal_line_id).filter(BankStatementLine.journal_line_id.isnot(None)).all()}
        out = []
        for l, e in (s.query(JournalEntryLine, JournalEntry).join(JournalEntry, JournalEntry.id == JournalEntryLine.entry_id)
                     .filter(JournalEntryLine.account_id == b.account_id).order_by(JournalEntry.date).all()):
            if l.id not in used:
                out.append({"id": l.id, "date": e.date, "number": e.number if hasattr(e, "number") else "", "description": l.description or e.description or "", "debit": l.debit or 0.0, "credit": l.credit or 0.0})
        return out


def reconciliation_summary(bank_account_id):
    lines = list_statement_lines(bank_account_id)
    with get_session() as s:
        b = s.get(BankAccount, int(bank_account_id))
        acc_id = b.account_id
        gl = sum((d or 0) - (c or 0) for d, c in s.query(JournalEntryLine.debit, JournalEntryLine.credit).filter(JournalEntryLine.account_id == acc_id).all())
    bank_bal = None
    if lines:
        with_bal = [l for l in lines if l.balance is not None]
        bank_bal = with_bal[-1].balance if with_bal else sum(l.credit - l.debit for l in lines)
    unm = [l for l in lines if l.status == "unmatched"]
    return {"gl_balance": gl, "bank_balance": bank_bal, "difference": None if bank_bal is None else bank_bal - gl, "count": len(lines), "unmatched": len(unm),
            "unmatched_in": sum(l.credit for l in unm), "unmatched_out": sum(l.debit for l in unm), "gl_unmatched": len(unmatched_gl_lines(bank_account_id))}


def post_statement_line(line_id, counter_account_id, description=""):
    """قيد لسطر كشف مالوش قيد (رسوم بنكية، فوائد، إيداع غير مسجل...) ويتعلّم مطابق."""
    services.require_action("banking", "post")
    with get_session() as s:
        l = s.get(BankStatementLine, int(line_id))
        if not l or l.status != "unmatched":
            raise ValueError("السطر غير موجود أو متطابق بالفعل")
        b = s.get(BankAccount, l.bank_account_id)
        bank_acc, amt, is_in, d, ref = b.account_id, l.credit or l.debit, bool(l.credit), l.date, l.reference or f"BS{l.id}"
        desc = description or l.description or "حركة بنكية"
    lines = ([{"account_id": bank_acc, "debit": amt}, {"account_id": int(counter_account_id), "credit": amt}] if is_in
             else [{"account_id": int(counter_account_id), "debit": amt}, {"account_id": bank_acc, "credit": amt}])
    je = services.create_journal_entry(lines, reference=ref, description=desc, source_type="bank_statement", source_id=int(line_id), entry_date=d)
    with get_session() as s:
        jl = s.query(JournalEntryLine).filter(JournalEntryLine.entry_id == je, JournalEntryLine.account_id == bank_acc).first()
        l = s.get(BankStatementLine, int(line_id)); l.journal_line_id = jl.id; l.status = "posted"; s.commit()
    return je


def set_line_status(line_id, status):
    """ignored = تجاهل السطر، unmatched = فك المطابقة (غير مسموح لو اتعمل له قيد)."""
    services.require_action("banking", "edit")
    with get_session() as s:
        l = s.get(BankStatementLine, int(line_id))
        if l.status == "posted":
            raise ValueError("السطر ده اتعمل له قيد - مينفعش يتغير")
        l.status = status; l.journal_line_id = None; s.commit()


# ---------------------------------------------------------------- ملفات التحويل البنكي
def _write_transfer_files(label, bank, header_info, rows, columns):
    import openpyxl
    out_dir = services.downloads_dir()
    stamp = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
    base = os.path.join(out_dir, f"{label}_{stamp}")
    with open(base + ".csv", "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f)
        for k, v in header_info:
            w.writerow([k, v])
        w.writerow([])
        w.writerow([c[0] for c in columns])
        for r in rows:
            w.writerow([r.get(c[1], "") for c in columns])
    wb = openpyxl.Workbook(); ws = wb.active; ws.sheet_view.rightToLeft = True
    for k, v in header_info:
        ws.append([k, v])
    ws.append([]); ws.append([c[0] for c in columns])
    for r in rows:
        ws.append([r.get(c[1], "") for c in columns])
    wb.save(base + ".xlsx")
    return base + ".xlsx", base + ".csv"


def salary_transfer_rows(run_id):
    run = services.get_payroll_run(int(run_id))
    if not run:
        raise ValueError("دورة الرواتب غير موجودة")
    rows = []
    for ps in run.payslips:
        e = ps.employee
        iban = re.sub(r"\s+", "", e.bank_account or "").upper()
        rows.append({"emp_id": e.id, "code": e.code, "national_id": e.national_id or "", "name": e.name, "iban": iban, "valid": validate_iban(iban), "net": round(ps.net_salary or 0, 2),
                     "basic": round(ps.basic_salary or 0, 2), "allowances": round((ps.housing_allowance or 0) + (ps.transport_allowance or 0) + (ps.other_allowances or 0), 2),
                     "deductions": round((ps.gosi_deduction or 0) + (ps.other_deductions or 0), 2)})
    return run, rows


def export_salary_file(run_id, bank_account_id):
    """ملف رواتب (WPS عام): رقم الهوية، الاسم، الآيبان، الأساسي، البدلات، الخصومات، الصافي. الموظف بدون آيبان صحيح بيتوقف التصدير وتتحدد أسماؤهم."""
    services.require_action("banking", "post")
    run, rows = salary_transfer_rows(run_id)
    bad = [r["name"] for r in rows if r["net"] > 0 and not r["valid"]]
    if bad:
        raise ValueError("آيبان غير صحيح أو ناقص للموظفين: " + "، ".join(bad[:10]) + (" ..." if len(bad) > 10 else "") + " - صحّحه من بيانات الموظف")
    rows = [r for r in rows if r["net"] > 0]
    if not rows:
        raise ValueError("مفيش رواتب صافية للتحويل")
    b = get_bank_account(bank_account_id)
    total = round(sum(r["net"] for r in rows), 2)
    head = [("المنشأة", services.get_company_settings().name or ""), ("رقم المنشأة لدى البنك", b.establishment_id), ("آيبان حساب الخصم", b.iban), ("البنك", b.bank_name),
            ("شهر الرواتب", f"{run.month:02d}/{run.year}"), ("عدد الموظفين", len(rows)), ("إجمالي المبلغ", f"{total:.2f}"), ("تاريخ الإنشاء", datetime.date.today().isoformat())]
    cols = [("رقم الهوية/الإقامة", "national_id"), ("اسم الموظف", "name"), ("الآيبان", "iban"), ("الراتب الأساسي", "basic"), ("البدلات", "allowances"), ("الخصومات", "deductions"), ("صافي الراتب", "net")]
    return _write_transfer_files(f"salary_{run.year}{run.month:02d}", b, head, rows, cols) + (total, len(rows))


def get_bank_account(bank_account_id):
    with get_session() as s:
        b = s.query(BankAccount).options(joinedload(BankAccount.account)).filter(BankAccount.id == int(bank_account_id)).first()
        if not b:
            raise ValueError("الحساب البنكي غير موجود")
        s.expunge_all()
        return b


def supplier_balance(supplier_id):
    try:
        return float(services.party_statement("supplier", supplier_id)["closing"])
    except Exception:
        return 0.0


def export_supplier_transfers(items, bank_account_id, purpose="سداد مستحقات"):
    """items = [{supplier_id, amount, iban}] - بيحدّث آيبان المورد لو اتكتب، وبيرفض آيبان غلط أو مبلغ صفر."""
    services.require_action("banking", "post")
    if not items:
        raise ValueError("اختر مورد واحد على الأقل")
    rows = []
    for it in items:
        amt = round(float(it.get("amount") or 0), 2)
        if amt <= 0:
            continue
        iban = re.sub(r"\s+", "", it.get("iban") or "").upper()
        with get_session() as s:
            sup = s.get(Supplier, int(it["supplier_id"]))
            if not iban:
                iban = (sup.iban or "")
            if not validate_iban(iban):
                raise ValueError(f"آيبان المورد «{sup.name}» غير صحيح أو ناقص")
            sup.iban = iban; s.commit()
            rows.append({"name": sup.name, "iban": iban, "bank": sup.bank_name or "", "amount": amt, "purpose": purpose, "ref": f"SUP{sup.id}"})
    if not rows:
        raise ValueError("كل المبالغ صفر")
    b = get_bank_account(bank_account_id)
    total = round(sum(r["amount"] for r in rows), 2)
    head = [("المنشأة", services.get_company_settings().name or ""), ("رقم المنشأة لدى البنك", b.establishment_id), ("آيبان حساب الخصم", b.iban), ("البنك", b.bank_name),
            ("عدد التحويلات", len(rows)), ("إجمالي المبلغ", f"{total:.2f}"), ("تاريخ الإنشاء", datetime.date.today().isoformat())]
    cols = [("اسم المستفيد", "name"), ("آيبان المستفيد", "iban"), ("البنك", "bank"), ("المبلغ", "amount"), ("الغرض", "purpose"), ("المرجع", "ref")]
    return _write_transfer_files("supplier_transfers", b, head, rows, cols) + (total, len(rows))
