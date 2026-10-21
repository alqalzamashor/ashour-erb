"""تقرير الرواتب (مدفوع / غير مدفوع بالشهور) وصرف غير المدفوع منه."""
import datetime
from unittest.mock import MagicMock

import pytest
from flet.controls.context import _context_page

D = datetime.datetime


def _emp(S, name):
    a = [x.id for x in S.list_accounts(postable_only=True)]
    eid = S.create_employee(name, hire_date=D(2031, 1, 1))
    S.create_contract(eid, D(2031, 1, 1), 3000, housing_allowance=0, transport_allowance=0, other_allowances=0,
                      gosi_rate_employee=0, gosi_rate_employer=0, leave_provision_rate=0, eos_provision_rate=0,
                      basic_salary_account_id=a[0], gosi_deduction_account_id=a[1], other_deductions_account_id=a[1])
    return eid


def test_report_filters_and_multi_month_payment(S, tmp_path):
    e1, e2 = _emp(S, "رواتب أ"), _emp(S, "رواتب ب")
    r1 = S.create_payroll_run(1, 2031); S.post_payroll_run(r1)
    r2 = S.create_payroll_run(2, 2031)                       # فبراير مسودة
    rep = S.payroll_payments_report(2031)
    assert rep["totals"]["count"] == 4 and rep["totals"]["unpaid"] == rep["totals"]["total"] and len(rep["months"]) == 2
    assert all(len(S.payroll_payments_report(2031, m)["rows"]) == 2 for m in (1, 2))
    unpaid = [r for r in rep["rows"] if r["month"] == 1]
    # مينفعش صرف قسيمة من دورة مش مترحّلة
    feb = [r["payslip_id"] for r in rep["rows"] if r["month"] == 2]
    with pytest.raises(ValueError, match="ما اترحّلتش"):
        S.pay_payslips_multi(feb, 1)
    S.post_payroll_run(r2)
    acc = [a.id for a in S.list_accounts(postable_only=True)][0]
    pm = S.create_payment_method("بنك تقرير الرواتب", account_id=acc)
    pm_id = pm if isinstance(pm, int) else pm.id
    ids = [r["payslip_id"] for r in S.payroll_payments_report(2031, status="unpaid")["rows"]]
    one_each = [unpaid[0]["payslip_id"]] + feb[:1]
    res = S.pay_payslips_multi(one_each, pm_id, payment_date=D(2031, 3, 1))
    assert res["employee_count"] == 2 and len(res["payments"]) == 2 and res["amount"] > 0
    paid = S.payroll_payments_report(2031, status="paid")
    assert paid["totals"]["count"] == 2 and all(r["payment_number"] != "-" for r in paid["rows"])
    left = S.payroll_payments_report(2031, status="unpaid")
    assert left["totals"]["count"] == 2 and set(r["payslip_id"] for r in left["rows"]) == set(ids) - set(one_each)
    with pytest.raises(ValueError, match="اتصرفت بالفعل"):
        S.pay_payslips_multi(one_each, pm_id)
    with pytest.raises(ValueError):
        S.payroll_payments_report(2031, status="x")

    from erp_app.views.payroll_report import build_payroll_report, export_payroll_report_xlsx
    from tests.test_view_controls import _bad
    pg = MagicMock(); pg.overlay = []; _context_page.set(pg)
    root = build_payroll_report(pg)
    assert _bad(root) == []
    from openpyxl import load_workbook
    wb = load_workbook(export_payroll_report_xlsx(S.payroll_payments_report(2031), str(tmp_path / "p.xlsx")))
    assert wb.worksheets[0].max_row >= 6 and wb.worksheets[1].max_row == 3


def test_voucher_preview(S):
    import os
    import flet as ft
    e = _emp(S, "معاينة سند")
    rid = S.create_payroll_run(5, 2032); S.post_payroll_run(rid)
    acc = [a.id for a in S.list_accounts(postable_only=True)][0]
    pm = S.create_payment_method("بنك معاينة", account_id=acc)
    S.pay_payroll_run(rid, pm if isinstance(pm, int) else pm.id)
    rows = [r for r in S.payroll_payments_report(2032)["rows"] if r["paid"]]
    assert rows and all(r["payment_id"] for r in rows)
    pdf = S.generate_payment_pdf_file(rows[0]["payment_id"])
    assert os.path.getsize(pdf) > 800
    from erp_app.views.payroll_report import build_payroll_report
    from tests.test_drafts import _walk
    pg = MagicMock(); pg.overlay = []; _context_page.set(pg)
    root = build_payroll_report(pg)
    year = [c for c in _walk(root) if isinstance(c, ft.TextField) and c.label == "السنة"][0]
    year.value = "2032"; year.on_change(None)
    btns = [c for c in _walk(root) if isinstance(c, ft.IconButton) and c.tooltip == "معاينة سند الصرف"]
    assert btns


def test_open_or_download_uses_coroutine(monkeypatch):
    """page.launch_url مغلّف بـ @deprecated (ليس coroutine) فيجب تمريره لـ run_task داخل async."""
    import inspect
    from unittest.mock import MagicMock
    from erp_app import theme, files
    monkeypatch.setattr(files, 'publish_download', lambda p: 'http://x/y.pdf')
    page = MagicMock()
    theme.open_or_download_file(page, '/tmp/a.pdf')
    fn = page.run_task.call_args[0][0]
    assert inspect.iscoroutinefunction(fn)
    page.launch_url.return_value = None
    import asyncio
    asyncio.run(fn())
    page.launch_url.assert_called_once_with('http://x/y.pdf')


def test_payroll_report_pdf(S, tmp_path):
    import os
    from erp_app.views.payroll_report import export_payroll_report_pdf
    p = export_payroll_report_pdf(S.payroll_payments_report(2031), str(tmp_path / "p.pdf"))
    assert os.path.getsize(p) > 800
