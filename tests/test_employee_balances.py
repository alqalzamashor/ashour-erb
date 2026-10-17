"""تقرير أرصدة الموظفين (إجازات ونهاية خدمة)."""
import datetime
from unittest.mock import MagicMock

from flet.controls.context import _context_page

D = datetime.datetime


def test_report_and_view_and_export(S, tmp_path):
    accs = [a.id for a in S.list_accounts(postable_only=True)]
    eid = S.create_employee("موظف تقرير", hire_date=D(2024, 1, 1))
    S.create_contract(eid, D(2024, 1, 1), 3000, housing_allowance=0, transport_allowance=0, other_allowances=0,
                      gosi_rate_employee=0, gosi_rate_employer=0, leave_provision_rate=0, eos_provision_rate=0,
                      basic_salary_account_id=accs[0], gosi_deduction_account_id=accs[1], other_deductions_account_id=accs[1])
    nocon = S.create_employee("موظف بلا عقد", hire_date=D(2024, 1, 1))
    rid = S.create_leave_request(eid, "annual", D(2025, 3, 1), D(2025, 3, 5)); S.decide_leave_request(rid, True)
    rows = {r["employee_id"]: r for r in S.employee_balances_report(D(2025, 12, 31))}
    r = rows[eid]
    assert r["has_contract"] and r["leave_used"] == 5
    assert r["leave_days"] == round(r["leave_earned"] - 5, 2) and r["leave_amount"] == round(r["leave_days"] * 100, 2)
    assert r["eos_termination"] >= r["eos_resignation"] >= 0
    assert not rows[nocon]["has_contract"] and rows[nocon]["leave_amount"] == 0

    from erp_app.views.employee_balances import build_employee_balances, export_balances_xlsx
    pg = MagicMock(); pg.overlay = []; _context_page.set(pg)
    assert build_employee_balances(pg) is not None
    out = export_balances_xlsx(list(rows.values()), D(2025, 12, 31), str(tmp_path / "b.xlsx"))
    from openpyxl import load_workbook
    ws = load_workbook(out).active
    assert ws.max_row >= len(rows) + 3
