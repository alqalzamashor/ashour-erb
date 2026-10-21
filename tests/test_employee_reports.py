"""تقارير الموظفين."""
import datetime
from unittest.mock import MagicMock

import pytest
from flet.controls.context import _context_page

D = datetime.datetime


def _setup(S):
    accs = [a.id for a in S.list_accounts(postable_only=True)]
    eid = S.create_employee("موظف تقارير", hire_date=D(2024, 1, 1))
    end = D.combine(datetime.date.today() + datetime.timedelta(days=20), datetime.time())
    S.create_contract(eid, D(2024, 1, 1), 3000, end_date=end, housing_allowance=500, transport_allowance=0, other_allowances=0,
                      gosi_rate_employee=0, gosi_rate_employer=0, leave_provision_rate=0, eos_provision_rate=0,
                      basic_salary_account_id=accs[0], housing_allowance_account_id=accs[0],
                      gosi_deduction_account_id=accs[1], other_deductions_account_id=accs[1])
    rid = S.create_leave_request(eid, "annual", D(datetime.date.today().year, 1, 10), D(datetime.date.today().year, 1, 12))
    S.decide_leave_request(rid, True)
    return eid


def test_all_reports(S, tmp_path):
    eid = _setup(S)
    t = datetime.date.today()
    for k, _ in S.EMPLOYEE_REPORTS:
        rep = S.employee_report(k, employee_id=eid, date_from=D(t.year, 1, 1), date_to=D(t.year, 12, 31), year=t.year, days=60)
        assert rep["title"] and rep["columns"]
        assert all(len(r) == len(rep["columns"]) for r in rep["rows"]), k
        if rep.get("totals"):
            assert len(rep["totals"]) == len(rep["columns"]), k
    r = S.employee_report("roster")
    assert any(row[1] == "موظف تقارير" and row[8] == 3500 for row in r["rows"])
    h = S.employee_report("headcount")
    assert h["totals"][2] >= 1 and h["totals"][4] >= 3500
    c = S.employee_report("contracts_expiring", days=30)
    assert any(row[1] == "موظف تقارير" and 19 <= row[6] <= 21 for row in c["rows"])
    assert not any(row[1] == "موظف تقارير" for row in S.employee_report("contracts_expiring", days=5)["rows"])
    lv = S.employee_report("leaves", date_from=D(t.year, 1, 1), date_to=D(t.year, 12, 31))
    assert lv["rows"] and lv["totals"][-1] == 3
    mv = S.employee_report("movements", date_from=D(2024, 1, 1), date_to=D(2024, 1, 31))
    assert any(row[1] == "تعيين" for row in mv["rows"])
    with pytest.raises(ValueError):
        S.employee_report("statement")
    with pytest.raises(ValueError):
        S.employee_report("nope")
    with pytest.raises(ValueError):
        S.employee_report("roster", date_from=D(2025, 2, 1), date_to=D(2025, 1, 1))


def test_view_and_export(S, tmp_path):
    _setup(S)
    from erp_app.views.employee_reports import build_employee_reports, export_report_xlsx
    pg = MagicMock(); pg.overlay = []; _context_page.set(pg)
    assert build_employee_reports(pg) is not None
    from openpyxl import load_workbook
    rep = S.employee_report("roster")
    ws = load_workbook(export_report_xlsx(rep, str(tmp_path / "r.xlsx"))).active
    assert ws.max_row >= len(rep["rows"]) + 3


def test_every_report_exports_pdf(S, tmp_path):
    """كل أنواع تقارير الموظفين بتتصدّر PDF سليم (أعمدة + صفوف + إجمالي)."""
    import os
    eid = _setup(S)
    t = datetime.date.today()
    from erp_app.views.employee_reports import export_report_pdf
    for kind, _label in S.EMPLOYEE_REPORTS:
        rep = S.employee_report(kind, employee_id=eid, date_from=D(t.year, 1, 1), date_to=D(t.year, 12, 31), year=t.year, days=60)
        p = export_report_pdf(rep, str(tmp_path / f"{kind}.pdf"))
        assert os.path.getsize(p) > 800, kind
